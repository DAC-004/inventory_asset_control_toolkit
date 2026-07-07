from pathlib import Path
from openpyxl import load_workbook

WORKBOOK_PATH = Path(r"C:\Users\cruzd\OneDrive\Desktop\Interview\Inventory_IT_Asset_Control_Toolkit.xlsx")
REPORT_PATH = Path(r"C:\Users\cruzd\OneDrive\Desktop\Interview\audit_report.txt")

EXPECTED_SHEETS = [
    "README",
    "Master Inventory",
    "Inventory Dashboard",
    "Aged Excess Analysis",
    "Markdown Planner",
    "Transfer Planner",
    "IT Asset Register",
    "Audit Reconciliation",
    "Software Licenses",
    "Mobile Provisioning",
    "Disposal Log",
    "Management Summary",
]

def write_line(lines, text=""):
    print(text)
    lines.append(text)

def audit_workbook():
    lines = []

    write_line(lines, "=" * 80)
    write_line(lines, "WORKBOOK AUDIT REPORT")
    write_line(lines, "=" * 80)
    write_line(lines, f"Workbook path: {WORKBOOK_PATH}")

    if not WORKBOOK_PATH.exists():
        write_line(lines, f"ERROR: Workbook not found: {WORKBOOK_PATH}")
        REPORT_PATH.write_text("\n".join(lines), encoding="utf-8")
        return

    write_line(lines, f"File size: {WORKBOOK_PATH.stat().st_size / 1024:.2f} KB")

    try:
        wb = load_workbook(WORKBOOK_PATH, data_only=False)
    except Exception as e:
        write_line(lines, f"ERROR opening workbook: {e}")
        REPORT_PATH.write_text("\n".join(lines), encoding="utf-8")
        return

    write_line(lines, "")
    write_line(lines, "SHEET ORDER")
    write_line(lines, "-" * 80)

    for i, sheet in enumerate(wb.sheetnames, start=1):
        expected = EXPECTED_SHEETS[i - 1] if i <= len(EXPECTED_SHEETS) else "Unexpected"
        status = "OK" if sheet == expected else f"CHECK - expected '{expected}'"
        write_line(lines, f"{i:02d}. {sheet} — {status}")

    missing = [s for s in EXPECTED_SHEETS if s not in wb.sheetnames]
    unexpected = [s for s in wb.sheetnames if s not in EXPECTED_SHEETS]

    write_line(lines, "")
    write_line(lines, "MISSING EXPECTED SHEETS")
    write_line(lines, "-" * 80)
    write_line(lines, "None" if not missing else "\n".join(missing))

    write_line(lines, "")
    write_line(lines, "UNEXPECTED SHEETS")
    write_line(lines, "-" * 80)
    write_line(lines, "None" if not unexpected else "\n".join(unexpected))

    total_formulas = 0

    write_line(lines, "")
    write_line(lines, "SHEET DETAILS")
    write_line(lines, "-" * 80)

    for ws in wb.worksheets:
        formulas = []
        non_empty_cells = 0

        for row in ws.iter_rows():
            for cell in row:
                if cell.value not in (None, ""):
                    non_empty_cells += 1
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas.append((cell.coordinate, cell.value))

        total_formulas += len(formulas)

        write_line(lines, "")
        write_line(lines, f"Sheet: {ws.title}")
        write_line(lines, f"  Dimensions: {ws.max_row} rows x {ws.max_column} columns")
        write_line(lines, f"  Non-empty cells: {non_empty_cells}")
        write_line(lines, f"  Tables: {len(ws.tables)}")
        write_line(lines, f"  Charts: {len(ws._charts)}")
        write_line(lines, f"  Frozen panes: {ws.freeze_panes}")
        write_line(lines, f"  AutoFilter: {ws.auto_filter.ref}")
        write_line(lines, f"  Formulas: {len(formulas)}")

        if formulas:
            write_line(lines, "  Formula samples:")
            for coord, formula in formulas[:10]:
                write_line(lines, f"    {coord}: {formula}")
            if len(formulas) > 10:
                write_line(lines, f"    ... {len(formulas) - 10} more formulas")

    write_line(lines, "")
    write_line(lines, "SUMMARY")
    write_line(lines, "-" * 80)
    write_line(lines, f"Total sheets: {len(wb.sheetnames)}")
    write_line(lines, f"Total formulas: {total_formulas}")
    write_line(lines, "Audit complete.")

    REPORT_PATH.write_text("\n".join(lines), encoding="utf-8")
    write_line(lines, "")
    print(f"Report saved to: {REPORT_PATH}")

if __name__ == "__main__":
    print("Starting workbook audit...")
    audit_workbook()
