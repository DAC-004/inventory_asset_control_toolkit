"""Expanded tests for specs.md §16 data quality and sheet structure."""

from config.workbook_config import SHEET_ORDER
from openpyxl import load_workbook

from src.data_generation.generate_disposal import generate_disposal_data
from src.data_generation.generate_inventory import generate_inventory_data
from src.data_generation.generate_it_assets import generate_it_asset_data
from src.main import generate_all_data
from src.sheets.master_inventory_sheet import HEADERS as INVENTORY_HEADERS, HEADER_ROW as INV_HEADER_ROW
from src.sheets.it_asset_register_sheet import HEADERS as IT_HEADERS, HEADER_ROW as IT_HEADER_ROW
from src.workbook.builder import build_workbook


def test_inventory_duplicate_primary_keys():
    """Inventory item IDs must be unique."""
    df = generate_inventory_data(row_count=120, seed=42)
    assert df["item_id"].is_unique


def test_it_assets_multiple_lifecycle_stages():
    """IT asset data should include multiple lifecycle stages."""
    df = generate_it_asset_data(row_count=120, seed=42)
    assert df["lifecycle_stage"].nunique() >= 4


def test_disposal_includes_pending_wipe():
    """Disposal data must include at least one pending wipe record."""
    df = generate_disposal_data(row_count=25, seed=42)
    assert "Pending Wipe" in df["disposal_status"].values


def test_required_sheets_have_headers(tmp_path):
    """Operational sheets should expose styled header rows."""
    data = generate_all_data()
    output = tmp_path / "workbook.xlsx"
    build_workbook(data, output_path=output)
    wb = load_workbook(output)

    assert wb["Master Inventory"].cell(row=INV_HEADER_ROW, column=1).value == INVENTORY_HEADERS[0]
    assert wb["IT Asset Register"].cell(row=IT_HEADER_ROW, column=1).value == IT_HEADERS[0]
    assert wb.sheetnames == SHEET_ORDER
