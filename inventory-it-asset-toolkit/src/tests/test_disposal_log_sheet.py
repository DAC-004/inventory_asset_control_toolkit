"""Tests for Disposal Log sheet builder."""

from openpyxl import Workbook

from src.data_generation.generate_disposal import generate_disposal_data
from src.sheets.disposal_log_sheet import (
    DATA_START_ROW,
    HEADER_ROW,
    HEADERS,
    SUMMARY_TITLE_ROW,
    TABLE_NAME,
    _compute_summary,
    build,
)


def test_compute_summary_metrics():
    """Summary should calculate disposal pipeline counts."""
    df = generate_disposal_data(row_count=25, seed=42)
    summary = _compute_summary(df)

    assert summary["total_records"] == 25
    assert summary["pending_wipe"] >= 1
    assert summary["certificate_missing"] >= 1
    assert summary["disposed"] >= 1
    assert summary["hold_for_review"] >= 1
    assert (
        summary["pending_wipe"]
        + summary["certificate_missing"]
        + summary["disposed"]
        + summary["hold_for_review"]
        <= summary["total_records"]
    )


def test_disposal_log_sheet_structure():
    """Sheet should include summary, table, formatting, and validation."""
    wb = Workbook()
    ws = wb.active
    df = generate_disposal_data(row_count=25, seed=42)
    build(ws, {"data": {"disposal": df}})

    assert ws.cell(row=HEADER_ROW, column=1).value == HEADERS[0]
    assert ws.cell(row=SUMMARY_TITLE_ROW, column=1).value
    assert TABLE_NAME in ws.tables
    assert len(ws.conditional_formatting) >= 1
    assert len(ws.data_validations.dataValidation) >= 4
    assert ws[f"E{DATA_START_ROW}"].number_format
    assert ws[f"L{DATA_START_ROW}"].number_format
    assert ws.sheet_view.showGridLines is False
