"""Tests for workbook utility modules."""

from openpyxl import Workbook
from openpyxl.chart import Reference

from config import style_config as sc
from src.workbook.builder import SHEET_BUILDERS, create_workbook
from src.workbook.charts import create_bar_chart, create_pie_chart, make_category_reference
from src.workbook.formulas import count_if_range, margin_pct_formula, sum_if_numeric, sum_range
from src.workbook.utils import (
    apply_auto_filter,
    autosize_columns,
    create_excel_table,
    format_currency_columns,
    format_date_columns,
    format_percentage_columns,
    set_print_layout,
)
from src.workbook.validations import (
    add_compliance_status_validation,
    add_disposal_status_validation,
    add_inventory_status_validation,
    add_lifecycle_stage_validation,
)


def test_create_workbook_has_no_default_sheet():
    """create_workbook should return a workbook without the default sheet."""
    wb = create_workbook()
    assert len(wb.sheetnames) == 0
    assert wb.properties.title is not None


def test_formula_helpers_return_strings():
    """Formula helpers should return Excel formula strings."""
    assert sum_range("Master Inventory", "N", 3, 100).startswith("=SUM(")
    assert count_if_range("Master Inventory", "U", 3, 100, "Healthy").startswith("=COUNTIF(")
    assert margin_pct_formula("Master Inventory", "M", "L", 3).startswith("=IF(")


def test_sum_if_numeric_quotes_expression_criteria():
    """SUMIF numeric criteria must be quoted for Excel compatibility."""
    formula = sum_if_numeric("Master Inventory", "Q", ">180", "N", 3, 122)
    assert '">180"' in formula
    assert ",>180," not in formula


def test_utils_format_and_table_helpers():
    """Utility helpers should apply formats and create tables without error."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Test"
    ws["A1"] = "Header"
    ws["A2"] = 100.5
    ws["B2"] = 0.25
    ws["C2"] = "2026-07-01"

    format_currency_columns(ws, ["A"], 2, 2)
    format_percentage_columns(ws, ["B"], 2, 2)
    format_date_columns(ws, ["C"], 2, 2)
    create_excel_table(ws, "TestTable", "A1:C2")
    apply_auto_filter(ws, "A1:C2")
    autosize_columns(ws)
    set_print_layout(ws, orientation="landscape", repeat_header_rows="1:1")

    assert ws["A2"].number_format == sc.NUMBER_FORMATS["currency"]
    assert len(ws.tables) == 1


def test_validation_helpers_register_rules():
    """Validation helpers should add data validation rules to the worksheet."""
    wb = Workbook()
    ws = wb.active
    add_inventory_status_validation(ws, "U3:U50")
    add_lifecycle_stage_validation(ws, "O3:O50")
    add_compliance_status_validation(ws, "L3:L50")
    add_disposal_status_validation(ws, "N3:N50")
    assert len(ws.data_validations.dataValidation) == 4


def test_chart_helpers_build_charts():
    """Chart helpers should return configured chart objects."""
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Category"
    ws["B1"] = "Value"
    ws["A2"] = "Healthy"
    ws["B2"] = 10
    ws["A3"] = "Excess"
    ws["B3"] = 5

    cats = make_category_reference(ws, 1, 2, 3)
    data = Reference(ws, min_col=2, min_row=1, max_col=2, max_row=3)
    bar = create_bar_chart("Test Bar", data, cats)
    pie = create_pie_chart("Test Pie", data, cats)
    assert bar.title.tx.rich.paragraphs[0].r[0].t == "Test Bar"
    assert pie.title.tx.rich.paragraphs[0].r[0].t == "Test Pie"


def test_sheet_builders_cover_all_required_sheets():
    """Every sheet in SHEET_ORDER should have a registered builder."""
    from config.workbook_config import SHEET_ORDER

    assert set(SHEET_BUILDERS.keys()) == set(SHEET_ORDER)
