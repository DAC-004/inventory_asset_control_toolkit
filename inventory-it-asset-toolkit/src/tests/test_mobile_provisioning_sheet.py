"""Tests for Mobile Provisioning sheet builder."""

from openpyxl import Workbook

from src.data_generation.generate_mobile import generate_mobile_data
from src.sheets.mobile_provisioning_sheet import (
    DATA_START_ROW,
    HEADER_ROW,
    HEADERS,
    SUMMARY_TITLE_ROW,
    TABLE_NAME,
    _compute_summary,
    build,
)


def test_compute_summary_metrics():
    """Summary should calculate fleet checklist counts."""
    df = generate_mobile_data(row_count=45, seed=42)
    summary = _compute_summary(df)

    assert summary["total_devices"] == 45
    assert summary["assigned"] >= 1
    assert summary["pending_setup"] >= 1
    assert summary["missing_agreement"] >= 1
    assert summary["returned"] >= 1
    assert (
        summary["assigned"]
        + summary["pending_setup"]
        + summary["missing_agreement"]
        + summary["returned"]
        <= summary["total_devices"]
    )


def test_mobile_provisioning_sheet_structure():
    """Sheet should include summary, table, formatting, and validation."""
    wb = Workbook()
    ws = wb.active
    df = generate_mobile_data(row_count=45, seed=42)
    build(ws, {"data": {"mobile": df}})

    assert ws.cell(row=HEADER_ROW, column=1).value == HEADERS[0]
    assert ws.cell(row=SUMMARY_TITLE_ROW, column=1).value
    assert TABLE_NAME in ws.tables
    assert len(ws.conditional_formatting) >= 1
    assert len(ws.data_validations.dataValidation) >= 5
    assert ws[f"M{DATA_START_ROW}"].number_format
    assert ws[f"N{DATA_START_ROW}"].number_format
    assert ws.sheet_view.showGridLines is False
