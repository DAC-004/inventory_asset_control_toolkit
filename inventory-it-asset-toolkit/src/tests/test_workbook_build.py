"""
Tests for workbook build process.

TODO: Expand coverage per specs.md §16 once sheet builders are implemented.
"""

from pathlib import Path

from config.workbook_config import SHEET_ORDER, WORKBOOK_PATH
from openpyxl import load_workbook

from src.main import generate_all_data
from src.workbook.builder import build_workbook


def test_build_workbook_creates_file(tmp_path: Path):
    """build_workbook should write an .xlsx file to the output path."""
    data = generate_all_data()
    output = tmp_path / "test_workbook.xlsx"
    result = build_workbook(data, output_path=output)
    assert result.exists()
    assert result.suffix == ".xlsx"


def test_workbook_contains_required_sheets(tmp_path: Path):
    """Built workbook should include every sheet from SHEET_ORDER."""
    data = generate_all_data()
    output = tmp_path / "test_workbook.xlsx"
    build_workbook(data, output_path=output)
    wb = load_workbook(output, read_only=True)
    assert wb.sheetnames == SHEET_ORDER
