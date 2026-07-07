"""Tests for Master Inventory sheet builder."""

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from src.data_generation.generate_inventory import generate_inventory_data
from src.sheets.master_inventory_sheet import (
    COL_COUNT,
    DATA_START_ROW,
    HEADER_ROW,
    HEADERS,
    TABLE_NAME,
    build,
)


def test_master_inventory_sheet_structure():
    """Master Inventory sheet should have title, headers, table, and data."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Master Inventory"
    df = generate_inventory_data(row_count=50, seed=42)

    build(ws, {"data": {"inventory": df}})

    assert ws.cell(row=HEADER_ROW, column=1).value == HEADERS[0]
    assert ws.cell(row=HEADER_ROW, column=COL_COUNT).value == HEADERS[-1]
    assert ws.cell(row=DATA_START_ROW, column=1).value is not None
    assert TABLE_NAME in ws.tables
    assert len(ws.data_validations.dataValidation) >= 2
    assert len(ws.conditional_formatting) >= 2


def test_master_inventory_column_formats():
    """Currency, date, and percentage columns should have number formats applied."""
    wb = Workbook()
    ws = wb.active
    df = generate_inventory_data(row_count=10, seed=42)
    build(ws, {"data": {"inventory": df}})

    row = DATA_START_ROW
    assert ws[f"L{row}"].number_format  # unit cost
    assert "$" in ws[f"L{row}"].number_format or "#" in ws[f"L{row}"].number_format
    assert "%" in ws[f"S{row}"].number_format
    assert ws[f"O{row}"].number_format  # date


def test_master_inventory_freeze_panes():
    """Header row should be frozen below row 2."""
    wb = Workbook()
    ws = wb.active
    df = generate_inventory_data(row_count=5, seed=42)
    build(ws, {"data": {"inventory": df}})

    assert ws.freeze_panes == f"A{DATA_START_ROW}"
