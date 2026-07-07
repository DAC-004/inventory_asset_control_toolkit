"""Tests for Management Summary sheet builder."""

from openpyxl import Workbook

from src.data_generation.generate_inventory import generate_inventory_data
from src.data_generation.generate_it_assets import generate_it_asset_data
from src.main import generate_all_data
from src.sheets.management_summary_sheet import (
    KPI1_VALUE_ROW,
    RISK_DATA_START_ROW,
    SECTION1_ROW,
    SECTION3C_ROW,
    SECTION3B_ROW,
    SECTION5_ROW,
    _top_inventory_risks,
    _top_it_asset_risks,
    build,
)


def test_top_risk_helpers_return_rows():
    """Top risk helpers should return ranked summary rows."""
    inventory = generate_inventory_data(row_count=120, seed=42)
    assets = generate_it_asset_data(row_count=120, seed=42)

    inv_risks = _top_inventory_risks(inventory)
    it_risks = _top_it_asset_risks(assets)

    assert 1 <= len(inv_risks) <= 5
    assert 1 <= len(it_risks) <= 5
    assert "SKU" in inv_risks[0]
    assert "Asset Tag" in it_risks[0]


def test_management_summary_sheet_structure():
    """Sheet should include all executive sections and print settings."""
    wb = Workbook()
    ws = wb.active
    data = generate_all_data()
    build(ws, {"data": data})

    assert "Inventory Health" in str(ws.cell(row=SECTION1_ROW, column=1).value)
    assert "Disposal / Lifecycle" in str(ws.cell(row=SECTION3B_ROW, column=1).value)
    assert "Executive Summaries" in str(ws.cell(row=SECTION3C_ROW, column=1).value)
    assert "30 / 60 / 90" in str(ws.cell(row=SECTION5_ROW, column=1).value)
    assert str(ws.cell(row=KPI1_VALUE_ROW, column=1).value).startswith("=")
    assert ws.cell(row=RISK_DATA_START_ROW, column=1).value
    assert ws.cell(row=RISK_DATA_START_ROW, column=6).value
    assert ws.sheet_view.showGridLines is False
    assert ws.page_setup.orientation == ws.ORIENTATION_LANDSCAPE
    assert ws.page_setup.fitToWidth == 1
    assert ws.page_setup.fitToHeight == 1
