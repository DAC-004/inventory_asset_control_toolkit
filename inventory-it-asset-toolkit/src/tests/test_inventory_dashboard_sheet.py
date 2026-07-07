"""Tests for Inventory Dashboard sheet builder."""

from openpyxl import Workbook

from src.data_generation.generate_inventory import generate_inventory_data
from src.sheets.inventory_dashboard_sheet import build


def test_inventory_dashboard_kpi_cards():
    """Dashboard should render 10 KPI cards with formulas."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventory Dashboard"
    df = generate_inventory_data(row_count=50, seed=42)

    build(ws, {"data": {"inventory": df}})

    assert ws.cell(row=3, column=1).value == "Total Inventory Value"
    assert str(ws.cell(row=4, column=1).value).startswith("=SUM(")
    assert ws.cell(row=6, column=9).value == "Average Gross Margin %"
    assert len(ws._charts) >= 4


def test_inventory_dashboard_summary_tables():
    """Dashboard should include all required summary table sections."""
    wb = Workbook()
    ws = wb.active
    df = generate_inventory_data(row_count=50, seed=42)
    build(ws, {"data": {"inventory": df}})

    values = [
        ws.cell(row=r, column=1).value
        for r in range(1, ws.max_row + 1)
        if ws.cell(row=r, column=1).value
    ]
    combined = " ".join(str(v) for v in values)
    assert "Inventory Value by Location" in combined
    assert "Inventory Status Breakdown" in combined
    assert "Aging Bucket Summary" in combined
    assert "Recommended Action Summary" in combined
    assert "Top 10 Excess Inventory Items" in combined


def test_inventory_dashboard_gridlines_hidden():
    """Dashboard should hide gridlines for a clean executive layout."""
    wb = Workbook()
    ws = wb.active
    df = generate_inventory_data(row_count=20, seed=42)
    build(ws, {"data": {"inventory": df}})

    assert ws.sheet_view.showGridLines is False
