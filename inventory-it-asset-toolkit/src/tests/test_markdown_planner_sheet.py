"""Tests for Markdown Planner sheet builder."""

from openpyxl import Workbook

from config.workbook_config import MARKDOWN_THRESHOLDS
from src.data_generation.generate_inventory import generate_inventory_data
from src.sheets.markdown_planner_sheet import (
    DATA_START_ROW,
    HEADER_ROW,
    HEADERS,
    TABLE_NAME,
    _assign_markdown_plan,
    _build_markdown_dataframe,
    build,
)


def test_assign_markdown_plan_rules():
    """Markdown rules should follow age and demand thresholds."""
    assert _assign_markdown_plan(400, 0, "Obsolete") == (
        MARKDOWN_THRESHOLDS["liquidate_markdown_pct"],
        "Liquidate",
    )
    assert _assign_markdown_plan(300, 5, "Slow-Moving") == (
        MARKDOWN_THRESHOLDS["markdown_20_pct"],
        "20% Markdown",
    )
    assert _assign_markdown_plan(200, 5, "Excess / Aged") == (
        MARKDOWN_THRESHOLDS["markdown_10_pct"],
        "10% Markdown",
    )
    assert _assign_markdown_plan(90, 10, "Excess") == (0.0, "Transfer First")
    assert _assign_markdown_plan(90, 10, "Slow-Moving") == (0.0, "Hold")


def test_build_markdown_dataframe_filters_and_calculates():
    """Planner should include eligible statuses and compute markdown price."""
    df = generate_inventory_data(row_count=120, seed=42)
    planner = _build_markdown_dataframe(df)

    assert list(planner.columns) == HEADERS
    assert planner["Recommended Disposition"].notna().all()
    assert set(planner["SKU"]).issubset(set(df["sku"]))
    assert planner["Suggested Markdown %"].between(0, 1).all()
    assert (planner["Markdown Price"] <= planner["Current Selling Price"]).all()
    assert (planner["Estimated Recovery Value"] >= 0).all()

    sample = planner.iloc[0]
    expected_price = round(
        sample["Current Selling Price"] * (1 - sample["Suggested Markdown %"]),
        2,
    )
    assert sample["Markdown Price"] == expected_price


def test_markdown_planner_sheet_structure():
    """Sheet should include table and disposition conditional formatting."""
    wb = Workbook()
    ws = wb.active
    df = generate_inventory_data(row_count=50, seed=42)
    build(ws, {"data": {"inventory": df}})

    assert ws.cell(row=HEADER_ROW, column=1).value == "SKU"
    assert TABLE_NAME in ws.tables
    assert len(ws.conditional_formatting) >= 1
    assert ws.cell(row=DATA_START_ROW, column=15).value is not None
