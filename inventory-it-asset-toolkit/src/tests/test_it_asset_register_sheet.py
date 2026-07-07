"""Tests for IT Asset Register sheet builder."""

from openpyxl import Workbook

from src.data_generation.generate_it_assets import COLUMN_ORDER, generate_it_asset_data
from src.sheets.it_asset_register_sheet import (
    DATA_START_ROW,
    HEADER_ROW,
    HEADERS,
    TABLE_NAME,
    build,
)


def test_it_asset_register_sheet_structure():
    """IT Asset Register should include table, validations, and conditional formatting."""
    wb = Workbook()
    ws = wb.active
    df = generate_it_asset_data(row_count=50, seed=42)
    build(ws, {"data": {"it_assets": df}})

    assert ws.cell(row=HEADER_ROW, column=1).value == HEADERS[0]
    assert len(df.columns) == len(COLUMN_ORDER) == len(HEADERS)
    assert ws.cell(row=DATA_START_ROW, column=1).value is not None
    assert TABLE_NAME in ws.tables
    assert len(ws.data_validations.dataValidation) >= 2
    assert len(ws.conditional_formatting) >= 3
    assert ws.sheet_view.showGridLines is False


def test_it_asset_register_date_and_status_columns():
    """Date columns should be formatted; Missing status should exist in sample data."""
    wb = Workbook()
    ws = wb.active
    df = generate_it_asset_data(row_count=120, seed=42)
    build(ws, {"data": {"it_assets": df}})

    assert ws[f"J{DATA_START_ROW}"].number_format
    assert ws[f"K{DATA_START_ROW}"].number_format
    assert "Missing" in df["status"].values
