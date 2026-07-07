"""
Acceptance tests for core data generation and workbook build requirements.

Covers Prompt 19 / specs.md §16 essentials with practical, focused checks.
"""

from pathlib import Path

from openpyxl import load_workbook

from config.workbook_config import OUTPUT_FILENAME, SHEET_ORDER, WORKBOOK_PATH
from src.data_generation.generate_inventory import (
    COLUMN_ORDER as INVENTORY_COLUMNS,
    generate_inventory_data,
)
from src.data_generation.generate_it_assets import (
    COLUMN_ORDER as IT_ASSET_COLUMNS,
    generate_it_asset_data,
)
from src.data_generation.generate_software import generate_software_data
from src.main import generate_all_data
from src.workbook.builder import build_workbook


def test_inventory_data_is_not_empty():
    """Generated inventory data should contain rows."""
    df = generate_inventory_data(row_count=120, seed=42)
    assert not df.empty
    assert len(df) > 0


def test_it_asset_data_is_not_empty():
    """Generated IT asset data should contain rows."""
    df = generate_it_asset_data(row_count=120, seed=42)
    assert not df.empty
    assert len(df) > 0


def test_inventory_data_includes_required_columns():
    """Inventory data should include every required column in sheet order."""
    df = generate_inventory_data(row_count=50, seed=42)
    assert list(df.columns) == INVENTORY_COLUMNS


def test_it_asset_data_includes_required_columns():
    """IT asset data should include every required column in sheet order."""
    df = generate_it_asset_data(row_count=50, seed=42)
    assert list(df.columns) == IT_ASSET_COLUMNS


def test_software_data_includes_over_assigned_record():
    """Software data should include at least one over-assigned license."""
    df = generate_software_data(row_count=22, seed=42)
    assert "Over-Assigned" in df["compliance_status"].values


def test_workbook_file_is_created(tmp_path: Path):
    """Building the workbook should create a non-empty .xlsx file."""
    data = generate_all_data()
    output = tmp_path / "Inventory_IT_Asset_Control_Toolkit.xlsx"
    result = build_workbook(data, output_path=output)

    assert result.exists()
    assert result.suffix == ".xlsx"
    assert result.stat().st_size > 0


def test_workbook_contains_all_required_sheet_names(tmp_path: Path):
    """Workbook should contain every required sheet name."""
    data = generate_all_data()
    output = tmp_path / "workbook.xlsx"
    build_workbook(data, output_path=output)

    wb = load_workbook(output, read_only=True)
    assert set(SHEET_ORDER).issubset(set(wb.sheetnames))


def test_workbook_sheets_are_in_correct_order(tmp_path: Path):
    """Workbook tabs should appear in the required order."""
    data = generate_all_data()
    output = tmp_path / "workbook.xlsx"
    build_workbook(data, output_path=output)

    wb = load_workbook(output, read_only=True)
    assert wb.sheetnames == SHEET_ORDER


def test_workbook_output_path_matches_config():
    """Default build output should use the configured dist/ workbook path."""
    assert WORKBOOK_PATH.name == OUTPUT_FILENAME
    assert WORKBOOK_PATH.parent.name == "dist"


def test_workbook_sumif_formulas_are_excel_compatible(tmp_path: Path):
    """Dashboard and Management Summary SUMIF formulas must use quoted criteria."""
    output = tmp_path / "workbook.xlsx"
    build_workbook(generate_all_data(), output_path=output)
    wb = load_workbook(output, data_only=False)

    for sheet_name in ("Inventory Dashboard", "Management Summary"):
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=1, max_row=20, min_col=1, max_col=9):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("=SUMIF"):
                    assert ",>180," not in cell.value, f"Invalid SUMIF in {sheet_name} {cell.coordinate}"

    wb.close()


def test_workbook_can_be_opened_with_openpyxl(tmp_path: Path):
    """Built workbook should open without corruption using openpyxl."""
    data = generate_all_data()
    output = tmp_path / "workbook.xlsx"
    build_workbook(data, output_path=output)

    wb = load_workbook(output)
    try:
        assert len(wb.sheetnames) == len(SHEET_ORDER)
        assert wb["README"].cell(row=1, column=1).value is not None
    finally:
        wb.close()
