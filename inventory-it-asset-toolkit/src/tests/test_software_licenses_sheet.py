"""Tests for Software Licenses sheet builder."""

from openpyxl import Workbook

from src.data_generation.generate_software import generate_software_data
from src.sheets.software_licenses_sheet import (
    DATA_START_ROW,
    HEADER_ROW,
    HEADERS,
    SUMMARY_TITLE_ROW,
    TABLE_NAME,
    _compute_summary,
    build,
)


def test_compute_summary_metrics():
    """Summary should calculate portfolio counts and total annual cost."""
    df = generate_software_data(row_count=22, seed=42)
    summary = _compute_summary(df)

    assert summary["total_records"] == 22
    assert summary["over_assigned"] >= 1
    assert summary["renewals_30"] >= 1
    assert summary["renewals_90"] >= summary["renewals_30"]
    assert summary["total_annual_cost"] > 0


def test_software_licenses_sheet_structure():
    """Sheet should include summary, table, formatting, and validation."""
    wb = Workbook()
    ws = wb.active
    df = generate_software_data(row_count=22, seed=42)
    build(ws, {"data": {"software": df}})

    assert ws.cell(row=HEADER_ROW, column=1).value == HEADERS[0]
    assert ws.cell(row=SUMMARY_TITLE_ROW, column=1).value
    assert TABLE_NAME in ws.tables
    assert len(ws.conditional_formatting) >= 1
    assert len(ws.data_validations.dataValidation) >= 1
    assert ws[f"K{DATA_START_ROW}"].number_format
    assert ws[f"G{DATA_START_ROW}"].number_format
    assert ws.sheet_view.showGridLines is False
