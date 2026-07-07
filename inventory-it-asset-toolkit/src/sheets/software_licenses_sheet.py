"""Software Licenses sheet — license compliance and renewal tracking."""

from __future__ import annotations

from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from config import style_config as sc
from config.workbook_config import SOFTWARE_RENEWAL_THRESHOLDS
from src.data_generation.generate_software import COLUMN_ORDER
from src.workbook.styles import (
    apply_kpi_card_style,
    apply_risk_conditional_formatting,
    apply_section_header_style,
    apply_table_header_style,
    freeze_panes,
)
from src.workbook.utils import (
    autosize_columns,
    create_excel_table,
    format_currency_columns,
    format_date_columns,
    format_integer_columns,
    set_landscape_print,
)
from src.workbook.validations import add_compliance_status_validation

HEADERS = [
    "Software Name",
    "Vendor",
    "License Type",
    "Purchased Licenses",
    "Assigned Licenses",
    "Available Licenses",
    "Renewal Date",
    "Days Until Renewal",
    "Department",
    "Owner",
    "Annual Cost",
    "Compliance Status",
]

TITLE_ROW = 1
SUMMARY_TITLE_ROW = 3
SUMMARY_VALUE_ROW = 4
HEADER_ROW = 6
DATA_START_ROW = 7
COL_COUNT = len(HEADERS)
TABLE_NAME = "SoftwareLicensesTable"

COL_INTEGER = ["D", "E", "F", "H"]
COL_CURRENCY = ["K"]
COL_DATES = ["G"]
COL_COMPLIANCE = "L"

# Compliance conditional formatting (Over-Assigned red, Due Soon orange, Watch yellow)
COMPLIANCE_CF_MAP = {
    "Compliant": "healthy",
    "Renewal Watch": "watch",
    "Renewal Due Soon": "slow_moving",
    "Over-Assigned": "over_assigned",
}


def _coerce_cell_value(value: Any) -> Any:
    """Normalize pandas values for openpyxl."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if pd.isna(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.date()
    return value


def _compute_summary(df: pd.DataFrame) -> dict[str, Any]:
    """Compute summary metrics for the license portfolio."""
    if df.empty:
        return {
            "total_records": 0,
            "over_assigned": 0,
            "renewals_30": 0,
            "renewals_90": 0,
            "total_annual_cost": 0.0,
        }

    due_soon_days = SOFTWARE_RENEWAL_THRESHOLDS["renewal_due_soon_days"]
    watch_days = SOFTWARE_RENEWAL_THRESHOLDS["renewal_watch_days"]

    return {
        "total_records": len(df),
        "over_assigned": int((df["compliance_status"] == "Over-Assigned").sum()),
        "renewals_30": int((df["days_until_renewal"] <= due_soon_days).sum()),
        "renewals_90": int((df["days_until_renewal"] <= watch_days).sum()),
        "total_annual_cost": round(float(df["annual_cost"].sum()), 2),
    }


def _write_title(ws: Worksheet, record_count: int) -> None:
    """Render sheet title."""
    ws.merge_cells(start_row=TITLE_ROW, start_column=1, end_row=TITLE_ROW, end_column=COL_COUNT)
    cell = ws.cell(
        row=TITLE_ROW,
        column=1,
        value=f"Software Licenses — {record_count:,} Applications Tracked",
    )
    cell.font = Font(
        name=sc.FONTS["default_name"],
        bold=True,
        color=sc.COLORS["white"],
        size=sc.FONTS["title_size"],
    )
    cell.fill = sc.HEADER_FILL
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[TITLE_ROW].height = 28


def _write_summary_section(ws: Worksheet, summary: dict[str, Any]) -> None:
    """Write KPI summary cards above the license table."""
    apply_section_header_style(ws, SUMMARY_TITLE_ROW - 1, 1, "License Portfolio Summary", span_cols=10)

    kpis = [
        ("Total Software Records", summary["total_records"], sc.NUMBER_FORMATS["integer"]),
        ("Over-Assigned Count", summary["over_assigned"], sc.NUMBER_FORMATS["integer"]),
        ("Renewals Due ≤ 30 Days", summary["renewals_30"], sc.NUMBER_FORMATS["integer"]),
        ("Renewals Due ≤ 90 Days", summary["renewals_90"], sc.NUMBER_FORMATS["integer"]),
        ("Total Annual Cost", summary["total_annual_cost"], sc.NUMBER_FORMATS["currency_compact"]),
    ]

    card_cols = [1, 3, 5, 7, 9]
    for idx, (title, value, fmt) in enumerate(kpis):
        col = card_cols[idx]
        apply_kpi_card_style(
            ws, SUMMARY_TITLE_ROW, col, SUMMARY_VALUE_ROW, col, title, value, span_cols=2
        )
        ws.cell(row=SUMMARY_VALUE_ROW, column=col).number_format = fmt


def _write_headers(ws: Worksheet) -> None:
    """Write and style table headers."""
    for col_idx, header in enumerate(HEADERS, start=1):
        ws.cell(row=HEADER_ROW, column=col_idx, value=header)
    apply_table_header_style(ws, HEADER_ROW, COL_COUNT)
    ws.row_dimensions[HEADER_ROW].height = 22


def _write_data_rows(ws: Worksheet, df: pd.DataFrame) -> int:
    """Write license rows and return last row."""
    if df.empty:
        return HEADER_ROW

    for row_offset, (_, row) in enumerate(df.iterrows()):
        excel_row = DATA_START_ROW + row_offset
        for col_idx, col_name in enumerate(COLUMN_ORDER, start=1):
            ws.cell(
                row=excel_row,
                column=col_idx,
                value=_coerce_cell_value(row.get(col_name)),
            )
    return DATA_START_ROW + len(df) - 1


def _apply_column_formats(ws: Worksheet, last_row: int) -> None:
    """Apply integer, currency, and date formats."""
    if last_row < DATA_START_ROW:
        return
    format_integer_columns(ws, COL_INTEGER, DATA_START_ROW, last_row)
    format_currency_columns(ws, COL_CURRENCY, DATA_START_ROW, last_row)
    format_date_columns(ws, COL_DATES, DATA_START_ROW, last_row)


def _create_license_table(ws: Worksheet, last_row: int) -> None:
    """Create filterable Excel table."""
    end_col = get_column_letter(COL_COUNT)
    if last_row < DATA_START_ROW:
        last_row = DATA_START_ROW
    create_excel_table(ws, TABLE_NAME, f"A{HEADER_ROW}:{end_col}{last_row}")


def _apply_compliance_formatting(ws: Worksheet, last_row: int) -> None:
    """Apply compliance status conditional formatting."""
    if last_row < DATA_START_ROW:
        return

    status_range = f"{COL_COMPLIANCE}{DATA_START_ROW}:{COL_COMPLIANCE}{last_row}"
    apply_risk_conditional_formatting(
        ws,
        status_range,
        COL_COMPLIANCE,
        DATA_START_ROW,
        COMPLIANCE_CF_MAP,
    )


def _apply_validations(ws: Worksheet, last_row: int) -> None:
    """Add compliance status dropdown validation."""
    if last_row < DATA_START_ROW:
        return
    add_compliance_status_validation(
        ws, f"{COL_COMPLIANCE}{DATA_START_ROW}:{COL_COMPLIANCE}{last_row}"
    )


def build(ws: Worksheet, context: dict[str, Any]) -> None:
    """Build the Software Licenses sheet from software license data."""
    df: pd.DataFrame = context["data"].get("software", pd.DataFrame())
    summary = _compute_summary(df)

    _write_title(ws, len(df))
    _write_summary_section(ws, summary)
    _write_headers(ws)
    last_row = _write_data_rows(ws, df)

    _apply_column_formats(ws, last_row)
    _create_license_table(ws, last_row)
    _apply_compliance_formatting(ws, last_row)
    _apply_validations(ws, last_row)

    freeze_panes(ws, row=DATA_START_ROW, col=1)
    autosize_columns(ws, min_width=10, max_width=36)
    ws.column_dimensions["A"].width = min(ws.column_dimensions["A"].width, 28)

    ws.sheet_view.showGridLines = False
    set_landscape_print(ws, fit_width=1, repeat_header_rows=f"{HEADER_ROW}:{HEADER_ROW}")
