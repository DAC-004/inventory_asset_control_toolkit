"""Mobile Provisioning sheet — device issuance and recovery checklist."""

from __future__ import annotations

from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from config import style_config as sc
from src.data_generation.generate_mobile import COLUMN_ORDER
from src.workbook.styles import (
    apply_kpi_card_style,
    apply_risk_conditional_formatting,
    apply_section_header_style,
    apply_table_header_style,
    freeze_panes,
)
from src.workbook.utils import (
    autosize_columns,
    create_excel_table,
    format_date_columns,
    set_landscape_print,
)
from src.workbook.validations import add_mobile_status_validation, add_yes_no_validation

HEADERS = [
    "Employee Name",
    "Department",
    "Device Type",
    "Asset Tag",
    "IMEI",
    "SIM Number",
    "Phone Number",
    "Carrier",
    "MDM Enrolled",
    "Security Configured",
    "Required Apps Installed",
    "User Agreement Signed",
    "Date Issued",
    "Return Date",
    "Status",
]

TITLE_ROW = 1
SUMMARY_TITLE_ROW = 3
SUMMARY_VALUE_ROW = 4
HEADER_ROW = 6
DATA_START_ROW = 7
COL_COUNT = len(HEADERS)
TABLE_NAME = "MobileProvisioningTable"

COL_YES_NO = ["I", "J", "K", "L"]
COL_DATES = ["M", "N"]
COL_STATUS = "O"

MOBILE_STATUS_CF_MAP = {
    "Ready": "healthy",
    "Assigned": "healthy",
    "Pending Setup": "pending",
    "Missing Agreement": "missing",
    "Returned": "neutral",
    "Disabled": "watch",
}


def _coerce_cell_value(value: Any) -> Any:
    """Normalize pandas values for openpyxl."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if pd.isna(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.date()
    return value


def _compute_summary(df: pd.DataFrame) -> dict[str, int]:
    """Compute checklist summary metrics for the mobile fleet."""
    if df.empty:
        return {
            "total_devices": 0,
            "assigned": 0,
            "pending_setup": 0,
            "missing_agreement": 0,
            "returned": 0,
        }

    status = df["status"]
    return {
        "total_devices": len(df),
        "assigned": int((status == "Assigned").sum()),
        "pending_setup": int((status == "Pending Setup").sum()),
        "missing_agreement": int((status == "Missing Agreement").sum()),
        "returned": int((status == "Returned").sum()),
    }


def _write_title(ws: Worksheet, record_count: int) -> None:
    """Render sheet title."""
    ws.merge_cells(start_row=TITLE_ROW, start_column=1, end_row=TITLE_ROW, end_column=COL_COUNT)
    cell = ws.cell(
        row=TITLE_ROW,
        column=1,
        value=f"Mobile Provisioning — {record_count:,} Devices Tracked",
    )
    cell.font = Font(
        name=sc.FONTS["default_name"],
        bold=True,
        color=sc.COLORS["white"],
        size=sc.FONTS["title_size"],
    )
    cell.fill = sc.HEADER_FILL
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[TITLE_ROW].height = 28


def _write_summary_section(ws: Worksheet, summary: dict[str, int]) -> None:
    """Write checklist-style KPI summary above the provisioning table."""
    apply_section_header_style(
        ws, SUMMARY_TITLE_ROW - 1, 1, "Provisioning Checklist Summary", span_cols=10
    )

    kpis = [
        ("Total Mobile Devices", summary["total_devices"]),
        ("Assigned Devices", summary["assigned"]),
        ("Pending Setup", summary["pending_setup"]),
        ("Missing User Agreement", summary["missing_agreement"]),
        ("Returned Devices", summary["returned"]),
    ]

    card_cols = [1, 4, 7, 10, 13]
    for idx, (title, value) in enumerate(kpis):
        col = card_cols[idx]
        apply_kpi_card_style(
            ws, SUMMARY_TITLE_ROW, col, SUMMARY_VALUE_ROW, col, title, value, span_cols=2
        )
        ws.cell(row=SUMMARY_VALUE_ROW, column=col).number_format = sc.NUMBER_FORMATS["integer"]


def _write_headers(ws: Worksheet) -> None:
    """Write and style table headers."""
    for col_idx, header in enumerate(HEADERS, start=1):
        ws.cell(row=HEADER_ROW, column=col_idx, value=header)
    apply_table_header_style(ws, HEADER_ROW, COL_COUNT)
    ws.row_dimensions[HEADER_ROW].height = 22


def _write_data_rows(ws: Worksheet, df: pd.DataFrame) -> int:
    """Write mobile provisioning rows and return last row."""
    if df.empty:
        return HEADER_ROW

    for row_offset, (_, row) in enumerate(df.iterrows()):
        excel_row = DATA_START_ROW + row_offset
        for col_idx, col_name in enumerate(COLUMN_ORDER, start=1):
            ws.cell(
                row=excel_row,
                column=col_idx,
                value=_coerce_cell_value(row.get(col_name)),
            )
    return DATA_START_ROW + len(df) - 1


def _apply_column_formats(ws: Worksheet, last_row: int) -> None:
    """Apply date formatting to issue and return date columns."""
    if last_row < DATA_START_ROW:
        return
    format_date_columns(ws, COL_DATES, DATA_START_ROW, last_row)


def _create_provisioning_table(ws: Worksheet, last_row: int) -> None:
    """Create filterable Excel table."""
    end_col = get_column_letter(COL_COUNT)
    if last_row < DATA_START_ROW:
        last_row = DATA_START_ROW
    create_excel_table(ws, TABLE_NAME, f"A{HEADER_ROW}:{end_col}{last_row}")


def _apply_status_formatting(ws: Worksheet, last_row: int) -> None:
    """Apply status conditional formatting (Pending Setup orange, Missing Agreement red)."""
    if last_row < DATA_START_ROW:
        return

    status_range = f"{COL_STATUS}{DATA_START_ROW}:{COL_STATUS}{last_row}"
    apply_risk_conditional_formatting(
        ws,
        status_range,
        COL_STATUS,
        DATA_START_ROW,
        MOBILE_STATUS_CF_MAP,
    )


def _apply_validations(ws: Worksheet, last_row: int) -> None:
    """Add Yes/No and status dropdown validations."""
    if last_row < DATA_START_ROW:
        return

    for col in COL_YES_NO:
        add_yes_no_validation(ws, f"{col}{DATA_START_ROW}:{col}{last_row}")

    add_mobile_status_validation(
        ws, f"{COL_STATUS}{DATA_START_ROW}:{COL_STATUS}{last_row}"
    )


def build(ws: Worksheet, context: dict[str, Any]) -> None:
    """Build the Mobile Provisioning sheet from mobile provisioning data."""
    df: pd.DataFrame = context["data"].get("mobile", pd.DataFrame())
    summary = _compute_summary(df)

    _write_title(ws, len(df))
    _write_summary_section(ws, summary)
    _write_headers(ws)
    last_row = _write_data_rows(ws, df)

    _apply_column_formats(ws, last_row)
    _create_provisioning_table(ws, last_row)
    _apply_status_formatting(ws, last_row)
    _apply_validations(ws, last_row)

    freeze_panes(ws, row=DATA_START_ROW, col=1)
    autosize_columns(ws, min_width=10, max_width=36)
    ws.column_dimensions["A"].width = min(ws.column_dimensions["A"].width, 28)
    ws.column_dimensions["E"].width = min(ws.column_dimensions["E"].width, 22)
    ws.column_dimensions["F"].width = min(ws.column_dimensions["F"].width, 24)

    ws.sheet_view.showGridLines = False
    set_landscape_print(ws, fit_width=1, repeat_header_rows=f"{HEADER_ROW}:{HEADER_ROW}")
