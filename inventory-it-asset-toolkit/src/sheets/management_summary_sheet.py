"""Management Summary sheet — printable one-page executive overview."""

from __future__ import annotations

from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

from config import style_config as sc
from config.workbook_config import (
    DEMO_REFERENCE_DATE,
    INVENTORY_THRESHOLDS,
    RANDOM_SEED,
    ROW_COUNTS,
    SOFTWARE_RENEWAL_THRESHOLDS,
    VERSION,
)
from src.sheets.aged_excess_sheet import _build_analysis_dataframe
from src.sheets.audit_reconciliation_sheet import _build_exceptions_dataframe
from src.workbook.formulas import count_if_range, sum_if_numeric, sum_if_range, sum_range
from src.workbook.styles import (
    apply_kpi_card_style,
    apply_section_header_style,
    apply_table_header_range,
)
from src.workbook.utils import (
    format_currency_columns,
    set_column_widths,
    set_print_layout,
)

MASTER_SHEET = "Master Inventory"
IT_SHEET = "IT Asset Register"
SOFTWARE_SHEET = "Software Licenses"
DISPOSAL_SHEET = "Disposal Log"

MASTER_DATA_START = 3
IT_DATA_START = 3
SOFTWARE_DATA_START = 7
DISPOSAL_DATA_START = 7

COL_TOTAL_VALUE = "N"
COL_AGE_DAYS = "Q"
COL_STATUS = "U"
COL_RECOMMENDED_ACTION = "V"
COL_IT_STATUS = "L"
COL_DAYS_UNTIL_RENEWAL = "H"
COL_ANNUAL_COST = "K"
COL_COMPLIANCE = "L"
COL_DISPOSAL_STATUS = "N"

MARKDOWN_ACTIONS = ("Markdown Review", "Transfer or Markdown", "Liquidate")
TRANSFER_ACTIONS = ("Review Transfer", "Transfer or Markdown")
PENDING_DISPOSAL_STATUSES = (
    "Pending Wipe",
    "Wiped",
    "Pending Vendor Pickup",
    "Certificate Missing",
    "Hold for Review",
)

INVENTORY_RISK_HEADERS = ["SKU", "Location", "Issue", "Value"]
IT_RISK_HEADERS = ["Asset Tag", "Device", "Risk", "Location"]

IT_STATUS_PRIORITY = {
    "Missing": 0,
    "In Repair": 1,
    "Retired": 2,
    "Returned": 3,
    "In Stock": 4,
    "Assigned": 5,
    "Disposed": 6,
}

IT_RISK_LABELS = {
    "Missing": "Missing — Audit Exception",
    "In Repair": "In Repair — Service Backlog",
    "Retired": "Retired — Lifecycle Review",
    "Returned": "Returned — Pending Processing",
    "In Stock": "In Stock — Unassigned Asset",
    "Assigned": "Assigned — Monitor Compliance",
}

TITLE_ROW = 1
SUBTITLE_ROW = 2

SECTION1_ROW = 3
KPI1_TITLE_ROW = 4
KPI1_VALUE_ROW = 5

SECTION2_ROW = 7
KPI2_TITLE_ROW = 8
KPI2_VALUE_ROW = 9

SECTION3_ROW = 11
KPI3_TITLE_ROW = 12
KPI3_VALUE_ROW = 13

SECTION3B_ROW = 15
KPI3B_TITLE_ROW = 16
KPI3B_VALUE_ROW = 17

SECTION3C_ROW = 19
SUMMARY_HEADER_ROW = 20
SUMMARY_DATA_START_ROW = 21

SECTION4_ROW = 24
RISK_HEADER_ROW = 25
RISK_DATA_START_ROW = 26

SECTION5_ROW = 32
PLAN_HEADER_ROW = 33
PLAN_DATA_START_ROW = 34

ACTION_PLAN = {
    "30 Days": [
        "Validate asset and inventory records",
        "Review high-risk exceptions",
        "Confirm current reporting needs",
    ],
    "60 Days": [
        "Standardize audit and reconciliation procedures",
        "Improve dashboard reporting",
        "Identify recurring inventory issues",
    ],
    "90 Days": [
        "Automate recurring reports",
        "Build KPI review cadence",
        "Recommend process improvements",
    ],
}


def _data_end(start_row: int, record_count: int) -> int:
    """Return the last data row for a sheet section."""
    if record_count <= 0:
        return start_row
    return start_row + record_count - 1


def _inventory_kpis(inv_end: int) -> list[tuple[str, str, str]]:
    """Build inventory highlight KPI definitions."""
    start = MASTER_DATA_START
    aged_threshold = INVENTORY_THRESHOLDS["excess_aged_age_days"]
    excess_formula = "=" + "+".join(
        sum_if_range(MASTER_SHEET, COL_STATUS, status, COL_TOTAL_VALUE, start, inv_end).lstrip("=")
        for status in ("Excess", "Excess / Aged")
    )
    transfer_formula = "=" + "+".join(
        count_if_range(MASTER_SHEET, COL_RECOMMENDED_ACTION, start, inv_end, action).lstrip("=")
        for action in TRANSFER_ACTIONS
    )
    markdown_formula = "=" + "+".join(
        count_if_range(MASTER_SHEET, COL_RECOMMENDED_ACTION, start, inv_end, action).lstrip("=")
        for action in MARKDOWN_ACTIONS
    )
    return [
        ("Total Inventory Value", sum_range(MASTER_SHEET, COL_TOTAL_VALUE, start, inv_end), "currency"),
        (
            "Aged Inventory Value",
            sum_if_numeric(MASTER_SHEET, COL_AGE_DAYS, f">{aged_threshold}", COL_TOTAL_VALUE, start, inv_end),
            "currency",
        ),
        ("Excess Inventory Value", excess_formula, "currency"),
        ("Transfer Candidates", transfer_formula, "integer"),
        ("Markdown Candidates", markdown_formula, "integer"),
    ]


def _it_asset_kpis(it_end: int, disposal_end: int, audit_exceptions: int) -> list[tuple[str, str, str]]:
    """Build IT asset control KPI definitions."""
    start = IT_DATA_START
    disposal_start = DISPOSAL_DATA_START
    pending_disposal = "=" + "+".join(
        count_if_range(DISPOSAL_SHEET, COL_DISPOSAL_STATUS, disposal_start, disposal_end, status).lstrip("=")
        for status in PENDING_DISPOSAL_STATUSES
    )
    return [
        (
            "Total IT Assets",
            f"=COUNTA('{IT_SHEET}'!A{start}:A{it_end})",
            "integer",
        ),
        ("Assigned Assets", count_if_range(IT_SHEET, COL_IT_STATUS, start, it_end, "Assigned"), "integer"),
        ("Missing Assets", count_if_range(IT_SHEET, COL_IT_STATUS, start, it_end, "Missing"), "integer"),
        ("Audit Exceptions", str(audit_exceptions), "integer"),
        ("Assets Pending Disposal", pending_disposal, "integer"),
    ]


def _software_kpis(sw_end: int) -> list[tuple[str, str, str]]:
    """Build software license risk KPI definitions."""
    start = SOFTWARE_DATA_START
    watch_days = SOFTWARE_RENEWAL_THRESHOLDS["renewal_watch_days"]
    return [
        (
            "Total Software Records",
            f"=COUNTA('{SOFTWARE_SHEET}'!A{start}:A{sw_end})",
            "integer",
        ),
        (
            "Over-Assigned Licenses",
            count_if_range(SOFTWARE_SHEET, COL_COMPLIANCE, start, sw_end, "Over-Assigned"),
            "integer",
        ),
        (
            "Renewals Due ≤ 90 Days",
            f"=COUNTIF('{SOFTWARE_SHEET}'!{COL_DAYS_UNTIL_RENEWAL}{start}:{COL_DAYS_UNTIL_RENEWAL}{sw_end},\"<={watch_days}\")",
            "integer",
        ),
        (
            "Total Annual Software Cost",
            sum_range(SOFTWARE_SHEET, COL_ANNUAL_COST, start, sw_end),
            "currency",
        ),
    ]


def _disposal_kpis(disp_end: int) -> list[tuple[str, str, str]]:
    """Build disposal and lifecycle risk KPI definitions."""
    start = DISPOSAL_DATA_START
    return [
        (
            "Pending Wipe",
            count_if_range(DISPOSAL_SHEET, COL_DISPOSAL_STATUS, start, disp_end, "Pending Wipe"),
            "integer",
        ),
        (
            "Certificate Missing",
            count_if_range(
                DISPOSAL_SHEET, COL_DISPOSAL_STATUS, start, disp_end, "Certificate Missing"
            ),
            "integer",
        ),
        (
            "Disposed",
            count_if_range(DISPOSAL_SHEET, COL_DISPOSAL_STATUS, start, disp_end, "Disposed"),
            "integer",
        ),
        (
            "Hold for Review",
            count_if_range(DISPOSAL_SHEET, COL_DISPOSAL_STATUS, start, disp_end, "Hold for Review"),
            "integer",
        ),
    ]


def _recommended_action_summary(inventory: pd.DataFrame) -> list[tuple[str, int]]:
    """Summarize recommended action counts for the executive summary."""
    if inventory.empty:
        return []
    counts = inventory["recommended_action"].value_counts()
    return [(str(action), int(count)) for action, count in counts.items()]


def _audit_exception_summary(exceptions: pd.DataFrame) -> list[tuple[str, int]]:
    """Summarize audit exception counts excluding clean records."""
    if exceptions.empty:
        return []
    filtered = exceptions[exceptions["Exception Type"] != "No Exception"]
    if filtered.empty:
        return [("No Exceptions", 0)]
    counts = filtered["Exception Type"].value_counts().head(5)
    return [(str(exc_type), int(count)) for exc_type, count in counts.items()]


def _compliance_summary(software: pd.DataFrame) -> list[tuple[str, int]]:
    """Summarize software compliance status counts."""
    if software.empty:
        return []
    counts = software["compliance_status"].value_counts()
    return [(str(status), int(count)) for status, count in counts.items()]


def _top_inventory_risks(inventory: pd.DataFrame, limit: int = 5) -> list[dict[str, Any]]:
    """Return top inventory risk rows for the executive summary."""
    analysis = _build_analysis_dataframe(inventory)
    if analysis.empty:
        return []

    rows: list[dict[str, Any]] = []
    for _, row in analysis.head(limit).iterrows():
        rows.append(
            {
                "SKU": row["SKU"],
                "Location": row["Location"],
                "Issue": row["Issue Type"],
                "Value": row["Inventory Value"],
            }
        )
    return rows


def _top_it_asset_risks(assets: pd.DataFrame, limit: int = 5) -> list[dict[str, Any]]:
    """Return top IT asset risk rows for the executive summary."""
    if assets.empty:
        return []

    working = assets[assets["status"] != "Disposed"].copy()
    working["_priority"] = working["status"].map(lambda status: IT_STATUS_PRIORITY.get(status, 99))
    working = working.sort_values(["_priority", "asset_tag"]).head(limit)

    rows: list[dict[str, Any]] = []
    for _, row in working.iterrows():
        status = str(row["status"])
        rows.append(
            {
                "Asset Tag": row["asset_tag"],
                "Device": row["device_type"],
                "Risk": IT_RISK_LABELS.get(status, status),
                "Location": row["location"] or row["department"] or "Unassigned",
            }
        )
    return rows


def _write_title_banner(ws: Worksheet) -> None:
    """Render executive summary title."""
    ws.merge_cells(start_row=TITLE_ROW, start_column=1, end_row=TITLE_ROW, end_column=10)
    cell = ws.cell(row=TITLE_ROW, column=1, value="Management Summary")
    cell.font = Font(
        name=sc.FONTS["default_name"],
        bold=True,
        color=sc.COLORS["white"],
        size=sc.FONTS["title_size"],
    )
    cell.fill = sc.HEADER_FILL
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[TITLE_ROW].height = 30

    ws.merge_cells(start_row=SUBTITLE_ROW, start_column=1, end_row=SUBTITLE_ROW, end_column=10)
    subtitle = ws.cell(
        row=SUBTITLE_ROW,
        column=1,
        value=(
            f"{VERSION} | Executive Overview | Reference Date: "
            f"{DEMO_REFERENCE_DATE.strftime('%B %d, %Y')}"
        ),
    )
    subtitle.font = Font(
        name=sc.FONTS["default_name"],
        italic=True,
        color=sc.COLORS["dark_blue"],
        size=sc.FONTS["body_size"],
    )
    subtitle.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[SUBTITLE_ROW].height = 18


def _write_kpi_cards(
    ws: Worksheet,
    title_row: int,
    value_row: int,
    kpis: list[tuple[str, str, str]],
    card_cols: list[int] | None = None,
) -> None:
    """Render a row of KPI cards with number formats."""
    card_cols = card_cols or [1, 3, 5, 7, 9]
    format_map = {
        "currency": sc.NUMBER_FORMATS["currency_compact"],
        "integer": sc.NUMBER_FORMATS["integer"],
        "percentage": sc.NUMBER_FORMATS["percentage"],
    }

    for idx, (title, value, fmt_key) in enumerate(kpis):
        col = card_cols[idx]
        apply_kpi_card_style(ws, title_row, col, value_row, col, title, value, span_cols=2)
        ws.cell(row=value_row, column=col).number_format = format_map[fmt_key]


def _write_executive_summaries(
    ws: Worksheet,
    action_rows: list[tuple[str, int]],
    audit_rows: list[tuple[str, int]],
    compliance_rows: list[tuple[str, int]],
) -> None:
    """Write compact Recommended Action, Audit Exception, and Compliance summaries."""
    apply_section_header_style(
        ws, SECTION3C_ROW, 1, "Executive Summaries", span_cols=10
    )

    blocks = [
        (1, "Recommended Action Summary", ["Action", "Count"], action_rows),
        (4, "Audit Exception Summary", ["Exception Type", "Count"], audit_rows),
        (7, "Software Compliance Summary", ["Compliance Status", "Count"], compliance_rows),
    ]

    for start_col, title, headers, rows in blocks:
        end_col = start_col + 1
        ws.merge_cells(
            start_row=SUMMARY_HEADER_ROW,
            start_column=start_col,
            end_row=SUMMARY_HEADER_ROW,
            end_column=end_col,
        )
        ws.cell(row=SUMMARY_HEADER_ROW, column=start_col, value=title).font = sc.KPI_TITLE_FONT
        apply_table_header_range(ws, SUMMARY_HEADER_ROW + 1, start_col, end_col)
        for col_offset, header in enumerate(headers):
            ws.cell(row=SUMMARY_HEADER_ROW + 1, column=start_col + col_offset, value=header)

        for row_offset, (label, count) in enumerate(rows[:5]):
            excel_row = SUMMARY_DATA_START_ROW + row_offset
            ws.cell(row=excel_row, column=start_col, value=label)
            count_cell = ws.cell(row=excel_row, column=start_col + 1, value=count)
            count_cell.number_format = sc.NUMBER_FORMATS["integer"]


def _write_risk_tables(
    ws: Worksheet,
    inventory_risks: list[dict[str, Any]],
    it_risks: list[dict[str, Any]],
) -> None:
    """Write side-by-side top risk mini tables."""
    apply_section_header_style(ws, SECTION4_ROW, 1, "5. Top Risks", span_cols=10)

    ws.merge_cells(start_row=RISK_HEADER_ROW, start_column=1, end_row=RISK_HEADER_ROW, end_column=4)
    ws.cell(row=RISK_HEADER_ROW, column=1, value="Top 5 Inventory Risks").font = sc.KPI_TITLE_FONT
    apply_table_header_range(ws, RISK_HEADER_ROW + 1, 1, 4)
    for col_idx, header in enumerate(INVENTORY_RISK_HEADERS, start=1):
        ws.cell(row=RISK_HEADER_ROW + 1, column=col_idx, value=header)

    ws.merge_cells(start_row=RISK_HEADER_ROW, start_column=6, end_row=RISK_HEADER_ROW, end_column=9)
    ws.cell(row=RISK_HEADER_ROW, column=6, value="Top 5 IT Asset Risks").font = sc.KPI_TITLE_FONT
    apply_table_header_range(ws, RISK_HEADER_ROW + 1, 6, 9)
    for col_idx, header in enumerate(IT_RISK_HEADERS, start=6):
        ws.cell(row=RISK_HEADER_ROW + 1, column=col_idx, value=header)

    for offset in range(5):
        excel_row = RISK_DATA_START_ROW + offset
        if offset < len(inventory_risks):
            risk = inventory_risks[offset]
            ws.cell(row=excel_row, column=1, value=risk["SKU"])
            ws.cell(row=excel_row, column=2, value=risk["Location"])
            ws.cell(row=excel_row, column=3, value=risk["Issue"])
            value_cell = ws.cell(row=excel_row, column=4, value=risk["Value"])
            value_cell.number_format = sc.NUMBER_FORMATS["currency_compact"]
        if offset < len(it_risks):
            risk = it_risks[offset]
            ws.cell(row=excel_row, column=6, value=risk["Asset Tag"])
            ws.cell(row=excel_row, column=7, value=risk["Device"])
            ws.cell(row=excel_row, column=8, value=risk["Risk"])
            ws.cell(row=excel_row, column=9, value=risk["Location"])


def _write_action_plan(ws: Worksheet) -> None:
    """Write the 30 / 60 / 90 day action plan section."""
    apply_section_header_style(
        ws, SECTION5_ROW, 1, "6. 30 / 60 / 90 Day Action Plan", span_cols=10
    )

    plan_cols = [1, 4, 7]
    horizons = ["30 Days", "60 Days", "90 Days"]
    for idx, horizon in enumerate(horizons):
        col = plan_cols[idx]
        ws.merge_cells(
            start_row=PLAN_HEADER_ROW,
            start_column=col,
            end_row=PLAN_HEADER_ROW,
            end_column=col + 2,
        )
        header = ws.cell(row=PLAN_HEADER_ROW, column=col, value=horizon)
        header.font = sc.KPI_TITLE_FONT
        header.fill = sc.KPI_CARD_FILL
        header.alignment = sc.KPI_CARD_STYLE["alignment"]
        header.border = sc.NAVY_BORDER

        for bullet_idx, bullet in enumerate(ACTION_PLAN[horizon]):
            row = PLAN_DATA_START_ROW + bullet_idx
            ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 2)
            cell = ws.cell(row=row, column=col, value=f"• {bullet}")
            cell.font = sc.BODY_FONT
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            cell.border = sc.THIN_GRAY_BORDER
            ws.row_dimensions[row].height = 28


def build(ws: Worksheet, context: dict[str, Any]) -> None:
    """Build the printable Management Summary one-pager."""
    data: dict[str, pd.DataFrame] = context["data"]

    inventory = data.get("inventory", pd.DataFrame())
    it_assets = data.get("it_assets", pd.DataFrame())
    software = data.get("software", pd.DataFrame())
    disposal = data.get("disposal", pd.DataFrame())

    inv_count = len(inventory) or ROW_COUNTS["inventory"]["default"]
    it_count = len(it_assets) or ROW_COUNTS["it_assets"]["default"]
    sw_count = len(software) or ROW_COUNTS["software"]["default"]
    disp_count = len(disposal) or ROW_COUNTS["disposal"]["default"]

    inv_end = _data_end(MASTER_DATA_START, inv_count)
    it_end = _data_end(IT_DATA_START, it_count)
    sw_end = _data_end(SOFTWARE_DATA_START, sw_count)
    disp_end = _data_end(DISPOSAL_DATA_START, disp_count)

    exceptions = _build_exceptions_dataframe(it_assets, seed=RANDOM_SEED)
    audit_exception_count = int((exceptions["Exception Type"] != "No Exception").sum())
    inventory_risks = _top_inventory_risks(inventory)
    it_risks = _top_it_asset_risks(it_assets)

    _write_title_banner(ws)

    apply_section_header_style(ws, SECTION1_ROW, 1, "1. Inventory Health Highlights", span_cols=10)
    _write_kpi_cards(ws, KPI1_TITLE_ROW, KPI1_VALUE_ROW, _inventory_kpis(inv_end))

    apply_section_header_style(ws, SECTION2_ROW, 1, "2. IT Asset Control Highlights", span_cols=10)
    _write_kpi_cards(
        ws,
        KPI2_TITLE_ROW,
        KPI2_VALUE_ROW,
        _it_asset_kpis(it_end, disp_end, audit_exception_count),
    )

    apply_section_header_style(ws, SECTION3_ROW, 1, "3. Software License Risks", span_cols=10)
    _write_kpi_cards(
        ws,
        KPI3_TITLE_ROW,
        KPI3_VALUE_ROW,
        _software_kpis(sw_end),
        card_cols=[1, 3, 5, 7],
    )

    apply_section_header_style(
        ws, SECTION3B_ROW, 1, "4. Disposal / Lifecycle Risks", span_cols=10
    )
    _write_kpi_cards(
        ws,
        KPI3B_TITLE_ROW,
        KPI3B_VALUE_ROW,
        _disposal_kpis(disp_end),
        card_cols=[1, 3, 5, 7],
    )

    _write_executive_summaries(
        ws,
        _recommended_action_summary(inventory),
        _audit_exception_summary(exceptions),
        _compliance_summary(software),
    )

    _write_risk_tables(ws, inventory_risks, it_risks)
    _write_action_plan(ws)

    format_currency_columns(ws, ["D"], RISK_DATA_START_ROW, RISK_DATA_START_ROW + 4)

    set_column_widths(
        ws,
        {
            "A": 16,
            "B": 18,
            "C": 22,
            "D": 14,
            "E": 2,
            "F": 14,
            "G": 16,
            "H": 24,
            "I": 18,
            "J": 8,
        },
    )

    ws.sheet_view.showGridLines = False
    set_print_layout(
        ws,
        orientation="landscape",
        fit_width=1,
        fit_height=1,
        repeat_header_rows=f"{TITLE_ROW}:{SUBTITLE_ROW}",
    )
