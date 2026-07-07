"""README sheet — toolkit guide for interviewers and reviewers."""

from __future__ import annotations

from typing import Any

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from config import style_config as sc
from config.workbook_config import AUTHOR, VERSION, WORKBOOK_TITLE
from src.workbook.styles import apply_section_header_style
from src.workbook.utils import set_column_widths, set_portrait_print

# Layout constants
CONTENT_START_COL = 1
CONTENT_END_COL = 4
SPAN_COLS = CONTENT_END_COL - CONTENT_START_COL + 1


def _merge_and_write(
    ws: Worksheet,
    row: int,
    text: str,
    *,
    font: Font | None = None,
    fill: PatternFill | None = None,
    alignment: Alignment | None = None,
    row_height: float | None = None,
) -> None:
    """Write wrapped text into a merged block across the content columns."""
    ws.merge_cells(
        start_row=row,
        start_column=CONTENT_START_COL,
        end_row=row,
        end_column=CONTENT_END_COL,
    )
    cell = ws.cell(row=row, column=CONTENT_START_COL, value=text)
    cell.font = font or sc.BODY_FONT
    cell.alignment = alignment or Alignment(
        horizontal="left",
        vertical="top",
        wrap_text=True,
    )
    if fill:
        cell.fill = fill
    if row_height:
        ws.row_dimensions[row].height = row_height


def _apply_title_banner(ws: Worksheet, row: int, title: str) -> None:
    """Render the main toolkit title in a navy banner."""
    ws.merge_cells(
        start_row=row,
        start_column=CONTENT_START_COL,
        end_row=row,
        end_column=CONTENT_END_COL,
    )
    cell = ws.cell(row=row, column=CONTENT_START_COL, value=title)
    cell.font = Font(
        name=sc.FONTS["default_name"],
        bold=True,
        color=sc.COLORS["white"],
        size=18,
    )
    cell.fill = PatternFill(
        start_color=sc.COLORS["navy"],
        end_color=sc.COLORS["navy"],
        fill_type="solid",
    )
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 36


def _apply_subtitle_bar(ws: Worksheet, row: int, text: str) -> None:
    """Render version and author line on a light-blue bar."""
    ws.merge_cells(
        start_row=row,
        start_column=CONTENT_START_COL,
        end_row=row,
        end_column=CONTENT_END_COL,
    )
    cell = ws.cell(row=row, column=CONTENT_START_COL, value=text)
    cell.font = Font(
        name=sc.FONTS["default_name"],
        bold=True,
        color=sc.COLORS["dark_blue"],
        size=sc.FONTS["body_size"],
    )
    cell.fill = sc.KPI_CARD_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 22


def _apply_disclaimer(ws: Worksheet, row: int, text: str) -> None:
    """Render the sample-data disclaimer with subtle emphasis."""
    ws.merge_cells(
        start_row=row,
        start_column=CONTENT_START_COL,
        end_row=row,
        end_column=CONTENT_END_COL,
    )
    cell = ws.cell(row=row, column=CONTENT_START_COL, value=text)
    cell.font = Font(
        name=sc.FONTS["default_name"],
        italic=True,
        color=sc.COLORS["dark_blue"],
        size=sc.FONTS["body_size"],
    )
    cell.fill = PatternFill(
        start_color=sc.COLORS["gray"],
        end_color=sc.COLORS["gray"],
        fill_type="solid",
    )
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 28


def _write_section(ws: Worksheet, row: int, title: str) -> int:
    """Write a section header and return the next available row."""
    apply_section_header_style(ws, row, CONTENT_START_COL, title, span_cols=SPAN_COLS)
    ws.row_dimensions[row].height = 22
    return row + 1


def _write_body(ws: Worksheet, row: int, text: str, row_height: float = 48) -> int:
    """Write wrapped body text and return the next available row."""
    _merge_and_write(ws, row, text, row_height=row_height)
    return row + 1


def _write_spacer(ws: Worksheet, row: int, height: float = 8) -> int:
    """Insert a blank spacer row and return the next row."""
    ws.row_dimensions[row].height = height
    return row + 1


def build(ws: Worksheet, context: dict[str, Any]) -> None:
    """Populate the README tab with toolkit overview, demo paths, and interview notes."""
    row = 1

    # --- Title area ---
    _apply_title_banner(ws, row, WORKBOOK_TITLE)
    row += 1
    _apply_subtitle_bar(ws, row, f"Version: {VERSION}    |    Author: {AUTHOR}")
    row += 1
    _apply_disclaimer(
        ws,
        row,
        "Disclaimer: All data in this workbook is fictional sample data created for "
        "demonstration and interview purposes. No real company, customer, or employee "
        "information is represented.",
    )
    row += 1
    row = _write_spacer(ws, row)

    # --- Purpose ---
    row = _write_section(ws, row, "Purpose of This Workbook")
    row = _write_body(
        ws,
        row,
        "This toolkit demonstrates a disciplined approach to inventory optimization and "
        "IT asset lifecycle control. It combines clean data capture, exception reporting, "
        "audit readiness, and management-level visibility in a single interview-ready "
        "workbook generated entirely from Python source code.",
        row_height=56,
    )
    row = _write_spacer(ws, row)

    # --- Mavis demo path ---
    row = _write_section(ws, row, "Mavis / Inventory Optimization Demo Path")
    row = _write_body(
        ws,
        row,
        "Recommended tab order:\n"
        "1. Inventory Dashboard\n"
        "2. Aged Excess Analysis\n"
        "3. Transfer Planner\n"
        "4. Markdown Planner\n"
        "5. Management Summary\n\n"
        "Suggested narrative:\n"
        '"This section focuses on inventory optimization. The dashboard shows inventory '
        "value, aged stock, excess inventory, transfer candidates, markdown candidates, "
        "and estimated recovery value. From there, I drill into aged and excess inventory, "
        "evaluate whether a transfer should happen before markdown, and use the markdown "
        'planner to protect margin while exiting slow-moving inventory."',
        row_height=120,
    )
    row = _write_spacer(ws, row)

    # --- IT demo path ---
    row = _write_section(ws, row, "IT Inventory Control Demo Path")
    row = _write_body(
        ws,
        row,
        "Recommended tab order:\n"
        "1. IT Asset Register\n"
        "2. Audit Reconciliation\n"
        "3. Software Licenses\n"
        "4. Mobile Provisioning\n"
        "5. Disposal Log\n"
        "6. Management Summary\n\n"
        "Suggested narrative:\n"
        '"This section focuses on IT asset lifecycle control. The asset register tracks '
        "equipment from receipt through assignment, audit, return, repair, data wipe, and "
        "disposal. Audit reconciliation compares system inventory against physical audit "
        "results and flags exceptions. The license tracker and mobile provisioning tabs "
        'help standardize compliance and equipment control."',
        row_height=120,
    )
    row = _write_spacer(ws, row)

    # --- Key capabilities ---
    row = _write_section(ws, row, "Key Capabilities")
    row = _write_body(
        ws,
        row,
        "Inventory & Supply Chain\n"
        "• Inventory health dashboards with KPI cards and charts\n"
        "• Aged, excess, slow-moving, and obsolete stock analysis\n"
        "• Transfer planning between locations to balance stock\n"
        "• Markdown modeling to protect margin on end-of-life inventory\n\n"
        "IT Asset Control\n"
        "• Full lifecycle asset register from receipt through disposal\n"
        "• Physical audit reconciliation with exception reporting\n"
        "• Software license compliance and renewal risk tracking\n"
        "• Mobile provisioning and recovery checklists\n"
        "• Disposal log with data wipe and certificate tracking\n\n"
        "Reporting\n"
        "• Printable management summary for leadership review\n"
        "• Conditional formatting to highlight operational risk\n"
        "• Filterable tables, formulas, and professional styling throughout",
        row_height=160,
    )
    row = _write_spacer(ws, row)

    # --- Interview talking points ---
    row = _write_section(ws, row, "Recommended Interview Talking Points")
    row = _write_body(
        ws,
        row,
        "Opening statement:\n"
        '"I built this toolkit to demonstrate how I approach inventory and asset control: '
        "clean data capture, lifecycle tracking, exception reporting, audit readiness, and "
        "management visibility. The same disciplined process applies whether I am analyzing "
        'warehouse inventory, retail stock, IT assets, software licenses, or equipment lifecycle status."\n\n'
        "Points to emphasize:\n"
        "• Business rules drive status, recommended actions, and compliance flags\n"
        "• Sample data is realistic but fictional — safe for live demo and portfolio review\n"
        "• Workbook is regenerated from code — reproducible, auditable, and easy to extend\n"
        "• Dashboards translate operational detail into decisions leaders can act on\n"
        "• Exception tabs (aged inventory, audit gaps, license over-assignment) show proactive control\n"
        "• Management Summary ties both use cases into a single leadership view",
        row_height=150,
    )
    row = _write_spacer(ws, row)

    # --- Navigation note ---
    row = _write_section(ws, row, "Workbook Navigation")
    row = _write_body(
        ws,
        row,
        "Use the sheet tabs at the bottom to move between modules. Start with the demo path "
        "that matches your interview focus, then close with Management Summary. All operational "
        "tabs include filterable headers — click the dropdown arrows to explore by status, "
        "location, or department.",
        row_height=48,
    )

    # --- Layout ---
    set_column_widths(ws, {"A": 4, "B": 28, "C": 28, "D": 28})
    ws.sheet_view.showGridLines = False
    set_portrait_print(ws, fit_width=1, fit_height=0, repeat_header_rows="1:3")
