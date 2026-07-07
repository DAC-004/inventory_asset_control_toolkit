"""Disposal Log sheet — asset retirement and data destruction tracking."""

from __future__ import annotations

from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from config import style_config as sc
from src.data_generation.generate_disposal import COLUMN_ORDER
from src.workbook.styles import (
    apply_kpi_card_style,
    apply_risk_conditional_formatting,
    apply_section_header_style,
    apply_table_header_style,
    freeze_panes,
)
from src.workbook.utils import (
    autosize_columns,
    create_excel_table,
    format_date_columns,
    set_landscape_print,
)
from src.workbook.validations import add_disposal_status_validation, add_yes_no_validation

HEADERS = [
    "Asset Tag",
    "Serial Number",
    "Device Type",
    "Assigned User",
    "Return Date",
    "Condition",
    "Data Wipe Required",
    "Data Wipe Completed",
    "Wipe Method",
    "Disposal Vendor",
    "Certificate Received",
    "Disposal Date",
    "Approved By",
    "Disposal Status",
]

TITLE_ROW = 1
SUMMARY_TITLE_ROW = 3
SUMMARY_VALUE_ROW = 4
HEADER_ROW = 6
DATA_START_ROW = 7
COL_COUNT = len(HEADERS)
TABLE_NAME = "DisposalLogTable"

COL_YES_NO = ["G", "H", "K"]
COL_DATES = ["E", "L"]
COL_STATUS = "N"

DISPOSAL_STATUS_CF_MAP = {
    "Pending Wipe": "pending",
    "Wiped": "watch",
    "Pending Vendor Pickup": "watch",
    "Disposed": "healthy",
    "Certificate Missing": "critical",
    "Hold for Review": "neutral",
}


def _coerce_cell_value(value: Any) -> Any:
    """Normalize pandas values for openpyxl."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if pd.isna(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.date()
    return value


def _compute_summary(df: pd.DataFrame) -> dict[str, int]:
    """Compute disposal pipeline summary metrics."""
    if df.empty:
        return {
            "total_records": 0,
            "pending_wipe": 0,
            "certificate_missing": 0,
            "disposed": 0,
            "hold_for_review": 0,
        }

    status = df["disposal_status"]
    return {
        "total_records": len(df),
        "pending_wipe": int((status == "Pending Wipe").sum()),
        "certificate_missing": int((status == "Certificate Missing").sum()),
        "disposed": int((status == "Disposed").sum()),
        "hold_for_review": int((status == "Hold for Review").sum()),
    }


def _write_title(ws: Worksheet, record_count: int) -> None:
    """Render sheet title."""
    ws.merge_cells(start_row=TITLE_ROW, start_column=1, end_row=TITLE_ROW, end_column=COL_COUNT)
    cell = ws.cell(
        row=TITLE_ROW,
        column=1,
        value=f"Disposal Log — {record_count:,} Retirement Records",
    )
    cell.font = Font(
        name=sc.FONTS["default_name"],
        bold=True,
        color=sc.COLORS["white"],
        size=sc.FONTS["title_size"],
    )
    cell.fill = sc.HEADER_FILL
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[TITLE_ROW].height = 28


def _write_summary_section(ws: Worksheet, summary: dict[str, int]) -> None:
    """Write KPI summary above the disposal table."""
    apply_section_header_style(
        ws, SUMMARY_TITLE_ROW - 1, 1, "Disposal Pipeline Summary", span_cols=10
    )

    kpis = [
        ("Total Disposal Records", summary["total_records"]),
        ("Pending Wipe", summary["pending_wipe"]),
        ("Certificate Missing", summary["certificate_missing"]),
        ("Disposed", summary["disposed"]),
        ("Hold for Review", summary["hold_for_review"]),
    ]

    card_cols = [1, 4, 7, 10, 13]
    for idx, (title, value) in enumerate(kpis):
        col = card_cols[idx]
        apply_kpi_card_style(
            ws, SUMMARY_TITLE_ROW, col, SUMMARY_VALUE_ROW, col, title, value, span_cols=2
        )
        ws.cell(row=SUMMARY_VALUE_ROW, column=col).number_format = sc.NUMBER_FORMATS["integer"]


def _write_headers(ws: Worksheet) -> None:
    """Write and style table headers."""
    for col_idx, header in enumerate(HEADERS, start=1):
        ws.cell(row=HEADER_ROW, column=col_idx, value=header)
    apply_table_header_style(ws, HEADER_ROW, COL_COUNT)
    ws.row_dimensions[HEADER_ROW].height = 22


def _write_data_rows(ws: Worksheet, df: pd.DataFrame) -> int:
    """Write disposal log rows and return last row."""
    if df.empty:
        return HEADER_ROW

    for row_offset, (_, row) in enumerate(df.iterrows()):
        excel_row = DATA_START_ROW + row_offset
        for col_idx, col_name in enumerate(COLUMN_ORDER, start=1):
            ws.cell(
                row=excel_row,
                column=col_idx,
                value=_coerce_cell_value(row.get(col_name)),
            )
    return DATA_START_ROW + len(df) - 1


def _apply_column_formats(ws: Worksheet, last_row: int) -> None:
    """Apply date formatting to return and disposal date columns."""
    if last_row < DATA_START_ROW:
        return
    format_date_columns(ws, COL_DATES, DATA_START_ROW, last_row)


def _create_disposal_table(ws: Worksheet, last_row: int) -> None:
    """Create filterable Excel table."""
    end_col = get_column_letter(COL_COUNT)
    if last_row < DATA_START_ROW:
        last_row = DATA_START_ROW
    create_excel_table(ws, TABLE_NAME, f"A{HEADER_ROW}:{end_col}{last_row}")


def _apply_status_formatting(ws: Worksheet, last_row: int) -> None:
    """Apply disposal status conditional formatting."""
    if last_row < DATA_START_ROW:
        return

    status_range = f"{COL_STATUS}{DATA_START_ROW}:{COL_STATUS}{last_row}"
    apply_risk_conditional_formatting(
        ws,
        status_range,
        COL_STATUS,
        DATA_START_ROW,
        DISPOSAL_STATUS_CF_MAP,
    )


def _apply_validations(ws: Worksheet, last_row: int) -> None:
    """Add Yes/No and disposal status dropdown validations."""
    if last_row < DATA_START_ROW:
        return

    for col in COL_YES_NO:
        add_yes_no_validation(ws, f"{col}{DATA_START_ROW}:{col}{last_row}")

    add_disposal_status_validation(
        ws, f"{COL_STATUS}{DATA_START_ROW}:{COL_STATUS}{last_row}"
    )


def build(ws: Worksheet, context: dict[str, Any]) -> None:
    """Build the Disposal Log sheet from disposal data."""
    df: pd.DataFrame = context["data"].get("disposal", pd.DataFrame())
    summary = _compute_summary(df)

    _write_title(ws, len(df))
    _write_summary_section(ws, summary)
    _write_headers(ws)
    last_row = _write_data_rows(ws, df)

    _apply_column_formats(ws, last_row)
    _create_disposal_table(ws, last_row)
    _apply_status_formatting(ws, last_row)
    _apply_validations(ws, last_row)

    freeze_panes(ws, row=DATA_START_ROW, col=1)
    autosize_columns(ws, min_width=10, max_width=36)
    ws.column_dimensions["D"].width = min(ws.column_dimensions["D"].width, 28)
    ws.column_dimensions["I"].width = min(ws.column_dimensions["I"].width, 24)
    ws.column_dimensions["J"].width = min(ws.column_dimensions["J"].width, 28)

    ws.sheet_view.showGridLines = False
    set_landscape_print(ws, fit_width=1, repeat_header_rows=f"{HEADER_ROW}:{HEADER_ROW}")
