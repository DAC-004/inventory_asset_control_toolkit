"""Aged Excess Analysis sheet — slow-moving, aged, excess, and obsolete inventory."""

from __future__ import annotations

from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from config import style_config as sc
from src.workbook.styles import (
    apply_risk_conditional_formatting,
    apply_table_header_style,
    freeze_panes,
)
from src.workbook.utils import (
    autosize_columns,
    create_excel_table,
    format_currency_columns,
    format_integer_columns,
    format_percentage_columns,
    set_landscape_print,
)
from src.workbook.validations import add_recommended_action_validation

HEADERS = [
    "SKU",
    "Product Name",
    "Location",
    "Quantity On Hand",
    "Max Stock",
    "Excess Quantity",
    "Age Days",
    "90 Day Demand",
    "Sell Through Rate",
    "Inventory Value",
    "Risk Level",
    "Issue Type",
    "Recommended Action",
    "Analyst Notes",
]

TITLE_ROW = 1
HEADER_ROW = 2
DATA_START_ROW = 3
COL_COUNT = len(HEADERS)
TABLE_NAME = "AgedExcessTable"

COL_INTEGER = ["D", "E", "F", "G", "H"]
COL_PERCENTAGE = ["I"]
COL_CURRENCY = ["J"]
COL_RISK = "K"
COL_RECOMMENDED_ACTION = "M"

RISK_LEVEL_ORDER = {"High": 0, "Medium": 1, "Low": 2}

RISK_LEVEL_CF_MAP = {
    "High": "critical",
    "Medium": "slow_moving",
    "Low": "healthy",
}

ISSUE_TYPE_MAP = {
    "Stockout Risk": "Stockout Risk",
    "Obsolete": "Obsolete",
    "Excess / Aged": "Excess / Aged",
    "Excess": "Excess",
    "Slow-Moving": "Slow-Moving",
    "Healthy": "Within Target",
}

ANALYST_NOTES = {
    "Stockout Risk": "Below minimum stock — prioritize replenishment to avoid lost sales.",
    "Obsolete": "No 90-day demand and aged over 365 days — recommend liquidation review.",
    "Excess / Aged": "Over max stock with extended age — evaluate transfer before markdown.",
    "Excess": "Quantity exceeds max stock — consider transfer to locations with shortages.",
    "Slow-Moving": "Sell-through below 15% over 180+ days — review markdown options.",
    "Healthy": "Inventory levels and movement are within target range.",
}


def _assign_risk_level(status: str) -> str:
    """Map inventory status to High / Medium / Low risk level."""
    if status in ("Obsolete", "Excess / Aged", "Stockout Risk"):
        return "High"
    if status in ("Excess", "Slow-Moving"):
        return "Medium"
    return "Low"


def _build_analysis_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    """Derive aged/excess analysis rows from master inventory data."""
    if df.empty:
        return pd.DataFrame(columns=HEADERS)

    working = df.copy()
    working["excess_quantity"] = (working["quantity_on_hand"] - working["max_stock"]).clip(lower=0)
    working["inventory_value"] = working["total_value"]
    working["risk_level"] = working["status"].map(_assign_risk_level)
    working["issue_type"] = working["status"].map(ISSUE_TYPE_MAP)
    working["analyst_notes"] = working["status"].map(ANALYST_NOTES)

    analysis = working[
        [
            "sku",
            "product_name",
            "location",
            "quantity_on_hand",
            "max_stock",
            "excess_quantity",
            "age_days",
            "demand_90_day",
            "sell_through_rate",
            "inventory_value",
            "risk_level",
            "issue_type",
            "recommended_action",
            "analyst_notes",
        ]
    ].rename(
        columns={
            "sku": "SKU",
            "product_name": "Product Name",
            "location": "Location",
            "quantity_on_hand": "Quantity On Hand",
            "max_stock": "Max Stock",
            "excess_quantity": "Excess Quantity",
            "age_days": "Age Days",
            "demand_90_day": "90 Day Demand",
            "sell_through_rate": "Sell Through Rate",
            "inventory_value": "Inventory Value",
            "risk_level": "Risk Level",
            "issue_type": "Issue Type",
            "recommended_action": "Recommended Action",
            "analyst_notes": "Analyst Notes",
        }
    )

    # Focus on items requiring analyst attention
    analysis = analysis[analysis["Risk Level"] != "Low"].copy()
    analysis["_risk_order"] = analysis["Risk Level"].map(RISK_LEVEL_ORDER)
    analysis = analysis.sort_values(
        ["_risk_order", "Inventory Value"],
        ascending=[True, False],
    ).drop(columns="_risk_order")
    return analysis.reset_index(drop=True)


def _write_title(ws: Worksheet, record_count: int) -> None:
    """Render sheet title with exception count."""
    ws.merge_cells(start_row=TITLE_ROW, start_column=1, end_row=TITLE_ROW, end_column=COL_COUNT)
    cell = ws.cell(
        row=TITLE_ROW,
        column=1,
        value=f"Aged Excess Analysis — {record_count:,} Exception Items",
    )
    cell.font = Font(
        name=sc.FONTS["default_name"],
        bold=True,
        color=sc.COLORS["white"],
        size=sc.FONTS["title_size"],
    )
    cell.fill = sc.HEADER_FILL
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[TITLE_ROW].height = 28


def _write_headers(ws: Worksheet) -> None:
    """Write and style table headers."""
    for col_idx, header in enumerate(HEADERS, start=1):
        ws.cell(row=HEADER_ROW, column=col_idx, value=header)
    apply_table_header_style(ws, HEADER_ROW, COL_COUNT)
    ws.row_dimensions[HEADER_ROW].height = 22


def _write_data_rows(ws: Worksheet, analysis: pd.DataFrame) -> int:
    """Write analysis rows and return the last populated row."""
    if analysis.empty:
        return HEADER_ROW

    for row_offset, (_, row) in enumerate(analysis.iterrows()):
        excel_row = DATA_START_ROW + row_offset
        for col_idx, header in enumerate(HEADERS, start=1):
            value = row[header]
            if pd.isna(value):
                value = None
            ws.cell(row=excel_row, column=col_idx, value=value)
    return DATA_START_ROW + len(analysis) - 1


def _apply_column_formats(ws: Worksheet, last_row: int) -> None:
    """Apply number formats to numeric columns."""
    if last_row < DATA_START_ROW:
        return
    format_integer_columns(ws, COL_INTEGER, DATA_START_ROW, last_row)
    format_percentage_columns(ws, COL_PERCENTAGE, DATA_START_ROW, last_row)
    format_currency_columns(ws, COL_CURRENCY, DATA_START_ROW, last_row)


def _create_analysis_table(ws: Worksheet, last_row: int) -> None:
    """Create filterable Excel table over the data range."""
    end_col = get_column_letter(COL_COUNT)
    if last_row < DATA_START_ROW:
        last_row = DATA_START_ROW
    create_excel_table(ws, TABLE_NAME, f"A{HEADER_ROW}:{end_col}{last_row}")


def _apply_risk_formatting(ws: Worksheet, last_row: int) -> None:
    """Apply conditional formatting to Risk Level column."""
    if last_row < DATA_START_ROW:
        return
    risk_range = f"{COL_RISK}{DATA_START_ROW}:{COL_RISK}{last_row}"
    apply_risk_conditional_formatting(
        ws,
        risk_range,
        COL_RISK,
        DATA_START_ROW,
        RISK_LEVEL_CF_MAP,
    )


def _apply_validations(ws: Worksheet, last_row: int) -> None:
    """Add dropdown validation for Recommended Action."""
    if last_row < DATA_START_ROW:
        return
    add_recommended_action_validation(
        ws, f"{COL_RECOMMENDED_ACTION}{DATA_START_ROW}:{COL_RECOMMENDED_ACTION}{last_row}"
    )


def build(ws: Worksheet, context: dict[str, Any]) -> None:
    """Build the Aged Excess Analysis sheet from inventory data."""
    df: pd.DataFrame = context["data"].get("inventory", pd.DataFrame())
    analysis = _build_analysis_dataframe(df)

    _write_title(ws, len(analysis))
    _write_headers(ws)
    last_row = _write_data_rows(ws, analysis)

    _apply_column_formats(ws, last_row)
    _create_analysis_table(ws, last_row)
    _apply_risk_formatting(ws, last_row)
    _apply_validations(ws, last_row)

    freeze_panes(ws, row=DATA_START_ROW, col=1)
    autosize_columns(ws, min_width=10, max_width=40)
    ws.column_dimensions["B"].width = min(ws.column_dimensions["B"].width, 34)
    ws.column_dimensions["N"].width = min(max(ws.column_dimensions["N"].width, 36), 52)

    ws.sheet_view.showGridLines = False
    set_landscape_print(ws, fit_width=1, repeat_header_rows=f"{HEADER_ROW}:{HEADER_ROW}")
