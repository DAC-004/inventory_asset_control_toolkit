"""
Excel styling constants: colors, fonts, number formats, and risk mappings.

Applied by workbook/styles.py and individual sheet builders.
Modify palette and formats here to restyle the entire workbook.
"""

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# ---------------------------------------------------------------------------
# Color palette (hex without # prefix — openpyxl format)
# ---------------------------------------------------------------------------

COLORS = {
    "navy": "1F4E78",
    "dark_blue": "17365D",
    "light_blue": "D9EAF7",
    "green": "70AD47",
    "yellow": "FFC000",
    "orange": "ED7D31",
    "red": "C00000",
    "gray": "F2F2F2",
    "white": "FFFFFF",
    "border_gray": "BFBFBF",
}

# Backward-compatible color aliases
NAVY = COLORS["navy"]
DARK_BLUE = COLORS["dark_blue"]
LIGHT_BLUE = COLORS["light_blue"]
GREEN = COLORS["green"]
YELLOW = COLORS["yellow"]
ORANGE = COLORS["orange"]
RED = COLORS["red"]
GRAY = COLORS["gray"]
WHITE = COLORS["white"]

# ---------------------------------------------------------------------------
# Fonts
# ---------------------------------------------------------------------------

FONTS = {
    "default_name": "Calibri",
    "header_size": 11,
    "title_size": 14,
    "body_size": 10,
    "kpi_value_size": 18,
}

DEFAULT_FONT_NAME = FONTS["default_name"]
HEADER_FONT_SIZE = FONTS["header_size"]
KPI_VALUE_FONT_SIZE = FONTS["kpi_value_size"]
BODY_FONT_SIZE = FONTS["body_size"]

# ---------------------------------------------------------------------------
# Number formats
# ---------------------------------------------------------------------------

NUMBER_FORMATS = {
    "currency": '"$"#,##0.00',
    "currency_compact": '"$"#,##0',
    "percentage": "0.0%",
    "percentage_whole": "0%",
    "date": "mm/dd/yyyy",
    "date_short": "m/d/yy",
    "integer": "#,##0",
    "decimal": "#,##0.00",
}

# ---------------------------------------------------------------------------
# Borders
# ---------------------------------------------------------------------------

THIN_GRAY_BORDER = Border(
    left=Side(style="thin", color=COLORS["border_gray"]),
    right=Side(style="thin", color=COLORS["border_gray"]),
    top=Side(style="thin", color=COLORS["border_gray"]),
    bottom=Side(style="thin", color=COLORS["border_gray"]),
)

NAVY_BORDER = Border(
    left=Side(style="thin", color=COLORS["navy"]),
    right=Side(style="thin", color=COLORS["navy"]),
    top=Side(style="thin", color=COLORS["navy"]),
    bottom=Side(style="thin", color=COLORS["navy"]),
)

# ---------------------------------------------------------------------------
# Header row style (specs §8.2 — navy fill, white bold text, centered)
# ---------------------------------------------------------------------------

HEADER_STYLE = {
    "fill": PatternFill(
        start_color=COLORS["navy"],
        end_color=COLORS["navy"],
        fill_type="solid",
    ),
    "font": Font(
        name=FONTS["default_name"],
        bold=True,
        color=COLORS["white"],
        size=FONTS["header_size"],
    ),
    "alignment": Alignment(horizontal="center", vertical="center", wrap_text=True),
    "border": THIN_GRAY_BORDER,
}

# Pre-built objects for direct cell assignment
HEADER_FILL = HEADER_STYLE["fill"]
HEADER_FONT = HEADER_STYLE["font"]
CENTER_ALIGN = HEADER_STYLE["alignment"]

# ---------------------------------------------------------------------------
# KPI card style (specs §8.3 — light blue fill, dark blue title, large value)
# ---------------------------------------------------------------------------

KPI_CARD_STYLE = {
    "fill": PatternFill(
        start_color=COLORS["light_blue"],
        end_color=COLORS["light_blue"],
        fill_type="solid",
    ),
    "title_font": Font(
        name=FONTS["default_name"],
        bold=True,
        color=COLORS["dark_blue"],
        size=FONTS["body_size"],
    ),
    "value_font": Font(
        name=FONTS["default_name"],
        bold=True,
        color=COLORS["navy"],
        size=FONTS["kpi_value_size"],
    ),
    "border": NAVY_BORDER,
    "alignment": Alignment(horizontal="center", vertical="center", wrap_text=True),
}

KPI_CARD_FILL = KPI_CARD_STYLE["fill"]
KPI_TITLE_FONT = KPI_CARD_STYLE["title_font"]
KPI_VALUE_FONT = KPI_CARD_STYLE["value_font"]

# ---------------------------------------------------------------------------
# General-purpose fonts and fills
# ---------------------------------------------------------------------------

TITLE_FONT = Font(
    name=FONTS["default_name"],
    bold=True,
    color=COLORS["dark_blue"],
    size=FONTS["title_size"],
)
BODY_FONT = Font(name=FONTS["default_name"], size=FONTS["body_size"])
ALT_ROW_FILL = PatternFill(
    start_color=COLORS["gray"],
    end_color=COLORS["gray"],
    fill_type="solid",
)
LEFT_ALIGN = Alignment(horizontal="left", vertical="center", wrap_text=True)

# ---------------------------------------------------------------------------
# Risk color mappings (specs §8.4)
# Maps operational labels → fill color for conditional formatting
# ---------------------------------------------------------------------------

RISK_LEVEL_COLORS = {
    "healthy": COLORS["green"],
    "compliant": COLORS["green"],
    "watch": COLORS["yellow"],
    "renewal_watch": COLORS["yellow"],
    "slow_moving": COLORS["orange"],
    "pending": COLORS["orange"],
    "critical": COLORS["red"],
    "obsolete": COLORS["red"],
    "missing": COLORS["red"],
    "over_assigned": COLORS["red"],
    "neutral": COLORS["gray"],
}

# Inventory status → risk level
INVENTORY_STATUS_RISK = {
    "Healthy": "healthy",
    "Stockout Risk": "critical",
    "Excess": "watch",
    "Slow-Moving": "slow_moving",
    "Excess / Aged": "slow_moving",
    "Obsolete": "obsolete",
}

# Backward-compatible alias
RISK_COLORS = RISK_LEVEL_COLORS


def risk_fill(status: str, mapping: dict[str, str]) -> PatternFill:
    """Return a PatternFill for a status label using the given status→risk mapping."""
    level = mapping.get(status, "neutral")
    color = RISK_LEVEL_COLORS.get(level, COLORS["gray"])
    return PatternFill(start_color=color, end_color=color, fill_type="solid")
