"""Inventory Classification sheet — ABC, turnover, DOH, and accuracy."""

from __future__ import annotations

from typing import Any

import pandas as pd
from openpyxl.chart import Reference
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from config import style_config as sc
from src.services.classification_service import (
    CLASSIFICATION_HEADERS,
    build_classification_dataframe,
)
from src.workbook.charts import add_bar_chart, add_pie_chart, make_category_reference
from src.workbook.styles import (
    apply_kpi_card_style,
    apply_risk_conditional_formatting,
    apply_section_header_style,
    apply_table_header_style,
    freeze_panes,
)
from src.workbook.utils import (
    autosize_columns,
    create_excel_table,
    format_currency_columns,
    format_date_columns,
    format_integer_columns,
    format_percentage_columns,
    set_landscape_print,
)

HEADERS = CLASSIFICATION_HEADERS

TITLE_ROW = 1
KPI_SECTION_ROW = 2
KPI_TITLE_ROW = 3
KPI_VALUE_ROW = 4
TABLE_SECTION_ROW = 6
HEADER_ROW = 7
DATA_START_ROW = 8
COL_COUNT = len(HEADERS)
TABLE_NAME = "InventoryClassificationTable"

COL_CURRENCY = ["F", "G", "I", "L", "V"]
COL_INTEGER = ["E", "H", "U", "W"]
COL_PERCENTAGE = ["J", "S"]
COL_DATES = ["X", "Y"]
COL_ABC = "K"

ABC_CF_MAP = {"A": "healthy", "B": "watch", "C": "neutral"}


def _write_title(ws: Worksheet, count: int) -> None:
    ws.merge_cells(start_row=TITLE_ROW, start_column=1, end_row=TITLE_ROW, end_column=8)
    cell = ws.cell(
        row=TITLE_ROW, column=1, value=f"Inventory Classification — {count:,} SKUs"
    )
    cell.font = Font(
        name=sc.FONTS["default_name"],
        bold=True,
        color=sc.COLORS["white"],
        size=sc.FONTS["title_size"],
    )
    cell.fill = sc.HEADER_FILL
    ws.row_dimensions[TITLE_ROW].height = 28


def _write_kpi_cards(ws: Worksheet, df: pd.DataFrame) -> None:
    apply_section_header_style(
        ws, KPI_SECTION_ROW, 1, "Classification KPIs", span_cols=10
    )
    if df.empty:
        return
    abc = df["ABC Class"].value_counts()
    kpis = [
        ("Class A SKUs", int(abc.get("A", 0)), "integer"),
        ("Class B SKUs", int(abc.get("B", 0)), "integer"),
        ("Class C SKUs", int(abc.get("C", 0)), "integer"),
        ("Avg Turnover", round(float(df["Inventory Turnover"].mean()), 2), "decimal"),
        ("Avg Unit Accuracy", float(df["Unit Accuracy %"].mean()), "percentage"),
    ]
    card_cols = [1, 3, 5, 7, 9]
    for idx, (title, value, fmt) in enumerate(kpis):
        col = card_cols[idx]
        apply_kpi_card_style(
            ws, KPI_TITLE_ROW, col, KPI_VALUE_ROW, col, title, value, span_cols=2
        )
        cell = ws.cell(row=KPI_VALUE_ROW, column=col)
        if fmt == "percentage":
            cell.number_format = sc.NUMBER_FORMATS["percentage"]
        elif fmt == "integer":
            cell.number_format = sc.NUMBER_FORMATS["integer"]
        else:
            cell.number_format = "0.00"


def _write_headers(ws: Worksheet) -> None:
    apply_section_header_style(
        ws, TABLE_SECTION_ROW, 1, "Enterprise SKU Classification", span_cols=8
    )
    for col_idx, header in enumerate(HEADERS, start=1):
        ws.cell(row=HEADER_ROW, column=col_idx, value=header)
    apply_table_header_style(ws, HEADER_ROW, COL_COUNT)


def _write_data(ws: Worksheet, df: pd.DataFrame) -> int:
    if df.empty:
        return HEADER_ROW
    for offset, (_, row) in enumerate(df.iterrows()):
        excel_row = DATA_START_ROW + offset
        for col_idx, header in enumerate(HEADERS, start=1):
            val = row[header]
            if pd.isna(val):
                val = None
            ws.cell(row=excel_row, column=col_idx, value=val)
    return DATA_START_ROW + len(df) - 1


def _apply_formats(ws: Worksheet, last_row: int) -> None:
    if last_row < DATA_START_ROW:
        return
    format_currency_columns(ws, COL_CURRENCY, DATA_START_ROW, last_row)
    format_integer_columns(ws, COL_INTEGER, DATA_START_ROW, last_row)
    format_percentage_columns(ws, COL_PERCENTAGE, DATA_START_ROW, last_row)
    format_date_columns(ws, COL_DATES, DATA_START_ROW, last_row)


def _add_charts(ws: Worksheet, df: pd.DataFrame, last_row: int) -> None:
    if df.empty or last_row < DATA_START_ROW:
        return
    abc_summary = df["ABC Class"].value_counts().reset_index()
    abc_summary.columns = ["Class", "Count"]
    start = last_row + 3
    for i, row in abc_summary.iterrows():
        ws.cell(row=start + i, column=1, value=row["Class"])
        ws.cell(row=start + i, column=2, value=int(row["Count"]))
    if len(abc_summary) == 0:
        return
    cats = make_category_reference(ws, 1, start, start + len(abc_summary) - 1)
    data = Reference(
        ws, min_col=2, min_row=start, max_col=2, max_row=start + len(abc_summary) - 1
    )
    add_pie_chart(ws, "ABC Class Mix", data, cats, anchor="AA3", width=12, height=10)

    top = df.nsmallest(min(10, len(df)), "Value Rank")
    tstart = start + len(abc_summary) + 2
    for i, (_, row) in enumerate(top.iterrows()):
        ws.cell(row=tstart + i, column=1, value=row["SKU"])
        ws.cell(row=tstart + i, column=2, value=float(row["Annual Usage Value"]))
    cats2 = make_category_reference(ws, 1, tstart, tstart + len(top) - 1)
    data2 = Reference(
        ws, min_col=2, min_row=tstart, max_col=2, max_row=tstart + len(top) - 1
    )
    add_bar_chart(
        ws,
        "Top Usage Value SKUs",
        data2,
        cats2,
        anchor="AA18",
        width=14,
        height=10,
        y_axis_title="Usage Value ($)",
    )


def build(ws: Worksheet, context: dict[str, Any]) -> None:
    """Build Inventory Classification sheet."""
    df = build_classification_dataframe(context["data"])
    _write_title(ws, len(df))
    _write_kpi_cards(ws, df)
    _write_headers(ws)
    last_row = _write_data(ws, df)
    end_col = get_column_letter(COL_COUNT)
    if last_row >= DATA_START_ROW:
        create_excel_table(ws, TABLE_NAME, f"A{HEADER_ROW}:{end_col}{last_row}")
        abc_range = f"{COL_ABC}{DATA_START_ROW}:{COL_ABC}{last_row}"
        apply_risk_conditional_formatting(
            ws, abc_range, COL_ABC, DATA_START_ROW, ABC_CF_MAP
        )
    _apply_formats(ws, last_row)
    _add_charts(ws, df, last_row)
    freeze_panes(ws, row=DATA_START_ROW, col=1)
    autosize_columns(ws, min_width=10, max_width=36)
    ws.sheet_view.showGridLines = False
    set_landscape_print(
        ws, fit_width=1, repeat_header_rows=f"{HEADER_ROW}:{HEADER_ROW}"
    )
