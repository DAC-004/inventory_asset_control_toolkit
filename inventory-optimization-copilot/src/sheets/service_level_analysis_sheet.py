"""Service Level Analysis sheet — fill rates, gaps, and recommendations."""

from __future__ import annotations

from typing import Any

import pandas as pd
from openpyxl.chart import Reference
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from config import style_config as sc
from src.services.service_level_service import (
    SERVICE_LEVEL_HEADERS,
    build_service_level_dataframe,
    compute_service_level_kpis,
)
from src.workbook.charts import add_bar_chart, add_pie_chart, make_category_reference
from src.workbook.styles import (
    apply_kpi_card_style,
    apply_risk_conditional_formatting,
    apply_section_header_style,
    apply_table_header_style,
    freeze_panes,
)
from src.workbook.utils import (
    autosize_columns,
    create_excel_table,
    format_integer_columns,
    format_percentage_columns,
    set_landscape_print,
)

HEADERS = SERVICE_LEVEL_HEADERS

TITLE_ROW = 1
KPI_SECTION_ROW = 2
KPI_TITLE_ROW = 3
KPI_VALUE_ROW = 4
TABLE_SECTION_ROW = 6
HEADER_ROW = 7
DATA_START_ROW = 8
COL_COUNT = len(HEADERS)
TABLE_NAME = "ServiceLevelAnalysisTable"

STATUS_CF_MAP = {
    "On Target": "healthy",
    "Watch": "watch",
    "Below Target": "slow_moving",
    "Critical": "critical",
    "Data Review Required": "obsolete",
}


def _col(header: str) -> str:
    return str(get_column_letter(HEADERS.index(header) + 1))


def _write_title(ws: Worksheet, count: int) -> None:
    ws.merge_cells(
        start_row=TITLE_ROW, start_column=1, end_row=TITLE_ROW, end_column=10
    )
    cell = ws.cell(
        row=TITLE_ROW,
        column=1,
        value=f"Service Level Analysis — {count:,} Segments",
    )
    cell.font = Font(
        name=sc.FONTS["default_name"],
        bold=True,
        color=sc.COLORS["white"],
        size=sc.FONTS["title_size"],
    )
    cell.fill = sc.HEADER_FILL
    ws.row_dimensions[TITLE_ROW].height = 28


def _write_kpi_cards(ws: Worksheet, df: pd.DataFrame) -> None:
    apply_section_header_style(ws, KPI_SECTION_ROW, 1, "Service KPIs", span_cols=10)
    kpis = compute_service_level_kpis(df)
    cards = [
        ("Avg Unit Fill Rate", kpis["avg_unit_fill_rate"], "percentage"),
        ("Avg Line Fill Rate", kpis["avg_line_fill_rate"], "percentage"),
        ("On Target", kpis["on_target_count"], "integer"),
        ("Below Target", kpis["below_target_count"], "integer"),
        ("Critical", kpis["critical_count"], "integer"),
    ]
    card_cols = [1, 3, 5, 7, 9]
    for idx, (title, value, fmt) in enumerate(cards):
        col = card_cols[idx]
        apply_kpi_card_style(
            ws, KPI_TITLE_ROW, col, KPI_VALUE_ROW, col, title, value, span_cols=2
        )
        cell = ws.cell(row=KPI_VALUE_ROW, column=col)
        if fmt == "percentage":
            cell.number_format = sc.NUMBER_FORMATS["percentage"]
        elif fmt == "integer":
            cell.number_format = sc.NUMBER_FORMATS["integer"]


def _write_headers(ws: Worksheet) -> None:
    apply_section_header_style(
        ws, TABLE_SECTION_ROW, 1, "Service Performance by Segment", span_cols=10
    )
    for col_idx, header in enumerate(HEADERS, start=1):
        ws.cell(row=HEADER_ROW, column=col_idx, value=header)
    apply_table_header_style(ws, HEADER_ROW, COL_COUNT)


def _write_data(ws: Worksheet, df: pd.DataFrame) -> int:
    if df.empty:
        return HEADER_ROW
    for offset, (_, row) in enumerate(df.iterrows()):
        excel_row = DATA_START_ROW + offset
        for col_idx, header in enumerate(HEADERS, start=1):
            val = row[header]
            if pd.isna(val):
                val = None
            ws.cell(row=excel_row, column=col_idx, value=val)
    return DATA_START_ROW + len(df) - 1


def _apply_formats(ws: Worksheet, last_row: int) -> None:
    if last_row < DATA_START_ROW:
        return
    integer_headers = [
        "Order Count",
        "Line Count",
        "Units Ordered",
        "Units Fulfilled Immediately",
        "Backorder Units",
        "Lost Sales Units",
        "Stockout Weeks",
    ]
    pct_headers = [
        "Unit Fill Rate",
        "Line Fill Rate",
        "Order Fill Rate",
        "Backorder Rate",
        "Lost Sales Rate",
        "Stockout Frequency",
        "Service Level Target",
    ]
    format_integer_columns(
        ws, [_col(h) for h in integer_headers], DATA_START_ROW, last_row
    )
    format_percentage_columns(
        ws, [_col(h) for h in pct_headers], DATA_START_ROW, last_row
    )


def _add_charts(ws: Worksheet, df: pd.DataFrame, last_row: int) -> None:
    if df.empty or last_row < DATA_START_ROW:
        return
    status_summary = df["Status"].value_counts().reset_index()
    status_summary.columns = ["Status", "Count"]
    start = last_row + 3
    for i, row in status_summary.iterrows():
        ws.cell(row=start + i, column=1, value=row["Status"])
        ws.cell(row=start + i, column=2, value=int(row["Count"]))
    cats = make_category_reference(ws, 1, start, start + len(status_summary) - 1)
    data = Reference(
        ws,
        min_col=2,
        min_row=start,
        max_col=2,
        max_row=start + len(status_summary) - 1,
    )
    add_pie_chart(
        ws,
        "Service Status Mix",
        data,
        cats,
        anchor="AA3",
        width=12,
        height=10,
    )

    top_gap = df.nlargest(min(10, len(df)), "Service Gap")
    if not top_gap.empty:
        tstart = start + len(status_summary) + 2
        for i, (_, row) in enumerate(top_gap.iterrows()):
            label = f"{row['Location ID']}-{row['Category'][:8]}"
            ws.cell(row=tstart + i, column=1, value=label)
            ws.cell(row=tstart + i, column=2, value=float(row["Service Gap"]))
        cats2 = make_category_reference(ws, 1, tstart, tstart + len(top_gap) - 1)
        data2 = Reference(
            ws,
            min_col=2,
            min_row=tstart,
            max_col=2,
            max_row=tstart + len(top_gap) - 1,
        )
        add_bar_chart(
            ws,
            "Largest Service Gaps",
            data2,
            cats2,
            anchor="AA18",
            width=14,
            height=10,
            y_axis_title="Gap",
        )


def build(ws: Worksheet, context: dict[str, Any]) -> None:
    """Build Service Level Analysis sheet."""
    df = build_service_level_dataframe(context["data"])
    _write_title(ws, len(df))
    _write_kpi_cards(ws, df)
    _write_headers(ws)
    last_row = _write_data(ws, df)
    end_col = get_column_letter(COL_COUNT)
    if last_row >= DATA_START_ROW:
        create_excel_table(ws, TABLE_NAME, f"A{HEADER_ROW}:{end_col}{last_row}")
        status_col = _col("Status")
        apply_risk_conditional_formatting(
            ws,
            f"{status_col}{DATA_START_ROW}:{status_col}{last_row}",
            status_col,
            DATA_START_ROW,
            STATUS_CF_MAP,
        )
    _apply_formats(ws, last_row)
    _add_charts(ws, df, last_row)
    freeze_panes(ws, row=DATA_START_ROW, col=1)
    autosize_columns(ws, min_width=10, max_width=36)
    ws.sheet_view.showGridLines = False
    set_landscape_print(
        ws, fit_width=1, repeat_header_rows=f"{HEADER_ROW}:{HEADER_ROW}"
    )
