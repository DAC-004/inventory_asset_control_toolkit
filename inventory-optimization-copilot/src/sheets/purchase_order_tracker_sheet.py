"""Purchase Order Tracker sheet — open POs, lateness, and price variance."""

from __future__ import annotations

from typing import Any

import pandas as pd
from openpyxl.chart import Reference
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from config import style_config as sc
from src.services.purchase_order_service import (
    PO_TRACKER_HEADERS,
    build_po_tracker_dataframe,
    compute_po_tracker_kpis,
)
from src.workbook.charts import add_bar_chart, add_pie_chart, make_category_reference
from src.workbook.styles import (
    apply_kpi_card_style,
    apply_risk_conditional_formatting,
    apply_section_header_style,
    apply_table_header_style,
    freeze_panes,
)
from src.workbook.utils import (
    autosize_columns,
    create_excel_table,
    format_currency_columns,
    format_date_columns,
    format_integer_columns,
    set_landscape_print,
)

HEADERS = PO_TRACKER_HEADERS

TITLE_ROW = 1
KPI_SECTION_ROW = 2
KPI_TITLE_ROW = 3
KPI_VALUE_ROW = 4
TABLE_SECTION_ROW = 6
HEADER_ROW = 7
DATA_START_ROW = 8
COL_COUNT = len(HEADERS)
TABLE_NAME = "PurchaseOrderTrackerTable"

STATUS_CF_MAP = {
    "Late": "critical",
    "Quality Hold": "slow_moving",
    "Over-Received": "watch",
    "Open": "neutral",
    "Approved": "healthy",
    "Partially Received": "watch",
    "Received": "healthy",
    "Closed": "healthy",
    "Cancelled": "obsolete",
    "Draft": "neutral",
}


def _col(header: str) -> str:
    return str(get_column_letter(HEADERS.index(header) + 1))


def _write_title(ws: Worksheet, count: int) -> None:
    ws.merge_cells(
        start_row=TITLE_ROW, start_column=1, end_row=TITLE_ROW, end_column=10
    )
    cell = ws.cell(
        row=TITLE_ROW,
        column=1,
        value=f"Purchase Order Tracker — {count:,} PO Lines",
    )
    cell.font = Font(
        name=sc.FONTS["default_name"],
        bold=True,
        color=sc.COLORS["white"],
        size=sc.FONTS["title_size"],
    )
    cell.fill = sc.HEADER_FILL
    ws.row_dimensions[TITLE_ROW].height = 28


def _write_kpi_cards(ws: Worksheet, df: pd.DataFrame) -> None:
    apply_section_header_style(ws, KPI_SECTION_ROW, 1, "Open PO KPIs", span_cols=10)
    kpis = compute_po_tracker_kpis(df)
    cards = [
        ("Open Lines", kpis["open_lines"], "integer"),
        ("Open PO Value", kpis["open_value"], "currency"),
        ("Late Lines", kpis["late_lines"], "integer"),
        ("Quality Holds", kpis["quality_holds"], "integer"),
    ]
    card_cols = [1, 3, 5, 7]
    for idx, (title, value, fmt) in enumerate(cards):
        col = card_cols[idx]
        apply_kpi_card_style(
            ws, KPI_TITLE_ROW, col, KPI_VALUE_ROW, col, title, value, span_cols=2
        )
        cell = ws.cell(row=KPI_VALUE_ROW, column=col)
        if fmt == "integer":
            cell.number_format = sc.NUMBER_FORMATS["integer"]
        elif fmt == "currency":
            cell.number_format = sc.NUMBER_FORMATS["currency_compact"]


def _write_headers(ws: Worksheet) -> None:
    apply_section_header_style(
        ws, TABLE_SECTION_ROW, 1, "Purchase Order Line Detail", span_cols=10
    )
    for col_idx, header in enumerate(HEADERS, start=1):
        ws.cell(row=HEADER_ROW, column=col_idx, value=header)
    apply_table_header_style(ws, HEADER_ROW, COL_COUNT)


def _write_data(ws: Worksheet, df: pd.DataFrame) -> int:
    if df.empty:
        return HEADER_ROW
    for offset, (_, row) in enumerate(df.iterrows()):
        excel_row = DATA_START_ROW + offset
        for col_idx, header in enumerate(HEADERS, start=1):
            val = row[header]
            if pd.isna(val):
                val = None
            ws.cell(row=excel_row, column=col_idx, value=val)
    return DATA_START_ROW + len(df) - 1


def _apply_formats(ws: Worksheet, last_row: int) -> None:
    if last_row < DATA_START_ROW:
        return
    integer_headers = [
        "Ordered Quantity",
        "Received Quantity",
        "Accepted Quantity",
        "Rejected Quantity",
        "Open Quantity",
        "Days Late",
    ]
    currency_headers = [
        "Unit Cost",
        "Extended Cost",
        "Open PO Value",
        "Baseline Unit Cost",
        "Purchase Price Variance",
    ]
    date_headers = [
        "Order Date",
        "Approval Date",
        "Promised Date",
        "Expected Date",
        "First Receipt Date",
        "Final Receipt Date",
    ]
    format_integer_columns(
        ws, [_col(h) for h in integer_headers], DATA_START_ROW, last_row
    )
    format_currency_columns(
        ws, [_col(h) for h in currency_headers], DATA_START_ROW, last_row
    )
    format_date_columns(ws, [_col(h) for h in date_headers], DATA_START_ROW, last_row)


def _add_charts(ws: Worksheet, df: pd.DataFrame, last_row: int) -> None:
    if df.empty or last_row < DATA_START_ROW:
        return
    status_summary = df["PO Status"].value_counts().reset_index()
    status_summary.columns = ["Status", "Count"]
    start = last_row + 3
    for i, row in status_summary.iterrows():
        ws.cell(row=start + i, column=1, value=row["Status"])
        ws.cell(row=start + i, column=2, value=int(row["Count"]))
    cats = make_category_reference(ws, 1, start, start + len(status_summary) - 1)
    data = Reference(
        ws,
        min_col=2,
        min_row=start,
        max_col=2,
        max_row=start + len(status_summary) - 1,
    )
    add_pie_chart(
        ws,
        "PO Status Mix",
        data,
        cats,
        anchor="AF3",
        width=12,
        height=10,
    )

    late = df[df["Days Late"] > 0].nlargest(min(10, len(df)), "Days Late")
    if not late.empty:
        tstart = start + len(status_summary) + 2
        for i, (_, row) in enumerate(late.iterrows()):
            ws.cell(row=tstart + i, column=1, value=str(row["PO Line"]))
            ws.cell(row=tstart + i, column=2, value=int(row["Days Late"]))
        cats2 = make_category_reference(ws, 1, tstart, tstart + len(late) - 1)
        data2 = Reference(
            ws,
            min_col=2,
            min_row=tstart,
            max_col=2,
            max_row=tstart + len(late) - 1,
        )
        add_bar_chart(
            ws,
            "Days Late — Top PO Lines",
            data2,
            cats2,
            anchor="AF18",
            width=14,
            height=10,
            y_axis_title="Days",
        )


def build(ws: Worksheet, context: dict[str, Any]) -> None:
    """Build Purchase Order Tracker sheet."""
    df = build_po_tracker_dataframe(context["data"])
    _write_title(ws, len(df))
    _write_kpi_cards(ws, df)
    _write_headers(ws)
    last_row = _write_data(ws, df)
    end_col = get_column_letter(COL_COUNT)
    if last_row >= DATA_START_ROW:
        create_excel_table(ws, TABLE_NAME, f"A{HEADER_ROW}:{end_col}{last_row}")
        status_col = _col("PO Status")
        apply_risk_conditional_formatting(
            ws,
            f"{status_col}{DATA_START_ROW}:{status_col}{last_row}",
            status_col,
            DATA_START_ROW,
            STATUS_CF_MAP,
        )
    _apply_formats(ws, last_row)
    _add_charts(ws, df, last_row)
    freeze_panes(ws, row=DATA_START_ROW, col=1)
    autosize_columns(ws, min_width=10, max_width=36)
    ws.sheet_view.showGridLines = False
    set_landscape_print(
        ws, fit_width=1, repeat_header_rows=f"{HEADER_ROW}:{HEADER_ROW}"
    )
