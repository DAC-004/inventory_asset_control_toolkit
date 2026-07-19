"""Inventory Dashboard — executive A:V table/chart grid."""

from __future__ import annotations

from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from config import style_config as sc
from config.workbook_config import AS_OF_DATE
from src.domain.constants import MASTER_DATA_START
from src.services.classification_service import compute_dashboard_classification_kpis
from src.services.kpi_dashboard_service import (
    compute_abc_usage_summary,
    compute_action_summary,
    compute_aging_summary,
    compute_fill_rate_by_location,
    compute_location_summary,
    compute_planning_service_kpis,
    compute_procurement_network_kpis,
    compute_replenishment_status_summary,
    compute_status_summary,
    compute_top_excess,
    compute_transfer_benefit_summary,
    compute_vendor_risk_summary,
)
from src.services.kpi_service import dashboard_kpi_definitions
from src.sheets.dashboard_layout import (
    CHART_ANCHOR_COL,
    CHART_AREA_START_COL,
    CHART_LABEL_COL,
    CHART_WIDTH_CM,
    DASHBOARD_CANVAS_COLS,
    DASHBOARD_COLUMN_WIDTHS,
    DashboardSectionLayout,
    DATA_ROW_HEIGHT,
    FREEZE_PANE,
    SECTION_TITLE_HEIGHT,
    TABLE_AREA_END_COL,
    TABLE_HEADER_HEIGHT,
    chart_height_cm,
    compute_section_layouts,
)
from src.workbook.charts import (
    add_bar_chart,
    add_clustered_bar_chart,
    make_category_reference,
    make_data_reference,
)
from src.workbook.utils import (
    add_internal_sheet_link,
    create_excel_table,
    format_currency_columns,
    format_integer_columns,
    format_percentage_columns,
    set_column_widths,
    set_landscape_print,
)

DASHBOARD_TABLE_PREFIX = "Dashboard"


def _data_end_row(record_count: int) -> int:
    if record_count <= 0:
        return MASTER_DATA_START
    return MASTER_DATA_START + record_count - 1


def _clear_dashboard_charts(ws: Worksheet) -> int:
    count = len(ws._charts)
    ws._charts.clear()
    return count


def _apply_canvas_columns(ws: Worksheet) -> None:
    set_column_widths(ws, DASHBOARD_COLUMN_WIDTHS)
    for col_idx in range(1, DASHBOARD_CANVAS_COLS + 1):
        letter = get_column_letter(col_idx)
        dim = ws.column_dimensions[letter]
        dim.hidden = False
        if dim.width is None or dim.width <= 0:
            dim.width = DASHBOARD_COLUMN_WIDTHS.get(letter, 12)
    # Column H holds chart category helpers — keep off-screen for reviewers.
    ws.column_dimensions[get_column_letter(CHART_LABEL_COL)].hidden = True


def _apply_kpi_value_format(
    ws: Worksheet, row: int, col: int, label: str, fmt: str
) -> None:
    cell = ws.cell(row=row, column=col)
    if fmt == "currency":
        cell.number_format = sc.NUMBER_FORMATS["currency_compact"]
    elif fmt == "percentage":
        cell.number_format = sc.NUMBER_FORMATS["percentage"]
    elif fmt == "turnover":
        cell.number_format = '0.0"x"'
    elif fmt == "days":
        cell.number_format = '#,##0" days"'
    elif fmt == "decimal":
        cell.number_format = "0.00"
    elif "Turnover" in label:
        cell.number_format = '0.0"x"'
    elif "DOH" in label or "Coverage Days" in label:
        cell.number_format = '#,##0" days"'
    else:
        cell.number_format = sc.NUMBER_FORMATS["integer"]


def _write_header_and_kpis(ws: Worksheet, end_row: int, data: dict[str, Any]) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=21)
    title_cell = ws.cell(
        row=1,
        column=1,
        value=f"Inventory Dashboard — As of {AS_OF_DATE.strftime('%B %d, %Y')}",
    )
    title_cell.font = Font(
        name=sc.FONTS["default_name"],
        bold=True,
        color=sc.COLORS["white"],
        size=sc.FONTS["title_size"],
    )
    title_cell.fill = sc.HEADER_FILL
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 30
    add_internal_sheet_link(ws, 1, 22, "README", "← README")

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=22)
    desc = ws.cell(
        row=2,
        column=1,
        value=(
            "Executive inventory health, service, replenishment, supplier risk, and "
            "network optimization — where exposure is concentrated and where action is required."
        ),
    )
    desc.font = Font(name=sc.FONTS["default_name"], size=sc.FONTS["body_size"])
    desc.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 24

    kpi_cards: list[tuple[str, Any, str]] = [
        (title, formula, fmt)
        for title, formula, fmt in dashboard_kpi_definitions(MASTER_DATA_START, end_row)
    ]
    class_kpis = compute_dashboard_classification_kpis(data)
    planning = compute_planning_service_kpis(data)
    procurement = compute_procurement_network_kpis(data)
    kpi_cards.extend(
        [
            ("Inventory Turnover", class_kpis["enterprise_turnover"], "turnover"),
            ("Financial DOH", class_kpis["avg_financial_doh"], "days"),
            ("Unit Fill Rate", planning["unit_fill_rate"], "percentage"),
            ("Line Fill Rate", planning["line_fill_rate"], "percentage"),
            ("Open PO Value", procurement["open_po_value"], "currency"),
            ("Transfer Net Benefit", procurement["transfer_net_benefit"], "currency"),
            (
                "Recommended Order Value",
                planning["recommended_order_value"],
                "currency",
            ),
            ("High-Risk Vendors", procurement["high_risk_vendors"], "integer"),
        ]
    )

    cards_per_row = 6
    span_cols = 3
    start_row = 3
    for idx, (label, value, fmt) in enumerate(kpi_cards[:24]):
        band = idx // cards_per_row
        col = 1 + (idx % cards_per_row) * span_cols
        title_row = start_row + band * 2
        value_row = title_row + 1
        end_col = col + span_cols - 1
        ws.merge_cells(
            start_row=title_row, start_column=col, end_row=title_row, end_column=end_col
        )
        ws.merge_cells(
            start_row=value_row, start_column=col, end_row=value_row, end_column=end_col
        )
        tcell = ws.cell(row=title_row, column=col, value=label)
        tcell.font = sc.KPI_TITLE_FONT
        tcell.fill = sc.KPI_CARD_FILL
        tcell.alignment = sc.KPI_CARD_STYLE["alignment"]
        vcell = ws.cell(row=value_row, column=col, value=value)
        vcell.font = sc.KPI_VALUE_FONT
        vcell.fill = sc.KPI_CARD_FILL
        vcell.alignment = sc.KPI_CARD_STYLE["alignment"]
        _apply_kpi_value_format(ws, value_row, col, label, fmt)
        ws.row_dimensions[title_row].height = 18
        ws.row_dimensions[value_row].height = 24


def _write_section_title(ws: Worksheet, section_row: int, title: str) -> None:
    ws.merge_cells(
        start_row=section_row,
        start_column=1,
        end_row=section_row,
        end_column=DASHBOARD_CANVAS_COLS,
    )
    cell = ws.cell(row=section_row, column=1, value=title)
    cell.font = Font(
        name=sc.FONTS["default_name"],
        bold=True,
        color=sc.COLORS["white"],
        size=sc.FONTS["body_size"],
    )
    cell.fill = PatternFill(
        start_color=sc.COLORS["navy"],
        end_color=sc.COLORS["navy"],
        fill_type="solid",
    )
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[section_row].height = SECTION_TITLE_HEIGHT


def _write_section_context(ws: Worksheet, row: int, context: str) -> None:
    ws.merge_cells(
        start_row=row, start_column=1, end_row=row, end_column=TABLE_AREA_END_COL
    )
    left = ws.cell(row=row, column=1, value="Summary table")
    left.font = Font(
        name=sc.FONTS["default_name"], italic=True, size=9, color=sc.COLORS["dark_blue"]
    )
    ws.merge_cells(
        start_row=row,
        start_column=CHART_AREA_START_COL,
        end_row=row,
        end_column=DASHBOARD_CANVAS_COLS,
    )
    right = ws.cell(row=row, column=CHART_AREA_START_COL, value=context)
    right.font = Font(
        name=sc.FONTS["default_name"], italic=True, size=9, color=sc.COLORS["dark_blue"]
    )
    right.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 16


def _write_section_table(
    ws: Worksheet,
    section: DashboardSectionLayout,
    headers: list[str],
    rows: list[list[Any]],
    table_name: str,
    chart_labels: list[str] | None = None,
) -> dict[str, int]:
    _write_section_title(ws, section.start_row, section.title)
    context_row = section.start_row + 1
    _write_section_context(ws, context_row, section.context)
    header_row = section.start_row + 2

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    ws.row_dimensions[header_row].height = TABLE_HEADER_HEIGHT

    first_data_row = header_row + 1
    wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)
    num_align = Alignment(horizontal="right", vertical="center")
    center_align = Alignment(horizontal="center", vertical="center")
    for row_offset, row_values in enumerate(rows):
        excel_row = first_data_row + row_offset
        ws.row_dimensions[excel_row].height = DATA_ROW_HEIGHT
        ws.row_dimensions[excel_row].hidden = False
        for col_idx, value in enumerate(row_values, start=1):
            cell = ws.cell(row=excel_row, column=col_idx, value=value)
            if col_idx == 1 and isinstance(value, int):
                cell.alignment = center_align
            elif col_idx > 1:
                cell.alignment = num_align
            else:
                cell.alignment = wrap

    if chart_labels:
        for row_offset, label in enumerate(chart_labels):
            label_cell = ws.cell(
                row=first_data_row + row_offset, column=CHART_LABEL_COL, value=label
            )
            label_cell.alignment = wrap

    last_data_row = first_data_row + len(rows) - 1 if rows else header_row
    table_cols = min(
        len(headers),
        TABLE_AREA_END_COL - 1 if chart_labels else TABLE_AREA_END_COL,
    )
    table_end_col = get_column_letter(table_cols)
    create_excel_table(
        ws,
        table_name,
        f"A{header_row}:{table_end_col}{max(last_data_row, header_row)}",
        style="TableStyleMedium2",
    )
    return {
        "header_row": header_row,
        "first_data_row": first_data_row,
        "last_data_row": last_data_row,
        "table_cols": table_cols,
        "chart_label_col": CHART_LABEL_COL if chart_labels else 1,
    }


def _section_dataframes(
    df: pd.DataFrame, data: dict[str, Any]
) -> dict[str, tuple[list[str], pd.DataFrame, list[str], str | None]]:
    return {
        "location": (
            ["Location", "Inventory Value ($)", "% of Total"],
            compute_location_summary(df),
            ["location", "inventory_value", "percent_of_total"],
            None,
        ),
        "status": (
            [
                "Inventory Status",
                "SKU-Location Count",
                "Inventory Value ($)",
                "% of Total",
            ],
            compute_status_summary(df),
            [
                "inventory_status",
                "sku_location_count",
                "inventory_value",
                "percent_of_total",
            ],
            None,
        ),
        "aging": (
            [
                "Aging Bucket",
                "SKU-Location Count",
                "Inventory Value ($)",
                "% of Total",
            ],
            compute_aging_summary(df),
            [
                "aging_bucket",
                "sku_location_count",
                "inventory_value",
                "percent_of_total",
            ],
            None,
        ),
        "top_excess": (
            [
                "Rank",
                "SKU",
                "Product Name",
                "Location",
                "Excess Quantity",
                "Excess Inventory Value ($)",
            ],
            compute_top_excess(df),
            [
                "rank",
                "sku",
                "product_name",
                "location",
                "excess_quantity",
                "excess_value",
            ],
            "chart_label",
        ),
        "abc_usage": (
            [
                "ABC Class",
                "SKU Count",
                "Annual Usage Value ($)",
                "% of Annual Usage",
            ],
            compute_abc_usage_summary(data),
            [
                "abc_class",
                "sku_count",
                "annual_usage_value",
                "percent_of_annual_usage_value",
            ],
            None,
        ),
        "replenishment": (
            [
                "Replenishment Status",
                "SKU-Location Count",
                "Recommended Order Qty",
                "Recommended Order Value ($)",
            ],
            compute_replenishment_status_summary(data),
            [
                "replenishment_status",
                "sku_location_count",
                "recommended_order_quantity",
                "recommended_order_value",
            ],
            None,
        ),
        "fill_rate": (
            [
                "Location",
                "Unit FR (%)",
                "Line FR (%)",
                "Order FR (%)",
                "Service Status",
            ],
            compute_fill_rate_by_location(data),
            [
                "location",
                "unit_fill_rate",
                "line_fill_rate",
                "order_fill_rate",
                "service_status",
            ],
            None,
        ),
        "vendor_risk": (
            [
                "Vendor Risk Class",
                "Supplier Count",
                "Open PO Value ($)",
                "Average Vendor Score",
            ],
            compute_vendor_risk_summary(data),
            [
                "vendor_risk_class",
                "supplier_count",
                "open_po_value",
                "average_vendor_score",
            ],
            None,
        ),
        "transfer_benefit": (
            [
                "Rank",
                "Source Location",
                "Destination Location",
                "SKU",
                "Transfer Quantity",
                "Net Benefit ($)",
            ],
            compute_transfer_benefit_summary(data),
            [
                "rank",
                "source_location",
                "destination_location",
                "sku",
                "transfer_quantity",
                "net_benefit",
            ],
            "chart_label",
        ),
        "action": (
            [
                "Recommended Action",
                "SKU-Location Count",
                "Inventory Value ($)",
                "Estimated Financial Impact ($)",
            ],
            compute_action_summary(df),
            [
                "recommended_action",
                "record_count",
                "inventory_value",
                "estimated_financial_impact",
            ],
            None,
        ),
    }


def _format_section_table(ws: Worksheet, key: str, meta: dict[str, int]) -> None:
    first = meta["first_data_row"]
    last = meta["last_data_row"]
    if last < first:
        return

    currency_map = {
        "location": ["B"],
        "status": ["C"],
        "aging": ["C"],
        "top_excess": ["F"],
        "abc_usage": ["C"],
        "replenishment": ["D"],
        "vendor_risk": ["C"],
        "transfer_benefit": ["F"],
        "action": ["C", "D"],
    }
    integer_map = {
        "status": ["B"],
        "aging": ["B"],
        "top_excess": ["A", "E"],
        "abc_usage": ["B"],
        "replenishment": ["B", "C"],
        "vendor_risk": ["B"],
        "transfer_benefit": ["A", "E"],
        "action": ["B"],
    }
    pct_map = {
        "location": ["C"],
        "status": ["D"],
        "aging": ["D"],
        "abc_usage": ["D"],
        "fill_rate": ["B", "C", "D"],
    }

    if key in currency_map:
        format_currency_columns(ws, currency_map[key], first, last)
    if key in integer_map:
        format_integer_columns(ws, integer_map[key], first, last)
    if key in pct_map:
        format_percentage_columns(ws, pct_map[key], first, last)
    if key == "vendor_risk":
        for row in range(first, last + 1):
            ws.cell(row=row, column=4).number_format = "0.0"


def _category_count(meta: dict[str, int]) -> int:
    if meta["last_data_row"] < meta["first_data_row"]:
        return 0
    return meta["last_data_row"] - meta["first_data_row"] + 1


def _add_section_chart(
    ws: Worksheet, section: DashboardSectionLayout, meta: dict[str, int]
) -> None:
    if meta["last_data_row"] < meta["first_data_row"]:
        return

    header_row = meta["header_row"]
    first_row = meta["first_data_row"]
    last_row = meta["last_data_row"]
    anchor = f"{CHART_ANCHOR_COL}{header_row}"
    w = CHART_WIDTH_CM
    h = chart_height_cm(section.end_row, header_row)

    if section.key in {"top_excess", "transfer_benefit"}:
        cats = make_category_reference(ws, meta["chart_label_col"], first_row, last_row)
    else:
        cats = make_category_reference(ws, 1, first_row, last_row)

    if section.key == "location":
        data = make_data_reference(ws, 2, first_row, last_row)
        add_bar_chart(
            ws,
            section.chart_title,
            data,
            cats,
            anchor=anchor,
            width=w,
            height=h,
            horizontal=True,
            value_axis_title="Inventory Value ($)",
            category_axis_title="Location",
            reverse_category_order=True,
            show_data_labels=True,
            hide_legend=True,
        )
    elif section.key == "status":
        data = make_data_reference(ws, 3, first_row, last_row)
        add_bar_chart(
            ws,
            section.chart_title,
            data,
            cats,
            anchor=anchor,
            width=w,
            height=h,
            horizontal=True,
            value_axis_title="Inventory Value ($)",
            category_axis_title="Inventory Status",
            reverse_category_order=True,
            show_data_labels=True,
            hide_legend=True,
        )
    elif section.key == "aging":
        data = make_data_reference(ws, 3, first_row, last_row)
        add_bar_chart(
            ws,
            section.chart_title,
            data,
            cats,
            anchor=anchor,
            width=w,
            height=h,
            horizontal=False,
            value_axis_title="Inventory Value ($)",
            category_axis_title="Aging Bucket",
            reverse_category_order=False,
            show_data_labels=True,
            hide_legend=True,
        )
    elif section.key == "top_excess":
        data = make_data_reference(ws, 6, first_row, last_row)
        add_bar_chart(
            ws,
            section.chart_title,
            data,
            cats,
            anchor=anchor,
            width=w,
            height=h,
            horizontal=True,
            value_axis_title="Excess Inventory Value ($)",
            category_axis_title="Inventory Item",
            reverse_category_order=True,
            show_data_labels=True,
            hide_legend=True,
        )
    elif section.key == "abc_usage":
        data = make_data_reference(ws, 3, first_row, last_row)
        add_bar_chart(
            ws,
            section.chart_title,
            data,
            cats,
            anchor=anchor,
            width=w,
            height=h,
            horizontal=False,
            value_axis_title="Annual Usage Value ($)",
            category_axis_title="ABC Class",
            reverse_category_order=False,
            show_data_labels=True,
            hide_legend=True,
        )
    elif section.key == "replenishment":
        data = make_data_reference(ws, 4, first_row, last_row)
        add_bar_chart(
            ws,
            section.chart_title,
            data,
            cats,
            anchor=anchor,
            width=w,
            height=h,
            horizontal=True,
            value_axis_title="Recommended Order Value ($)",
            category_axis_title="Replenishment Status",
            reverse_category_order=True,
            show_data_labels=True,
            hide_legend=True,
        )
    elif section.key == "fill_rate":
        data = make_data_reference(ws, 2, header_row, last_row, max_col=4)
        add_clustered_bar_chart(
            ws,
            section.chart_title,
            data,
            cats,
            anchor=anchor,
            width=w,
            height=h,
            horizontal=True,
            value_axis_title="Fill Rate (%)",
            category_axis_title=None,
            value_axis_min=0,
            value_axis_max=1.05,
            reverse_category_order=True,
            show_data_labels=True,
            hide_legend=True,
            compact_labels=True,
            label_number_format="0%",
        )
    elif section.key == "vendor_risk":
        data = make_data_reference(ws, 3, first_row, last_row)
        add_bar_chart(
            ws,
            section.chart_title,
            data,
            cats,
            anchor=anchor,
            width=w,
            height=h,
            horizontal=True,
            value_axis_title="Open PO Value ($)",
            category_axis_title="Vendor Risk Class",
            reverse_category_order=True,
            show_data_labels=True,
            hide_legend=True,
        )
    elif section.key == "transfer_benefit":
        data = make_data_reference(ws, 6, first_row, last_row)
        add_bar_chart(
            ws,
            section.chart_title,
            data,
            cats,
            anchor=anchor,
            width=w,
            height=h,
            horizontal=True,
            value_axis_title="Net Benefit ($)",
            category_axis_title="Transfer Opportunity",
            reverse_category_order=True,
            show_data_labels=True,
            hide_legend=True,
        )
    elif section.key == "action":
        data = make_data_reference(ws, 3, first_row, last_row)
        add_bar_chart(
            ws,
            section.chart_title,
            data,
            cats,
            anchor=anchor,
            width=w,
            height=h,
            horizontal=True,
            value_axis_title="Inventory Value ($)",
            category_axis_title="Recommended Action",
            reverse_category_order=True,
            show_data_labels=True,
            hide_legend=True,
        )


def _apply_view_settings(ws: Worksheet, last_row: int) -> None:
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    ws.freeze_panes = FREEZE_PANE
    ws.sheet_view.selection[0].activeCell = "A1"
    ws.sheet_view.selection[0].sqref = "A1"
    ws.print_area = f"A1:V{last_row}"


def build(ws: Worksheet, context: dict[str, Any]) -> None:
    """Build the Inventory Dashboard from generated workbook data."""
    data = context["data"]
    df: pd.DataFrame = data.get("inventory", pd.DataFrame())
    end_row = _data_end_row(len(df))

    context["_dashboard_charts_before_cleanup"] = _clear_dashboard_charts(ws)
    _apply_canvas_columns(ws)
    _write_header_and_kpis(ws, end_row, data)

    datasets = _section_dataframes(df, data)
    data_row_counts = {
        key: max(1, len(summary_df)) for key, (_, summary_df, _, _) in datasets.items()
    }
    section_layouts = compute_section_layouts(data_row_counts)
    section_meta: dict[str, dict[str, int]] = {}

    for idx, section in enumerate(section_layouts, start=1):
        headers, summary_df, cols, label_col = datasets[section.key]
        rows = summary_df[cols].values.tolist() if not summary_df.empty else []
        chart_labels = None
        if label_col and label_col in summary_df.columns and not summary_df.empty:
            chart_labels = summary_df[label_col].astype(str).tolist()
        meta = _write_section_table(
            ws,
            section,
            headers,
            rows,
            f"{DASHBOARD_TABLE_PREFIX}{idx}Table",
            chart_labels=chart_labels,
        )
        section_meta[section.key] = meta
        _format_section_table(ws, section.key, meta)
        _add_section_chart(ws, section, meta)

    last_row = section_layouts[-1].end_row
    _apply_view_settings(ws, last_row)
    set_landscape_print(ws, fit_width=1, repeat_header_rows="1:2")
    context["_dashboard_section_meta"] = section_meta
    context["_dashboard_section_layouts"] = section_layouts
