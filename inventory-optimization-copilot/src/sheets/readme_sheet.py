"""README sheet — toolkit guide for interviewers and reviewers."""

from __future__ import annotations

from typing import Any

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from config import style_config as sc
from config.workbook_config import AUTHOR, VERSION, WORKBOOK_TITLE
from src.workbook.styles import apply_section_header_style
from src.workbook.utils import set_column_widths, set_portrait_print

CONTENT_START_COL = 1
CONTENT_END_COL = 4
SPAN_COLS = CONTENT_END_COL - CONTENT_START_COL + 1


def _merge_and_write(
    ws: Worksheet,
    row: int,
    text: str,
    *,
    font: Font | None = None,
    fill: PatternFill | None = None,
    alignment: Alignment | None = None,
    row_height: float | None = None,
) -> None:
    ws.merge_cells(
        start_row=row,
        start_column=CONTENT_START_COL,
        end_row=row,
        end_column=CONTENT_END_COL,
    )
    cell = ws.cell(row=row, column=CONTENT_START_COL, value=text)
    cell.font = font or sc.BODY_FONT
    cell.alignment = alignment or Alignment(
        horizontal="left", vertical="top", wrap_text=True
    )
    if fill:
        cell.fill = fill
    if row_height:
        ws.row_dimensions[row].height = row_height


def _apply_title_banner(ws: Worksheet, row: int, title: str) -> None:
    ws.merge_cells(
        start_row=row,
        start_column=CONTENT_START_COL,
        end_row=row,
        end_column=CONTENT_END_COL,
    )
    cell = ws.cell(row=row, column=CONTENT_START_COL, value=title)
    cell.font = Font(
        name=sc.FONTS["default_name"], bold=True, color=sc.COLORS["white"], size=18
    )
    cell.fill = PatternFill(
        start_color=sc.COLORS["navy"], end_color=sc.COLORS["navy"], fill_type="solid"
    )
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 36


def _apply_subtitle_bar(ws: Worksheet, row: int, text: str) -> None:
    ws.merge_cells(
        start_row=row,
        start_column=CONTENT_START_COL,
        end_row=row,
        end_column=CONTENT_END_COL,
    )
    cell = ws.cell(row=row, column=CONTENT_START_COL, value=text)
    cell.font = Font(
        name=sc.FONTS["default_name"],
        bold=True,
        color=sc.COLORS["dark_blue"],
        size=sc.FONTS["body_size"],
    )
    cell.fill = sc.KPI_CARD_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 22


def _apply_disclaimer(ws: Worksheet, row: int, text: str) -> None:
    ws.merge_cells(
        start_row=row,
        start_column=CONTENT_START_COL,
        end_row=row,
        end_column=CONTENT_END_COL,
    )
    cell = ws.cell(row=row, column=CONTENT_START_COL, value=text)
    cell.font = Font(
        name=sc.FONTS["default_name"],
        italic=True,
        color=sc.COLORS["dark_blue"],
        size=sc.FONTS["body_size"],
    )
    cell.fill = PatternFill(
        start_color=sc.COLORS["gray"], end_color=sc.COLORS["gray"], fill_type="solid"
    )
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 28


def _write_section(ws: Worksheet, row: int, title: str) -> int:
    apply_section_header_style(ws, row, CONTENT_START_COL, title, span_cols=SPAN_COLS)
    ws.row_dimensions[row].height = 22
    return row + 1


def _write_body(ws: Worksheet, row: int, text: str, row_height: float = 48) -> int:
    _merge_and_write(ws, row, text, row_height=row_height)
    return row + 1


def _write_spacer(ws: Worksheet, row: int, height: float = 8) -> int:
    ws.row_dimensions[row].height = height
    return row + 1


def build(ws: Worksheet, context: dict[str, Any]) -> None:
    """Populate the README tab with inventory optimization overview and demo guidance."""
    row = 1
    _apply_title_banner(ws, row, WORKBOOK_TITLE)
    row += 1
    _apply_subtitle_bar(ws, row, f"Version: {VERSION}    |    Author: {AUTHOR}")
    row += 1
    _apply_disclaimer(
        ws,
        row,
        "Disclaimer: All data in this workbook is fictional sample data created for "
        "demonstration and interview purposes.",
    )
    row += 1
    row = _write_spacer(ws, row)

    row = _write_section(ws, row, "Purpose of This Workbook")
    row = _write_body(
        ws,
        row,
        "Inventory Optimization Copilot is a decision-support workbook for supply chain, "
        "procurement, and distribution teams. It highlights aged and excess inventory, "
        "transfer opportunities, markdown priorities, and executive KPIs — generated "
        "entirely from reproducible Python source code.",
        row_height=56,
    )
    row = _write_spacer(ws, row)

    row = _write_section(ws, row, "Recommended Demo Path")
    row = _write_body(
        ws,
        row,
        "Tab order:\n"
        "1. Inventory Dashboard\n"
        "2. Inventory Classification\n"
        "3. Aged Excess Analysis\n"
        "4. Cycle Count Plan\n"
        "5. Replenishment Planning\n"
        "6. Demand Forecast\n"
        "7. Service Level Analysis\n"
        "8. Transfer Planner\n"
        "9. Markdown Planner\n"
        "10. Management Summary\n\n"
        "Suggested narrative:\n"
        '"This workbook focuses on inventory optimization. The dashboard shows total value, '
        "aged stock, excess exposure, and recovery opportunity. I drill into exceptions, "
        'evaluate transfer-before-markdown economics, and close with a leadership-ready summary."',
        row_height=120,
    )
    row = _write_spacer(ws, row)

    row = _write_section(ws, row, "Key Capabilities")
    row = _write_body(
        ws,
        row,
        "• Inventory health dashboards with KPI cards and charts\n"
        "• ABC classification, turnover, DOH, and inventory accuracy analytics\n"
        "• Cycle count planning with risk-based prioritization\n"
        "• Replenishment planning with safety stock, ROP, and EOQ\n"
        "• Statistical demand forecasting and service level analysis\n"
        "• Aged, excess, slow-moving, and obsolete stock analysis\n"
        "• Transfer planning between locations to balance stock\n"
        "• Markdown modeling to protect margin on end-of-life inventory\n"
        "• Printable management summary for leadership review",
        row_height=100,
    )
    row = _write_spacer(ws, row)

    row = _write_section(ws, row, "Workbook Navigation")
    row = _write_body(
        ws,
        row,
        "Use the sheet tabs at the bottom to move between modules. Operational tabs include "
        "filterable headers — click dropdown arrows to explore by status, location, or category.",
        row_height=48,
    )

    set_column_widths(ws, {"A": 4, "B": 28, "C": 28, "D": 28})
    ws.sheet_view.showGridLines = False
    set_portrait_print(ws, fit_width=1, fit_height=0, repeat_header_rows="1:3")
