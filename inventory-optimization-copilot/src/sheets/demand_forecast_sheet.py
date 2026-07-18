"""Demand Forecast sheet — method comparison, selection, and horizon projections."""

from __future__ import annotations

from typing import Any

import pandas as pd
from openpyxl.chart import Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from config import style_config as sc
from config.workbook_config import FORECAST_DEMO_CHART_SKU_COUNT
from src.services.forecast_service import (
    DEMAND_FORECAST_HEADERS,
    _weekly_series,
    build_demand_forecast_dataframe,
    select_demo_chart_skus,
)
from src.workbook.charts import add_line_chart, make_category_reference
from src.workbook.styles import (
    apply_kpi_card_style,
    apply_risk_conditional_formatting,
    apply_section_header_style,
    apply_table_header_style,
    freeze_panes,
)
from src.workbook.utils import (
    autosize_columns,
    create_excel_table,
    format_integer_columns,
    format_percentage_columns,
    set_landscape_print,
)

HEADERS = DEMAND_FORECAST_HEADERS

TITLE_ROW = 1
KPI_SECTION_ROW = 2
KPI_TITLE_ROW = 3
KPI_VALUE_ROW = 4
TABLE_SECTION_ROW = 6
HEADER_ROW = 7
DATA_START_ROW = 8
COL_COUNT = len(HEADERS)
TABLE_NAME = "DemandForecastTable"

STATUS_CF_MAP = {
    "Approved": "healthy",
    "Limited History": "watch",
    "High Error": "slow_moving",
    "No Recent Demand": "neutral",
}


def _col(header: str) -> str:
    return str(get_column_letter(HEADERS.index(header) + 1))


def _write_title(ws: Worksheet, count: int) -> None:
    ws.merge_cells(
        start_row=TITLE_ROW, start_column=1, end_row=TITLE_ROW, end_column=10
    )
    cell = ws.cell(
        row=TITLE_ROW,
        column=1,
        value=f"Demand Forecast — {count:,} SKU-Locations",
    )
    cell.font = Font(
        name=sc.FONTS["default_name"],
        bold=True,
        color=sc.COLORS["white"],
        size=sc.FONTS["title_size"],
    )
    cell.fill = sc.HEADER_FILL
    ws.row_dimensions[TITLE_ROW].height = 28


def _write_kpi_cards(ws: Worksheet, df: pd.DataFrame) -> None:
    apply_section_header_style(ws, KPI_SECTION_ROW, 1, "Forecast KPIs", span_cols=10)
    if df.empty:
        return
    kpis = [
        ("SKU-Locations", len(df), "integer"),
        ("Avg Selected WAPE", round(float(df["Selected WAPE"].mean()), 4), "decimal"),
        (
            "Approved Reviews",
            int((df["Review Status"] == "Approved").sum()),
            "integer",
        ),
        (
            "Most Selected Method",
            (
                str(df["Selected Method"].mode().iloc[0])
                if not df["Selected Method"].mode().empty
                else "N/A"
            ),
            "text",
        ),
    ]
    card_cols = [1, 3, 5, 7]
    for idx, (title, value, fmt) in enumerate(kpis):
        col = card_cols[idx]
        apply_kpi_card_style(
            ws, KPI_TITLE_ROW, col, KPI_VALUE_ROW, col, title, value, span_cols=2
        )
        cell = ws.cell(row=KPI_VALUE_ROW, column=col)
        if fmt == "integer":
            cell.number_format = sc.NUMBER_FORMATS["integer"]


def _write_headers(ws: Worksheet) -> None:
    apply_section_header_style(
        ws, TABLE_SECTION_ROW, 1, "SKU-Location Forecast Comparison", span_cols=10
    )
    for col_idx, header in enumerate(HEADERS, start=1):
        ws.cell(row=HEADER_ROW, column=col_idx, value=header)
    apply_table_header_style(ws, HEADER_ROW, COL_COUNT)


def _write_data(ws: Worksheet, df: pd.DataFrame) -> int:
    if df.empty:
        return HEADER_ROW
    for offset, (_, row) in enumerate(df.iterrows()):
        excel_row = DATA_START_ROW + offset
        for col_idx, header in enumerate(HEADERS, start=1):
            val = row[header]
            if pd.isna(val):
                val = None
            ws.cell(row=excel_row, column=col_idx, value=val)
    return DATA_START_ROW + len(df) - 1


def _apply_formats(ws: Worksheet, last_row: int) -> None:
    if last_row < DATA_START_ROW:
        return
    integer_headers = [
        "History Length (Weeks)",
        "Forecast 4-Week",
        "Forecast 8-Week",
        "Forecast 12-Week",
        "Stockout Constrained History",
    ]
    pct_headers = [
        "Naive WAPE",
        "4-Week MA WAPE",
        "WMA WAPE",
        "SES WAPE",
        "Selected WAPE",
        "Selected MAPE",
    ]
    format_integer_columns(
        ws, [_col(h) for h in integer_headers], DATA_START_ROW, last_row
    )
    format_percentage_columns(
        ws, [_col(h) for h in pct_headers if h in HEADERS], DATA_START_ROW, last_row
    )


def _add_demo_charts(
    ws: Worksheet, df: pd.DataFrame, data: dict[str, Any], last_row: int
) -> None:
    if df.empty or last_row < DATA_START_ROW:
        return

    demand = data.get("demand_history", pd.DataFrame())
    demo_skus = select_demo_chart_skus(df, FORECAST_DEMO_CHART_SKU_COUNT)
    start = last_row + 3

    for sku in demo_skus:
        sku_rows = df[df["SKU"] == sku].head(1)
        if sku_rows.empty:
            continue
        loc_id = str(sku_rows.iloc[0]["Location ID"])
        weekly = _weekly_series(demand, sku, loc_id)
        if not weekly:
            continue

        chart_start = start
        ws.cell(row=chart_start, column=1, value="Week")
        ws.cell(row=chart_start, column=2, value="Actual")
        for i, value in enumerate(weekly):
            ws.cell(row=chart_start + 1 + i, column=1, value=i + 1)
            ws.cell(row=chart_start + 1 + i, column=2, value=float(value))

        cats = make_category_reference(
            ws, 1, chart_start + 1, chart_start + len(weekly)
        )
        series_data = Reference(
            ws,
            min_col=2,
            min_row=chart_start,
            max_col=2,
            max_row=chart_start + len(weekly),
        )
        chart = add_line_chart(
            ws,
            f"Demand History — {sku}",
            series_data,
            cats,
            anchor=f"AU{max(3, start - last_row + 3)}",
            width=14,
            height=8,
        )
        chart.series[0].title = SeriesLabel(v=f"{sku} weekly demand")
        start += len(weekly) + 6


def build(ws: Worksheet, context: dict[str, Any]) -> None:
    """Build Demand Forecast sheet."""
    data = context["data"]
    df = build_demand_forecast_dataframe(data)
    _write_title(ws, len(df))
    _write_kpi_cards(ws, df)
    _write_headers(ws)
    last_row = _write_data(ws, df)
    end_col = get_column_letter(COL_COUNT)
    if last_row >= DATA_START_ROW:
        create_excel_table(ws, TABLE_NAME, f"A{HEADER_ROW}:{end_col}{last_row}")
        status_col = _col("Review Status")
        apply_risk_conditional_formatting(
            ws,
            f"{status_col}{DATA_START_ROW}:{status_col}{last_row}",
            status_col,
            DATA_START_ROW,
            STATUS_CF_MAP,
        )
    _apply_formats(ws, last_row)
    _add_demo_charts(ws, df, data, last_row)
    freeze_panes(ws, row=DATA_START_ROW, col=1)
    autosize_columns(ws, min_width=10, max_width=36)
    ws.sheet_view.showGridLines = False
    set_landscape_print(
        ws, fit_width=1, repeat_header_rows=f"{HEADER_ROW}:{HEADER_ROW}"
    )
