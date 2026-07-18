"""Management Summary sheet — printable executive overview (inventory only)."""

from __future__ import annotations

from typing import Any

from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.worksheet import Worksheet

from config import style_config as sc
from config.workbook_config import AS_OF_DATE, VERSION
from src.services.summary_service import build_management_summary_context
from src.workbook.styles import (
    apply_kpi_card_style,
    apply_section_header_style,
    apply_table_header_range,
)
from src.workbook.utils import (
    add_internal_sheet_link,
    format_currency_columns,
    format_integer_columns,
    format_percentage_columns,
    set_column_widths,
    set_print_layout,
)

INVENTORY_RISK_HEADERS = ["SKU", "Location", "Issue", "Value"]

TITLE_ROW = 1
SUBTITLE_ROW = 2
SECTION1_ROW = 4
KPI1_TITLE_ROW = 5
KPI1_VALUE_ROW = 6
SECTION2_ROW = 8
KPI2_TITLE_ROW = 9
KPI2_VALUE_ROW = 10
SECTION3_ROW = 12
KPI3_TITLE_ROW = 13
KPI3_VALUE_ROW = 14
SECTION4_ROW = 16
KPI4_TITLE_ROW = 17
KPI4_VALUE_ROW = 18
SECTION5_ROW = 20
KPI5_TITLE_ROW = 21
KPI5_VALUE_ROW = 22
PAGE_BREAK_ROW = 24
SECTION6_ROW = 25
SUMMARY_HEADER_ROW = 26
SUMMARY_DATA_START_ROW = 27
SECTION7_ROW = 34
RISK_HEADER_ROW = 35
RISK_DATA_START_ROW = 36
SECTION8_ROW = 42
PLAN_HEADER_ROW = 43
PLAN_DATA_START_ROW = 44

ACTION_PLAN = {
    "30 Days": [
        "Validate inventory records and reconcile cycle count exceptions",
        "Review stockout-risk and below-ROP SKU-locations for expediting",
        "Confirm transfer-before-markdown priorities on excess inventory",
    ],
    "60 Days": [
        "Execute approved network transfers and monitor source protection levels",
        "Apply markdown tiers with margin guardrails on aged inventory",
        "Replenish critical locations and tighten vendor follow-up on late POs",
    ],
    "90 Days": [
        "Measure inventory turnover, fill rate, and markdown recovery impact",
        "Review forecast WAPE/bias and recalibrate safety stock policies",
        "Present process improvements and KPI targets to leadership",
    ],
}


def _write_title_banner(ws: Worksheet) -> None:
    ws.merge_cells(start_row=TITLE_ROW, start_column=1, end_row=TITLE_ROW, end_column=8)
    cell = ws.cell(row=TITLE_ROW, column=1, value="Management Summary")
    cell.font = Font(
        name=sc.FONTS["default_name"],
        bold=True,
        color=sc.COLORS["white"],
        size=sc.FONTS["title_size"],
    )
    cell.fill = sc.HEADER_FILL
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[TITLE_ROW].height = 30

    ws.merge_cells(
        start_row=SUBTITLE_ROW, start_column=1, end_row=SUBTITLE_ROW, end_column=8
    )
    subtitle = ws.cell(
        row=SUBTITLE_ROW,
        column=1,
        value=(
            f"{VERSION} | Inventory Executive Overview | "
            f"As-of Date: {AS_OF_DATE.strftime('%B %d, %Y')}"
        ),
    )
    subtitle.font = Font(
        name=sc.FONTS["default_name"],
        italic=True,
        color=sc.COLORS["dark_blue"],
        size=sc.FONTS["body_size"],
    )
    subtitle.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[SUBTITLE_ROW].height = 18
    add_internal_sheet_link(ws, TITLE_ROW, 10, "README", "← README")


def _write_kpi_row(
    ws: Worksheet,
    title_row: int,
    value_row: int,
    cards: list[tuple[str, Any, str]],
) -> None:
    card_cols = [1, 3, 5, 7, 9]
    format_map = {
        "currency": sc.NUMBER_FORMATS["currency_compact"],
        "integer": sc.NUMBER_FORMATS["integer"],
        "percentage": sc.NUMBER_FORMATS["percentage"],
        "decimal": "0.00",
    }
    for idx, (title, value, fmt_key) in enumerate(cards[:5]):
        col = card_cols[idx]
        apply_kpi_card_style(
            ws, title_row, col, value_row, col, title, value, span_cols=2
        )
        ws.cell(row=value_row, column=col).number_format = format_map[fmt_key]


def _write_metric_table(
    ws: Worksheet,
    header_row: int,
    data_start: int,
    rows: list[tuple[str, Any]],
) -> None:
    apply_table_header_range(ws, header_row, 1, 2)
    ws.cell(row=header_row, column=1, value="Metric")
    ws.cell(row=header_row, column=2, value="Value")
    for offset, (label, value) in enumerate(rows[:8]):
        excel_row = data_start + offset
        ws.cell(row=excel_row, column=1, value=label)
        ws.cell(row=excel_row, column=2, value=value)


def _write_risk_table(ws: Worksheet, inventory_risks: list[dict[str, Any]]) -> None:
    apply_section_header_style(ws, SECTION7_ROW, 1, "Top Inventory Risks", span_cols=8)
    apply_table_header_range(ws, RISK_HEADER_ROW, 1, 4)
    for col_idx, header in enumerate(INVENTORY_RISK_HEADERS, start=1):
        ws.cell(row=RISK_HEADER_ROW, column=col_idx, value=header)
    for offset in range(5):
        excel_row = RISK_DATA_START_ROW + offset
        if offset < len(inventory_risks):
            risk = inventory_risks[offset]
            ws.cell(row=excel_row, column=1, value=risk["SKU"])
            ws.cell(row=excel_row, column=2, value=risk["Location"])
            ws.cell(row=excel_row, column=3, value=risk["Issue"])
            value_cell = ws.cell(row=excel_row, column=4, value=risk["Value"])
            value_cell.number_format = sc.NUMBER_FORMATS["currency_compact"]


def _write_action_plan(ws: Worksheet) -> None:
    apply_section_header_style(
        ws, SECTION8_ROW, 1, "30 / 60 / 90 Day Action Plan", span_cols=8
    )
    plan_cols = [1, 4, 7]
    for idx, horizon in enumerate(["30 Days", "60 Days", "90 Days"]):
        col = plan_cols[idx]
        ws.merge_cells(
            start_row=PLAN_HEADER_ROW,
            start_column=col,
            end_row=PLAN_HEADER_ROW,
            end_column=col + 2,
        )
        header = ws.cell(row=PLAN_HEADER_ROW, column=col, value=horizon)
        header.font = sc.KPI_TITLE_FONT
        header.fill = sc.KPI_CARD_FILL
        header.alignment = sc.KPI_CARD_STYLE["alignment"]
        header.border = sc.NAVY_BORDER
        for bullet_idx, bullet in enumerate(ACTION_PLAN[horizon]):
            row = PLAN_DATA_START_ROW + bullet_idx
            ws.merge_cells(
                start_row=row, start_column=col, end_row=row, end_column=col + 2
            )
            cell = ws.cell(row=row, column=col, value=f"• {bullet}")
            cell.font = sc.BODY_FONT
            cell.alignment = Alignment(
                horizontal="left", vertical="top", wrap_text=True
            )
            cell.border = sc.THIN_GRAY_BORDER
            ws.row_dimensions[row].height = 30


def build(ws: Worksheet, context: dict[str, Any]) -> None:
    """Build the printable Management Summary (one to two pages)."""
    metrics = build_management_summary_context(context["data"])
    class_kpis = metrics["class_kpis"]
    planning = metrics["planning"]
    procurement = metrics["procurement"]

    _write_title_banner(ws)

    apply_section_header_style(ws, SECTION1_ROW, 1, "Inventory Health", span_cols=8)
    _write_kpi_row(
        ws,
        KPI1_TITLE_ROW,
        KPI1_VALUE_ROW,
        [
            ("Total Inventory Value", metrics["total_value"], "currency"),
            ("Aged Inventory Value", metrics["aged_value"], "currency"),
            ("Excess SKU-Locations", metrics["excess_count"], "integer"),
            ("Stockout-Risk SKUs", metrics["stockout_count"], "integer"),
            ("Est. Recovery Value", metrics["recovery_value"], "currency"),
        ],
    )

    apply_section_header_style(
        ws, SECTION2_ROW, 1, "Classification & Accuracy", span_cols=8
    )
    _write_kpi_row(
        ws,
        KPI2_TITLE_ROW,
        KPI2_VALUE_ROW,
        [
            (
                "Class A / B / C",
                f"{class_kpis['abc_a_count']}/{class_kpis['abc_b_count']}/{class_kpis['abc_c_count']}",
                "integer",
            ),
            ("Enterprise Turnover", class_kpis["enterprise_turnover"], "decimal"),
            ("Unit Accuracy %", class_kpis["avg_unit_accuracy_pct"], "percentage"),
            ("Overdue Cycle Counts", class_kpis["overdue_count"], "integer"),
            ("Avg Financial DOH", class_kpis["avg_financial_doh"], "decimal"),
        ],
    )

    apply_section_header_style(
        ws, SECTION3_ROW, 1, "Planning & Replenishment", span_cols=8
    )
    _write_kpi_row(
        ws,
        KPI3_TITLE_ROW,
        KPI3_VALUE_ROW,
        [
            ("Below Reorder Point", planning["below_reorder_point"], "integer"),
            ("Order Required Lines", metrics["order_required"], "integer"),
            (
                "Recommended Order Value",
                planning["recommended_order_value"],
                "currency",
            ),
            ("Forecast WAPE", planning["forecast_wape"], "percentage"),
            ("Forecast Reviews", metrics["forecast_reviews"], "integer"),
        ],
    )

    apply_section_header_style(
        ws, SECTION4_ROW, 1, "Service & Procurement", span_cols=8
    )
    _write_kpi_row(
        ws,
        KPI4_TITLE_ROW,
        KPI4_VALUE_ROW,
        [
            ("Unit Fill Rate", planning["unit_fill_rate"], "percentage"),
            ("Critical Service Gaps", metrics["critical_service"], "integer"),
            ("Open PO Lines", metrics["open_po_lines"], "integer"),
            ("Late PO Count", procurement["late_po_count"], "integer"),
            ("Vendor OTIF", procurement["vendor_otif"], "percentage"),
        ],
    )

    apply_section_header_style(ws, SECTION5_ROW, 1, "Transfer & Markdown", span_cols=8)
    _write_kpi_row(
        ws,
        KPI5_TITLE_ROW,
        KPI5_VALUE_ROW,
        [
            ("Transfer Actions", metrics["transfer_actions"], "integer"),
            ("Transfer Net Benefit", procurement["transfer_net_benefit"], "currency"),
            ("Remaining Shortage", procurement["remaining_shortage"], "integer"),
            ("Markdown Candidates", metrics["markdown_candidates"], "integer"),
            ("High-Risk Vendors", procurement["high_risk_vendors"], "integer"),
        ],
    )

    ws.row_breaks.append(Break(id=PAGE_BREAK_ROW))

    apply_section_header_style(
        ws, SECTION6_ROW, 1, "Recommended Actions & Priorities", span_cols=8
    )
    _write_metric_table(
        ws,
        SUMMARY_HEADER_ROW,
        SUMMARY_DATA_START_ROW,
        metrics["action_summary"],
    )
    _write_risk_table(ws, metrics["top_risks"])
    _write_action_plan(ws)

    format_integer_columns(
        ws, ["B"], SUMMARY_DATA_START_ROW, SUMMARY_DATA_START_ROW + 7
    )
    format_currency_columns(ws, ["D"], RISK_DATA_START_ROW, RISK_DATA_START_ROW + 4)
    format_percentage_columns(ws, ["B"], KPI2_VALUE_ROW, KPI5_VALUE_ROW)
    set_column_widths(
        ws, {"A": 18, "B": 16, "C": 22, "D": 14, "E": 2, "F": 14, "G": 16, "H": 18}
    )
    ws.sheet_view.showGridLines = False
    set_print_layout(
        ws,
        orientation="landscape",
        fit_width=1,
        fit_height=0,
        repeat_header_rows=f"{TITLE_ROW}:{SUBTITLE_ROW}",
    )
