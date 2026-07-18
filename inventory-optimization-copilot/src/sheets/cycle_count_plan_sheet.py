"""Cycle Count Plan sheet — prioritized count schedule and risk scores."""

from __future__ import annotations

from typing import Any

import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from config import style_config as sc
from src.services.classification_service import (
    CYCLE_COUNT_PLAN_HEADERS,
    build_cycle_count_plan_dataframe,
)
from src.workbook.styles import (
    apply_risk_conditional_formatting,
    apply_table_header_style,
    freeze_panes,
)
from src.workbook.utils import (
    autosize_columns,
    create_excel_table,
    format_currency_columns,
    format_date_columns,
    format_integer_columns,
    format_percentage_columns,
    set_landscape_print,
)

HEADERS = CYCLE_COUNT_PLAN_HEADERS

TITLE_ROW = 1
HEADER_ROW = 2
DATA_START_ROW = 3
COL_COUNT = len(HEADERS)
TABLE_NAME = "CycleCountPlanTable"

COL_INTEGER = ["A", "G", "L", "N", "P"]
COL_CURRENCY = ["H", "I"]
COL_PERCENTAGE = ["O"]
COL_DATES = ["J", "K"]
COL_STATUS = "Q"
COL_RISK = "P"

STATUS_CF_MAP = {
    "Overdue": "critical",
    "Due": "slow_moving",
    "Due Soon": "watch",
    "Scheduled": "healthy",
    "Completed": "healthy",
    "Recount Required": "obsolete",
}


def _write_title(ws: Worksheet, count: int) -> None:
    ws.merge_cells(start_row=TITLE_ROW, start_column=1, end_row=TITLE_ROW, end_column=8)
    cell = ws.cell(
        row=TITLE_ROW,
        column=1,
        value=f"Cycle Count Plan — {count:,} SKU-Locations",
    )
    cell.font = Font(
        name=sc.FONTS["default_name"],
        bold=True,
        color=sc.COLORS["white"],
        size=sc.FONTS["title_size"],
    )
    cell.fill = sc.HEADER_FILL
    ws.row_dimensions[TITLE_ROW].height = 28


def _write_headers(ws: Worksheet) -> None:
    for col_idx, header in enumerate(HEADERS, start=1):
        ws.cell(row=HEADER_ROW, column=col_idx, value=header)
    apply_table_header_style(ws, HEADER_ROW, COL_COUNT)


def _write_data(ws: Worksheet, df: pd.DataFrame) -> int:
    if df.empty:
        return HEADER_ROW
    for offset, (_, row) in enumerate(df.iterrows()):
        excel_row = DATA_START_ROW + offset
        for col_idx, header in enumerate(HEADERS, start=1):
            val = row[header]
            if pd.isna(val):
                val = None
            ws.cell(row=excel_row, column=col_idx, value=val)
    return DATA_START_ROW + len(df) - 1


def _apply_formats(ws: Worksheet, last_row: int) -> None:
    if last_row < DATA_START_ROW:
        return
    format_integer_columns(ws, COL_INTEGER, DATA_START_ROW, last_row)
    format_currency_columns(ws, COL_CURRENCY, DATA_START_ROW, last_row)
    format_percentage_columns(ws, COL_PERCENTAGE, DATA_START_ROW, last_row)
    format_date_columns(ws, COL_DATES, DATA_START_ROW, last_row)


def build(ws: Worksheet, context: dict[str, Any]) -> None:
    """Build Cycle Count Plan sheet."""
    df = build_cycle_count_plan_dataframe(context["data"])
    _write_title(ws, len(df))
    _write_headers(ws)
    last_row = _write_data(ws, df)
    end_col = get_column_letter(COL_COUNT)
    if last_row >= DATA_START_ROW:
        create_excel_table(ws, TABLE_NAME, f"A{HEADER_ROW}:{end_col}{last_row}")
        status_range = f"{COL_STATUS}{DATA_START_ROW}:{COL_STATUS}{last_row}"
        apply_risk_conditional_formatting(
            ws, status_range, COL_STATUS, DATA_START_ROW, STATUS_CF_MAP
        )
    _apply_formats(ws, last_row)
    freeze_panes(ws, row=DATA_START_ROW, col=1)
    autosize_columns(ws, min_width=10, max_width=32)
    ws.sheet_view.showGridLines = False
    set_landscape_print(
        ws, fit_width=1, repeat_header_rows=f"{HEADER_ROW}:{HEADER_ROW}"
    )
