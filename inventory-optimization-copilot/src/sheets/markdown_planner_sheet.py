"""Markdown Planner sheet — markdown modeling for aged and excess inventory."""

from __future__ import annotations

from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from config import style_config as sc
from config.workbook_config import MARKDOWN_THRESHOLDS
from src.workbook.styles import (
    apply_risk_conditional_formatting,
    apply_table_header_style,
    freeze_panes,
)
from src.workbook.utils import (
    autosize_columns,
    create_excel_table,
    format_currency_columns,
    format_integer_columns,
    format_percentage_columns,
    set_landscape_print,
)

HEADERS = [
    "SKU",
    "Product Name",
    "Location",
    "Quantity On Hand",
    "Unit Cost",
    "Current Selling Price",
    "Current Margin %",
    "Age Days",
    "Demand Trend",
    "Suggested Markdown %",
    "Markdown Price",
    "Projected Sell Through %",
    "Estimated Recovery Value",
    "Margin Impact",
    "Recommended Disposition",
]

ELIGIBLE_STATUSES = ("Slow-Moving", "Excess / Aged", "Obsolete", "Excess")

TITLE_ROW = 1
HEADER_ROW = 2
DATA_START_ROW = 3
COL_COUNT = len(HEADERS)
TABLE_NAME = "MarkdownPlannerTable"

COL_INTEGER = ["D", "H"]
COL_CURRENCY = ["E", "F", "K", "M", "N"]
COL_PERCENTAGE = ["G", "J", "L"]
COL_DISPOSITION = "O"

DISPOSITION_CF_MAP = {
    "Liquidate": "critical",
    "20% Markdown": "slow_moving",
    "10% Markdown": "watch",
    "Transfer First": "watch",
    "Hold": "healthy",
}


def _assign_demand_trend(demand_90_day: int, sell_through_rate: float) -> str:
    """Classify demand trend for markdown planning."""
    if demand_90_day == 0:
        return "No Demand"
    if sell_through_rate < 0.10:
        return "Declining"
    if sell_through_rate < 0.20:
        return "Slow"
    return "Moderate"


def _assign_markdown_plan(
    age_days: int,
    demand_90_day: int,
    status: str,
) -> tuple[float, str]:
    """
    Apply markdown rules in priority order (specs §7.2).

    Returns:
        (suggested_markdown_pct, recommended_disposition)
    """
    thresholds = MARKDOWN_THRESHOLDS

    if age_days > thresholds["liquidate_age_days"] and demand_90_day == 0:
        return thresholds["liquidate_markdown_pct"], thresholds["liquidate_disposition"]

    if age_days > thresholds["markdown_20_age_days"]:
        return thresholds["markdown_20_pct"], thresholds["markdown_20_disposition"]

    if age_days > thresholds["markdown_10_age_days"]:
        return thresholds["markdown_10_pct"], thresholds["markdown_10_disposition"]

    if status == "Excess":
        return thresholds["hold_markdown_pct"], "Transfer First"

    return thresholds["hold_markdown_pct"], thresholds["hold_disposition"]


def _project_sell_through(current_rate: float, markdown_pct: float) -> float:
    """Estimate sell-through improvement after markdown."""
    if markdown_pct == 0:
        return round(current_rate, 4)
    uplift = markdown_pct * 0.45
    return round(min(current_rate + uplift, 0.90), 4)


def _build_markdown_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    """Build markdown planner rows from eligible inventory records."""
    if df.empty:
        return pd.DataFrame(columns=HEADERS)

    working = df[df["status"].isin(ELIGIBLE_STATUSES)].copy()
    if working.empty:
        return pd.DataFrame(columns=HEADERS)

    rows = []
    for _, record in working.iterrows():
        age_days = int(record["age_days"])
        demand = int(record["demand_90_day"])
        status = str(record["status"])
        qty = int(record["quantity_on_hand"])
        unit_cost = float(record["unit_cost"])
        selling_price = float(record["selling_price"])
        sell_through = float(record["sell_through_rate"])
        current_margin = float(record["gross_margin_pct"])

        markdown_pct, disposition = _assign_markdown_plan(age_days, demand, status)
        markdown_price = round(selling_price * (1 - markdown_pct), 2)
        projected_str = _project_sell_through(sell_through, markdown_pct)
        recovery_value = round(qty * markdown_price * projected_str, 2)

        current_margin_dollars = (selling_price - unit_cost) * qty * sell_through
        projected_margin_dollars = (markdown_price - unit_cost) * qty * projected_str
        margin_impact = round(projected_margin_dollars - current_margin_dollars, 2)

        rows.append(
            {
                "SKU": record["sku"],
                "Product Name": record["product_name"],
                "Location": record["location"],
                "Quantity On Hand": qty,
                "Unit Cost": unit_cost,
                "Current Selling Price": selling_price,
                "Current Margin %": current_margin,
                "Age Days": age_days,
                "Demand Trend": _assign_demand_trend(demand, sell_through),
                "Suggested Markdown %": markdown_pct,
                "Markdown Price": markdown_price,
                "Projected Sell Through %": projected_str,
                "Estimated Recovery Value": recovery_value,
                "Margin Impact": margin_impact,
                "Recommended Disposition": disposition,
            }
        )

    result = pd.DataFrame(rows, columns=HEADERS)
    disposition_order = {
        "Liquidate": 0,
        "20% Markdown": 1,
        "10% Markdown": 2,
        "Transfer First": 3,
        "Hold": 4,
    }
    result["_sort"] = result["Recommended Disposition"].map(disposition_order)
    return (
        result.sort_values(
            ["_sort", "Estimated Recovery Value"], ascending=[True, False]
        )
        .drop(columns="_sort")
        .reset_index(drop=True)
    )


def _write_title(ws: Worksheet, record_count: int) -> None:
    """Render sheet title."""
    ws.merge_cells(
        start_row=TITLE_ROW, start_column=1, end_row=TITLE_ROW, end_column=COL_COUNT
    )
    cell = ws.cell(
        row=TITLE_ROW,
        column=1,
        value=f"Markdown Planner — {record_count:,} Candidate Items",
    )
    cell.font = Font(
        name=sc.FONTS["default_name"],
        bold=True,
        color=sc.COLORS["white"],
        size=sc.FONTS["title_size"],
    )
    cell.fill = sc.HEADER_FILL
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[TITLE_ROW].height = 28


def _write_headers(ws: Worksheet) -> None:
    """Write and style table headers."""
    for col_idx, header in enumerate(HEADERS, start=1):
        ws.cell(row=HEADER_ROW, column=col_idx, value=header)
    apply_table_header_style(ws, HEADER_ROW, COL_COUNT)
    ws.row_dimensions[HEADER_ROW].height = 22


def _write_data_rows(ws: Worksheet, planner: pd.DataFrame) -> int:
    """Write planner rows and return last row."""
    if planner.empty:
        return HEADER_ROW

    for row_offset, (_, row) in enumerate(planner.iterrows()):
        excel_row = DATA_START_ROW + row_offset
        for col_idx, header in enumerate(HEADERS, start=1):
            value = row[header]
            if pd.isna(value):
                value = None
            ws.cell(row=excel_row, column=col_idx, value=value)
    return DATA_START_ROW + len(planner) - 1


def _apply_column_formats(ws: Worksheet, last_row: int) -> None:
    """Apply currency, percentage, and integer formats."""
    if last_row < DATA_START_ROW:
        return
    format_integer_columns(ws, COL_INTEGER, DATA_START_ROW, last_row)
    format_currency_columns(ws, COL_CURRENCY, DATA_START_ROW, last_row)
    format_percentage_columns(ws, COL_PERCENTAGE, DATA_START_ROW, last_row)


def _create_planner_table(ws: Worksheet, last_row: int) -> None:
    """Create filterable Excel table."""
    end_col = get_column_letter(COL_COUNT)
    if last_row < DATA_START_ROW:
        last_row = DATA_START_ROW
    create_excel_table(ws, TABLE_NAME, f"A{HEADER_ROW}:{end_col}{last_row}")


def _apply_disposition_formatting(ws: Worksheet, last_row: int) -> None:
    """Apply conditional formatting to Recommended Disposition."""
    if last_row < DATA_START_ROW:
        return
    disposition_range = f"{COL_DISPOSITION}{DATA_START_ROW}:{COL_DISPOSITION}{last_row}"
    apply_risk_conditional_formatting(
        ws,
        disposition_range,
        COL_DISPOSITION,
        DATA_START_ROW,
        DISPOSITION_CF_MAP,
    )


def build(ws: Worksheet, context: dict[str, Any]) -> None:
    """Build the Markdown Planner sheet from inventory data."""
    df: pd.DataFrame = context["data"].get("inventory", pd.DataFrame())
    planner = _build_markdown_dataframe(df)

    _write_title(ws, len(planner))
    _write_headers(ws)
    last_row = _write_data_rows(ws, planner)

    _apply_column_formats(ws, last_row)
    _create_planner_table(ws, last_row)
    _apply_disposition_formatting(ws, last_row)

    freeze_panes(ws, row=DATA_START_ROW, col=1)
    autosize_columns(ws, min_width=10, max_width=38)
    ws.column_dimensions["B"].width = min(ws.column_dimensions["B"].width, 32)

    ws.sheet_view.showGridLines = False
    set_landscape_print(
        ws, fit_width=1, repeat_header_rows=f"{HEADER_ROW}:{HEADER_ROW}"
    )
