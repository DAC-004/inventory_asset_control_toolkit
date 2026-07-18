"""Post-save workbook validation for production builds."""

from __future__ import annotations

import math
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

from config.workbook_config import OUTPUT_FILENAME, SHEET_ORDER
from src.exceptions import WorkbookValidationError

FORBIDDEN_SHEET_NAMES = frozenset(
    {"IT Assets", "Software", "Hardware", "License", "Disposal Log"}
)


def _scan_sheet_for_invalid_numbers(ws) -> list[str]:
    """Return cell addresses containing NaN or infinity."""
    problems: list[str] = []
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            value = cell.value
            if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
                problems.append(f"{ws.title}!{cell.coordinate}")
                if len(problems) >= 10:
                    return problems
    return problems


def validate_workbook(wb: Workbook) -> None:
    """Validate an open workbook instance."""
    if wb.sheetnames != list(SHEET_ORDER):
        raise WorkbookValidationError(
            f"Sheet order mismatch: expected {SHEET_ORDER}, got {wb.sheetnames}"
        )

    if len(wb.sheetnames) != 14:
        raise WorkbookValidationError(f"Expected 14 sheets, found {len(wb.sheetnames)}")

    forbidden = FORBIDDEN_SHEET_NAMES.intersection(wb.sheetnames)
    if forbidden:
        raise WorkbookValidationError(
            f"Forbidden IT-oriented sheets present: {sorted(forbidden)}"
        )

    table_names: list[str] = []
    for ws in wb.worksheets:
        table_names.extend(ws.tables.keys())
        invalid = _scan_sheet_for_invalid_numbers(ws)
        if invalid:
            raise WorkbookValidationError(
                f"Invalid numeric values in cells: {', '.join(invalid)}"
            )

    if len(table_names) != len(set(table_names)):
        duplicates = {n for n in table_names if table_names.count(n) > 1}
        raise WorkbookValidationError(
            f"Duplicate Excel table names: {sorted(duplicates)}"
        )


def validate_workbook_file(path: Path, *, require_standard_name: bool = True) -> None:
    """Reopen and validate a workbook file on disk."""
    if not path.exists():
        raise WorkbookValidationError(
            f"Workbook file not found: {path.name}", path=path
        )

    if require_standard_name and path.name != OUTPUT_FILENAME:
        raise WorkbookValidationError(
            f"Unexpected output filename: {path.name} (expected {OUTPUT_FILENAME})",
            path=path,
        )

    if path.stat().st_size <= 0:
        raise WorkbookValidationError("Workbook file is empty", path=path)

    wb = load_workbook(path, read_only=False, data_only=False)
    try:
        validate_workbook(wb)
    finally:
        wb.close()
