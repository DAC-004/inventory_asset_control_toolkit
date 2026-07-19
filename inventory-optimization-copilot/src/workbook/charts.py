"""
openpyxl chart builders for dashboard and management summary sheets.
"""

from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.worksheet.worksheet import Worksheet


def _apply_chart_title(chart, title: str) -> None:
    """Set chart title with consistent font styling."""
    chart.title = title
    if chart.title and chart.title.tx:
        chart.title.tx.rich.paragraphs[0].pPr = None  # use default


def create_bar_chart(
    title: str,
    data_ref: Reference,
    categories_ref: Reference,
    grouping: str = "clustered",
    style: int = 10,
    width: float = 15,
    height: float = 10,
    y_axis_title: str | None = None,
    x_axis_title: str | None = None,
    horizontal: bool = False,
    tick_label_skip: int = 1,
    hide_legend: bool = False,
    value_axis_min: float | None = None,
    value_axis_max: float | None = None,
    show_data_labels: bool = False,
) -> BarChart:
    """Build a styled bar/column chart (not yet anchored to a worksheet)."""
    chart = BarChart()
    chart.type = "bar" if horizontal else "col"
    chart.grouping = grouping
    chart.style = style
    chart.width = width
    chart.height = height
    _apply_chart_title(chart, title)
    chart.title.overlay = False
    chart.legend.overlay = False

    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(categories_ref)

    value_axis = chart.x_axis if horizontal else chart.y_axis
    category_axis = chart.y_axis if horizontal else chart.x_axis
    if y_axis_title:
        value_axis.title = y_axis_title
    if x_axis_title:
        category_axis.title = x_axis_title
    category_axis.tickLblSkip = max(1, tick_label_skip)

    if value_axis_min is not None:
        value_axis.scaling.min = value_axis_min
    if value_axis_max is not None:
        value_axis.scaling.max = value_axis_max

    if hide_legend:
        chart.legend = None
    else:
        chart.legend.position = "b"

    if show_data_labels:
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showVal = True
        chart.dataLabels.showCatName = False
        chart.dataLabels.showLegendKey = False

    return chart


def create_pie_chart(
    title: str,
    data_ref: Reference,
    categories_ref: Reference,
    style: int = 10,
    width: float = 12,
    height: float = 10,
    show_percent_labels: bool = False,
    legend_position: str = "r",
) -> PieChart:
    """Build a styled pie chart (not yet anchored to a worksheet)."""
    chart = PieChart()
    chart.style = style
    chart.width = width
    chart.height = height
    _apply_chart_title(chart, title)
    chart.title.overlay = False

    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(categories_ref)

    chart.legend.position = legend_position
    chart.legend.overlay = False

    if show_percent_labels:
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showPercent = True
        chart.dataLabels.showCatName = False
        chart.dataLabels.showVal = False
        chart.dataLabels.showLeaderLines = True

    return chart


def create_clustered_bar_chart(
    title: str,
    data_ref: Reference,
    categories_ref: Reference,
    style: int = 10,
    width: float = 26,
    height: float = 10,
    horizontal: bool = False,
    y_axis_title: str | None = None,
    x_axis_title: str | None = None,
    value_axis_min: float | None = None,
    value_axis_max: float | None = None,
) -> BarChart:
    """Build a multi-series clustered bar/column chart from one data block."""
    chart = BarChart()
    chart.type = "bar" if horizontal else "col"
    chart.grouping = "clustered"
    chart.style = style
    chart.width = width
    chart.height = height
    _apply_chart_title(chart, title)
    chart.title.overlay = False
    chart.legend.position = "b"
    chart.legend.overlay = False

    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(categories_ref)

    value_axis = chart.x_axis if horizontal else chart.y_axis
    category_axis = chart.y_axis if horizontal else chart.x_axis
    if y_axis_title:
        value_axis.title = y_axis_title
    if x_axis_title:
        category_axis.title = x_axis_title
    if value_axis_min is not None:
        value_axis.scaling.min = value_axis_min
    if value_axis_max is not None:
        value_axis.scaling.max = value_axis_max
    category_axis.tickLblSkip = 1
    return chart


def add_bar_chart(
    ws: Worksheet,
    title: str,
    data_ref: Reference,
    categories_ref: Reference,
    anchor: str = "J11",
    **kwargs,
) -> BarChart:
    """Create a bar chart and anchor it on the worksheet."""
    chart = create_bar_chart(title, data_ref, categories_ref, **kwargs)
    ws.add_chart(chart, anchor)
    return chart


def add_pie_chart(
    ws: Worksheet,
    title: str,
    data_ref: Reference,
    categories_ref: Reference,
    anchor: str = "J20",
    **kwargs,
) -> PieChart:
    """Create a pie chart and anchor it on the worksheet."""
    chart = create_pie_chart(title, data_ref, categories_ref, **kwargs)
    ws.add_chart(chart, anchor)
    return chart


def add_clustered_bar_chart(
    ws: Worksheet,
    title: str,
    data_ref: Reference,
    categories_ref: Reference,
    anchor: str = "J11",
    **kwargs,
) -> BarChart:
    """Create a clustered bar chart and anchor it on the worksheet."""
    chart = create_clustered_bar_chart(title, data_ref, categories_ref, **kwargs)
    ws.add_chart(chart, anchor)
    return chart


def add_line_chart(
    ws: Worksheet,
    title: str,
    data_ref: Reference,
    categories_ref: Reference,
    anchor: str = "J11",
    width: float = 14,
    height: float = 8,
    y_axis_title: str | None = None,
) -> LineChart:
    """Create a line chart and anchor it on the worksheet."""
    chart = LineChart()
    chart.style = 10
    chart.width = width
    chart.height = height
    _apply_chart_title(chart, title)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(categories_ref)
    if y_axis_title:
        chart.y_axis.title = y_axis_title
    chart.legend = None
    ws.add_chart(chart, anchor)
    return chart


def make_data_reference(
    ws: Worksheet,
    min_col: int,
    min_row: int,
    max_row: int,
    max_col: int | None = None,
) -> Reference:
    """Build a Reference for chart data on the given worksheet."""
    max_col = max_col or min_col
    return Reference(
        ws, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row
    )


def make_category_reference(
    ws: Worksheet,
    col: int,
    min_row: int,
    max_row: int,
) -> Reference:
    """Build a Reference for chart category labels on the given worksheet."""
    return Reference(ws, min_col=col, min_row=min_row, max_col=col, max_row=max_row)
