"""Expanded tests for inventory data quality and sheet structure."""

from config.workbook_config import SHEET_ORDER
from openpyxl import load_workbook

from src.data_generation.generate_inventory import generate_inventory_data
from src.main import generate_all_data
from src.sheets.master_inventory_sheet import (
    HEADERS as INVENTORY_HEADERS,
    HEADER_ROW as INV_HEADER_ROW,
)
from src.workbook.builder import build_workbook


def test_inventory_duplicate_primary_keys():
    df = generate_inventory_data(row_count=120, seed=42)
    assert df["item_id"].is_unique


def test_required_sheets_have_headers(tmp_path):
    data = generate_all_data()
    output = tmp_path / "workbook.xlsx"
    build_workbook(data, output_path=output)
    wb = load_workbook(output)

    assert (
        wb["Master Inventory"].cell(row=INV_HEADER_ROW, column=1).value
        == INVENTORY_HEADERS[0]
    )
    assert wb.sheetnames == SHEET_ORDER
    assert len(wb.sheetnames) == 10
