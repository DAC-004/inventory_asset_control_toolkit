"""Tests for Prompt 09 workbook UX, dashboard, and reporting."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from config.workbook_config import SHEET_ORDER
from src.main import generate_all_data
from src.workbook.builder import build_workbook, create_workbook


def test_exact_sheet_count_and_order(tmp_path: Path):
    output = tmp_path / "ux_workbook.xlsx"
    build_workbook(generate_all_data(), output_path=output)
    wb = load_workbook(output)
    try:
        assert len(wb.sheetnames) == 14
        assert wb.sheetnames == SHEET_ORDER
    finally:
        wb.close()


def test_workbook_recalculation_on_load():
    wb = create_workbook()
    assert wb.calculation.fullCalcOnLoad is True


def test_no_it_sheets(tmp_path: Path):
    output = tmp_path / "ux_workbook.xlsx"
    build_workbook(generate_all_data(), output_path=output)
    wb = load_workbook(output, read_only=True)
    try:
        forbidden = {"IT Assets", "Software", "Hardware", "License"}
        assert forbidden.isdisjoint(set(wb.sheetnames))
    finally:
        wb.close()


def test_unique_table_names_across_workbook(tmp_path: Path):
    output = tmp_path / "ux_workbook.xlsx"
    build_workbook(generate_all_data(), output_path=output)
    wb = load_workbook(output)
    try:
        names: list[str] = []
        for ws in wb.worksheets:
            names.extend(ws.tables.keys())
        assert len(names) == len(set(names))
    finally:
        wb.close()


def test_operational_sheets_have_tables_and_freeze(tmp_path: Path):
    output = tmp_path / "ux_workbook.xlsx"
    build_workbook(generate_all_data(), output_path=output)
    wb = load_workbook(output)
    try:
        table_sheets = [
            "Master Inventory",
            "Inventory Classification",
            "Replenishment Planning",
            "Transfer Planner",
            "Markdown Planner",
            "Demand Forecast",
            "Service Level Analysis",
            "Purchase Order Tracker",
            "Vendor Scorecards",
        ]
        for name in table_sheets:
            ws = wb[name]
            assert len(ws.tables) >= 1, name
            assert ws.freeze_panes is not None, name
    finally:
        wb.close()


def test_dashboard_charts_and_links(tmp_path: Path):
    output = tmp_path / "ux_workbook.xlsx"
    build_workbook(generate_all_data(), output_path=output)
    wb = load_workbook(output)
    try:
        dash = wb["Inventory Dashboard"]
        assert len(dash._charts) == 10
        assert dash.cell(row=1, column=22).hyperlink is not None
        assert dash.page_setup.orientation == dash.ORIENTATION_LANDSCAPE
    finally:
        wb.close()


def test_management_summary_print_layout(tmp_path: Path):
    output = tmp_path / "ux_workbook.xlsx"
    build_workbook(generate_all_data(), output_path=output)
    wb = load_workbook(output)
    try:
        summary = wb["Management Summary"]
        assert summary.page_setup.orientation == summary.ORIENTATION_LANDSCAPE
        assert summary.print_title_rows is not None
        assert len(summary.row_breaks) >= 1
    finally:
        wb.close()


def test_workbook_round_trip(tmp_path: Path):
    output = tmp_path / "round_trip.xlsx"
    build_workbook(generate_all_data(), output_path=output)
    wb = load_workbook(output)
    try:
        wb.save(tmp_path / "round_trip_copy.xlsx")
    finally:
        wb.close()
    copy = load_workbook(tmp_path / "round_trip_copy.xlsx", read_only=True)
    try:
        assert copy.sheetnames == SHEET_ORDER
    finally:
        copy.close()
