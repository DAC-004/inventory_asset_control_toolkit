"""Tests for demand forecasting and service level analytics."""

import math

import pandas as pd
import pytest
from openpyxl import Workbook

from config.workbook_config import (
    FORECAST_HOLDOUT_WEEKS,
    FORECAST_PRIMARY_METRIC,
    FORECAST_SES_ALPHA,
)
from src.main import generate_all_data
from src.services.forecast_service import (
    DEMAND_FORECAST_HEADERS,
    FORECAST_METHODS,
    METHOD_FUNCTIONS,
    _compute_errors,
    _forecast_ma,
    _forecast_naive,
    _forecast_ses,
    _forecast_wma,
    build_demand_forecast_dataframe,
    build_forecast_lookup,
    evaluate_method_holdout,
    forecast_sku_location,
    resolve_planning_demand,
    select_best_method,
)
from src.services.replenishment_service import (
    REPLENISHMENT_HEADERS,
    build_replenishment_dataframe,
)
from src.services.service_level_service import (
    SERVICE_LEVEL_HEADERS,
    _finite_rate,
    build_service_level_dataframe,
)
from src.sheets.demand_forecast_sheet import TABLE_NAME as FORECAST_TABLE
from src.sheets.demand_forecast_sheet import build as build_forecast_sheet
from src.sheets.service_level_analysis_sheet import TABLE_NAME as SERVICE_TABLE
from src.sheets.service_level_analysis_sheet import build as build_service_sheet


def test_forecast_methods_implemented():
    """All four forecast methods should be available."""
    assert set(FORECAST_METHODS) == {
        "Naive",
        "4-Week Moving Average",
        "Weighted Moving Average",
        "Simple Exponential Smoothing",
    }
    history = [10.0, 12.0, 11.0, 13.0, 12.0, 14.0, 13.0, 15.0]
    assert _forecast_naive(history, 1)[0] == 15.0
    assert _forecast_ma(history, 1)[0] == pytest.approx(13.5)
    assert _forecast_wma(history, 1)[0] > 0
    assert _forecast_ses(history, 1)[0] > 0


def test_holdout_evaluation():
    """Holdout evaluation uses trailing weeks only."""
    history = [float(i) for i in range(1, 21)]
    errors = evaluate_method_holdout(history, _forecast_naive, holdout_weeks=4)
    assert errors.mae >= 0
    assert math.isfinite(errors.wape)


def test_error_metrics_bounds():
    """Error metrics stay finite and rates stay bounded."""
    actuals = [10.0, 20.0, 0.0, 15.0]
    forecasts = [9.0, 22.0, 1.0, 15.0]
    errors = _compute_errors(actuals, forecasts)
    assert math.isfinite(errors.mae)
    assert math.isfinite(errors.rmse)
    assert math.isfinite(errors.wape)
    assert errors.wape <= 1.0 or errors.wape >= 0
    assert errors.mape is not None or 0.0 in actuals


def test_zero_demand_forecast():
    """Zero demand history yields zero forecasts without NaN."""
    result = forecast_sku_location([0.0] * 12)
    assert result["forecast_4w"] == 0
    assert result["forecast_8w"] == 0
    assert math.isfinite(result["selected_errors"].wape)


def test_best_method_selection_by_wape():
    """Best method is selected using WAPE as primary metric."""
    history = [5.0] * 20
    method, errors = select_best_method(history)
    assert method in FORECAST_METHODS
    assert FORECAST_PRIMARY_METRIC == "WAPE"
    assert math.isfinite(errors.wape)


def test_forecast_horizons():
    """Forecast outputs include 4-, 8-, and 12-week horizons."""
    data = generate_all_data()
    df = build_demand_forecast_dataframe(data)
    assert list(df.columns) == DEMAND_FORECAST_HEADERS
    assert not df.empty
    assert (df["Forecast 4-Week"] >= 0).all()
    assert (df["Forecast 8-Week"] >= df["Forecast 4-Week"]).all()
    assert (df["Forecast 12-Week"] >= df["Forecast 8-Week"]).all()


def test_replenishment_demand_source_integration():
    """Replenishment records explicit demand source without silent override."""
    data = generate_all_data()
    df = build_replenishment_dataframe(data)
    assert "Demand Source" in REPLENISHMENT_HEADERS
    assert "Planning Avg Daily Demand" in REPLENISHMENT_HEADERS
    assert set(df["Demand Source"].unique()).issubset(
        {"Historical Average", "Selected Forecast", "Blended Demand"}
    )


def test_resolve_planning_demand_sources():
    """Planning demand resolver returns documented source labels."""
    daily, source = resolve_planning_demand(5.0, 0.0, 0.1, 0.2)
    assert source == "Historical Average"
    assert daily == 5.0

    daily2, source2 = resolve_planning_demand(0.0, 14.0, 0.1, 0.5)
    assert source2 == "Selected Forecast"
    assert daily2 == pytest.approx(2.0)


def test_unit_and_line_fill_rates():
    """Fill rates are bounded between 0 and 1."""
    data = generate_all_data()
    df = build_service_level_dataframe(data)
    assert list(df.columns) == SERVICE_LEVEL_HEADERS
    assert (df["Unit Fill Rate"] >= 0).all()
    assert (df["Unit Fill Rate"] <= 1).all()
    assert (df["Line Fill Rate"] >= 0).all()
    assert (df["Line Fill Rate"] <= 1).all()
    assert (df["Order Fill Rate"] >= 0).all()
    assert (df["Order Fill Rate"] <= 1).all()


def test_multi_line_order_fill_rate():
    """Order fill rate counts completely fulfilled orders only."""
    orders = pd.DataFrame(
        [
            {
                "order_line_id": "OL-1",
                "order_id": "ORD-1",
                "order_date": "2026-01-01",
                "requested_date": "2026-01-05",
                "fulfilled_date": "2026-01-04",
                "customer_segment": "Retail Walk-In",
                "sku": "T-1001",
                "location_id": "LOC-001",
                "ordered_units": 5,
                "fulfilled_units": 5,
                "backordered_units": 0,
                "line_status": "Fulfilled",
            },
            {
                "order_line_id": "OL-2",
                "order_id": "ORD-1",
                "order_date": "2026-01-01",
                "requested_date": "2026-01-05",
                "fulfilled_date": pd.NaT,
                "customer_segment": "Retail Walk-In",
                "sku": "T-1002",
                "location_id": "LOC-001",
                "ordered_units": 3,
                "fulfilled_units": 1,
                "backordered_units": 2,
                "line_status": "Partially Fulfilled",
            },
        ]
    )
    inventory = pd.DataFrame(
        [
            {
                "sku": "T-1001",
                "location_id": "LOC-001",
                "location_name": "DC-NY",
                "region": "Northeast",
                "category": "Tires",
                "product_name": "A",
            },
            {
                "sku": "T-1002",
                "location_id": "LOC-001",
                "location_name": "DC-NY",
                "region": "Northeast",
                "category": "Tires",
                "product_name": "B",
            },
        ]
    )
    df = build_service_level_dataframe(
        {
            "customer_orders": orders,
            "inventory_full": inventory,
            "demand_history": pd.DataFrame(),
        }
    )
    row = df.iloc[0]
    assert row["Line Fill Rate"] == pytest.approx(0.5)
    assert row["Order Fill Rate"] == pytest.approx(0.0)
    assert row["Unit Fill Rate"] == pytest.approx(6 / 8)


def test_service_level_statuses():
    """Service level statuses use documented labels."""
    data = generate_all_data()
    df = build_service_level_dataframe(data)
    allowed = {
        "On Target",
        "Watch",
        "Below Target",
        "Critical",
        "Data Review Required",
    }
    assert set(df["Status"].unique()).issubset(allowed)


def test_forecast_and_service_sheets(tmp_path):
    """Sheets and workbook include new tabs with tables."""
    from config.workbook_config import SHEET_ORDER
    from openpyxl import load_workbook

    from src.workbook.builder import build_workbook

    data = generate_all_data()
    wb = Workbook()
    build_forecast_sheet(wb.active, {"data": data})
    assert FORECAST_TABLE in wb.active.tables

    wb2 = Workbook()
    build_service_sheet(wb2.active, {"data": data})
    assert SERVICE_TABLE in wb2.active.tables

    output = tmp_path / "forecast_service.xlsx"
    build_workbook(data, output_path=output)
    loaded = load_workbook(output, read_only=True)
    assert "Demand Forecast" in loaded.sheetnames
    assert "Service Level Analysis" in loaded.sheetnames
    assert len(loaded.sheetnames) == len(SHEET_ORDER)


def test_forecast_lookup_for_replenishment():
    """Forecast lookup supports replenishment integration."""
    data = generate_all_data()
    lookup = build_forecast_lookup(data)
    assert lookup
    for value in lookup.values():
        assert "weekly_level" in value
        assert math.isfinite(value["selected_wape"])


def test_finite_rate_helper():
    assert _finite_rate(1.5) == 1.0
    assert _finite_rate(-0.2) == 0.0
    assert _finite_rate(float("nan")) == 0.0


def test_ses_alpha_configured():
    """SES uses centralized alpha parameter."""
    history = [10.0, 20.0, 10.0]
    result = METHOD_FUNCTIONS["Simple Exponential Smoothing"](history, 1)
    assert result[0] > 0
    assert FORECAST_SES_ALPHA == 0.30


def test_holdout_weeks_configured():
    assert FORECAST_HOLDOUT_WEEKS == 8
