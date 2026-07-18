"""Tests for Inventory Classification and Cycle Count Plan sheets."""

from openpyxl import Workbook

from src.main import generate_all_data
from src.sheets.cycle_count_plan_sheet import (
    DATA_START_ROW as PLAN_DATA_START,
    HEADER_ROW as PLAN_HEADER_ROW,
    TABLE_NAME as PLAN_TABLE_NAME,
    build as build_plan,
)
from src.sheets.inventory_classification_sheet import (
    DATA_START_ROW,
    HEADER_ROW,
    HEADERS,
    TABLE_NAME,
    build as build_classification,
)


def test_inventory_classification_sheet_structure():
    """Classification sheet includes table, formatting, and charts."""
    wb = Workbook()
    ws = wb.active
    data = generate_all_data()
    build_classification(ws, {"data": data})

    assert ws.cell(row=HEADER_ROW, column=1).value == HEADERS[0]
    assert ws.cell(row=DATA_START_ROW, column=1).value is not None
    assert TABLE_NAME in ws.tables
    assert len(ws.conditional_formatting) >= 1
    assert len(ws._charts) >= 1
    assert ws.sheet_view.showGridLines is False


def test_cycle_count_plan_sheet_structure():
    """Cycle count plan sheet includes table and status formatting."""
    wb = Workbook()
    ws = wb.active
    data = generate_all_data()
    build_plan(ws, {"data": data})

    assert ws.cell(row=PLAN_HEADER_ROW, column=1).value == "Count Priority"
    assert ws.cell(row=PLAN_DATA_START, column=2).value is not None
    assert PLAN_TABLE_NAME in ws.tables
    assert len(ws.conditional_formatting) >= 1
    assert ws.sheet_view.showGridLines is False


def test_workbook_integrity_with_new_sheets(tmp_path):
    """Full workbook build includes classification and cycle count tabs."""
    from config.workbook_config import SHEET_ORDER
    from openpyxl import load_workbook

    from src.workbook.builder import build_workbook

    data = generate_all_data()
    output = tmp_path / "classification_workbook.xlsx"
    build_workbook(data, output_path=output)
    wb = load_workbook(output, read_only=True)
    assert "Inventory Classification" in wb.sheetnames
    assert "Cycle Count Plan" in wb.sheetnames
    assert len(wb.sheetnames) == len(SHEET_ORDER)
