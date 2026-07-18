"""Production readiness, reproducibility, and release audit tests."""

from __future__ import annotations

import math
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from config.workbook_config import (
    CSV_FILES,
    OUTPUT_FILENAME,
    RANDOM_SEED,
    ROW_COUNTS,
    SHEET_ORDER,
    VERSION,
    WORKBOOK_TITLE,
)
from src.data_generation.pipeline import generate_all_datasets
from src.main import generate_all_data
from src.services.kpi_dashboard_service import compute_planning_service_kpis
from src.workbook.builder import build_workbook
from src.workbook.validator import validate_workbook, validate_workbook_file


def test_product_identity_constants():
    assert WORKBOOK_TITLE == "Inventory Optimization Copilot"
    assert VERSION == "v2.0.0"
    assert OUTPUT_FILENAME == "Inventory_Optimization_Copilot.xlsx"
    assert len(SHEET_ORDER) == 14


def test_reproducible_dataset_generation():
    """Same seed should produce identical logical dataset contents."""
    row_count = ROW_COUNTS["inventory"]["default"]
    first = generate_all_datasets(seed=RANDOM_SEED, row_count=row_count, validate=True)
    second = generate_all_datasets(seed=RANDOM_SEED, row_count=row_count, validate=True)

    assert set(first.keys()) == set(second.keys())
    for key in first:
        pd.testing.assert_frame_equal(
            first[key].reset_index(drop=True),
            second[key].reset_index(drop=True),
        )


def test_reproducible_workbook_structure(tmp_path: Path):
    """Same-seed builds should produce identical sheet order and KPI values."""
    data = generate_all_data()
    out1 = tmp_path / "build_a.xlsx"
    out2 = tmp_path / "build_b.xlsx"
    build_workbook(data, output_path=out1)
    build_workbook(data, output_path=out2)

    wb1 = load_workbook(out1, read_only=True)
    wb2 = load_workbook(out2, read_only=True)
    try:
        assert wb1.sheetnames == wb2.sheetnames == SHEET_ORDER
    finally:
        wb1.close()
        wb2.close()

    kpis_a = compute_planning_service_kpis(data)
    kpis_b = compute_planning_service_kpis(data)
    assert kpis_a == kpis_b


def test_workbook_audit(tmp_path: Path):
    """Production workbook must pass structural audit checks."""
    data = generate_all_data()
    output = tmp_path / OUTPUT_FILENAME
    build_workbook(data, output_path=output)

    validate_workbook_file(output)
    wb = load_workbook(output)
    try:
        validate_workbook(wb)
        assert len(wb.sheetnames) == 14
        assert wb.sheetnames == SHEET_ORDER

        table_names: list[str] = []
        chart_count = 0
        for ws in wb.worksheets:
            table_names.extend(ws.tables.keys())
            chart_count += len(ws._charts)
            if ws.freeze_panes is not None:
                assert ws.freeze_panes != "A1" or ws.title == "README"

        assert len(table_names) == len(set(table_names))
        assert chart_count >= 10

        dash = wb["Inventory Dashboard"]
        assert dash.page_setup.orientation == dash.ORIENTATION_LANDSCAPE
        summary = wb["Management Summary"]
        assert summary.page_setup.orientation == summary.ORIENTATION_LANDSCAPE
    finally:
        wb.close()


def test_workbook_no_nan_or_infinity(tmp_path: Path):
    data = generate_all_data()
    output = tmp_path / OUTPUT_FILENAME
    build_workbook(data, output_path=output)
    wb = load_workbook(output)
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows(
                min_row=1, max_row=min(ws.max_row, 500), max_col=ws.max_column
            ):
                for cell in row:
                    if isinstance(cell.value, float):
                        assert not math.isnan(cell.value)
                        assert not math.isinf(cell.value)
    finally:
        wb.close()


def test_eight_dataset_files_generated():
    data = generate_all_data()
    assert len(CSV_FILES) == 8
    for key in CSV_FILES:
        assert key in data
        assert not data[key].empty


def test_atomic_build_cleans_temp_on_validation_failure(tmp_path: Path, monkeypatch):
    """Failed validation should not leave temp files behind."""
    from src.workbook import builder as builder_mod

    data = generate_all_data()
    output = tmp_path / OUTPUT_FILENAME

    def fail_validate(_path: Path) -> None:
        from src.exceptions import WorkbookValidationError

        raise WorkbookValidationError("simulated failure")

    monkeypatch.setattr(builder_mod, "validate_workbook_file", fail_validate)

    from src.exceptions import WorkbookBuildError

    with pytest.raises(WorkbookBuildError):
        build_workbook(data, output_path=output)

    assert not output.exists()
    assert list(tmp_path.glob(".build_*.xlsx")) == []
