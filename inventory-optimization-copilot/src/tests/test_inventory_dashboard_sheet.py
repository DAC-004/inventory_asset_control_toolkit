"""Tests for Inventory Dashboard sheet builder."""

from openpyxl import Workbook
from openpyxl.utils import column_index_from_string

from src.data_generation.generate_inventory import generate_inventory_data
from src.main import generate_all_data
from src.sheets.dashboard_layout import (
    CHART_AREA_START_COL,
    DASHBOARD_CANVAS_COLS,
)
from src.sheets.inventory_dashboard_sheet import build


def test_inventory_dashboard_kpi_sections():
    """Dashboard should render KPI cards in the header band."""
    wb = Workbook()
    ws = wb.active
    data = generate_all_data()
    build(ws, {"data": data})

    assert ws.cell(row=1, column=1).value.startswith("Inventory Dashboard")
    assert ws.cell(row=3, column=1).value == "Total Inventory Value"
    assert str(ws.cell(row=4, column=1).value).startswith("=SUM(")
    assert ws.freeze_panes == "A9"


def test_inventory_dashboard_summary_tables_and_charts():
    """Dashboard should include ten dynamic sections with side-by-side charts."""
    wb = Workbook()
    ws = wb.active
    ctx: dict = {"data": generate_all_data()}
    build(ws, ctx)

    layouts = ctx["_dashboard_section_layouts"]
    for section in layouts:
        assert ws.cell(row=section.start_row, column=1).value == section.title

    assert len(ws._charts) == 10
    for chart in ws._charts:
        assert 8 <= chart.height <= 12.5
        assert 22 <= chart.width <= 28.5

    anchor_cols = []
    for chart in ws._charts:
        anchor = chart.anchor
        if hasattr(anchor, "_from"):
            anchor_cols.append(anchor._from.col + 1)
        else:
            col_letter = "".join(ch for ch in str(anchor) if ch.isalpha())
            anchor_cols.append(column_index_from_string(col_letter))
    assert all(col >= CHART_AREA_START_COL for col in anchor_cols)

    table_names = {t.displayName for t in ws.tables.values()}
    assert sum(1 for name in table_names if name.startswith("Dashboard")) == 10


def test_inventory_dashboard_readme_link():
    """Dashboard should link back to README."""
    wb = Workbook()
    ws = wb.active
    build(ws, {"data": generate_all_data()})
    link_cell = ws.cell(row=1, column=DASHBOARD_CANVAS_COLS)
    assert link_cell.hyperlink is not None
    assert "README" in link_cell.hyperlink.target


def test_inventory_dashboard_gridlines_hidden():
    """Dashboard should hide gridlines for a clean executive layout."""
    wb = Workbook()
    ws = wb.active
    df = generate_inventory_data(row_count=20, seed=42)
    build(ws, {"data": {"inventory": df}})

    assert ws.sheet_view.showGridLines is False
