"""Acceptance tests for inventory data generation and workbook build."""

from pathlib import Path

from openpyxl import load_workbook

from config.workbook_config import OUTPUT_FILENAME, SHEET_ORDER, WORKBOOK_PATH
from src.data_generation.generate_inventory import (
    COLUMN_ORDER as INVENTORY_COLUMNS,
    generate_inventory_data,
)
from src.main import generate_all_data
from src.workbook.builder import build_workbook


def test_inventory_data_is_not_empty():
    df = generate_inventory_data(row_count=120, seed=42)
    assert not df.empty


def test_inventory_data_includes_required_columns():
    df = generate_inventory_data(row_count=50, seed=42)
    assert list(df.columns) == INVENTORY_COLUMNS


def test_workbook_file_is_created(tmp_path: Path):
    data = generate_all_data()
    output = tmp_path / OUTPUT_FILENAME
    result = build_workbook(data, output_path=output)
    assert result.exists()
    assert result.stat().st_size > 0


def test_workbook_contains_all_required_sheet_names(tmp_path: Path):
    data = generate_all_data()
    output = tmp_path / "workbook.xlsx"
    build_workbook(data, output_path=output)
    wb = load_workbook(output, read_only=True)
    assert set(SHEET_ORDER).issubset(set(wb.sheetnames))


def test_workbook_sheets_are_in_correct_order(tmp_path: Path):
    data = generate_all_data()
    output = tmp_path / "workbook.xlsx"
    build_workbook(data, output_path=output)
    wb = load_workbook(output, read_only=True)
    assert wb.sheetnames == SHEET_ORDER


def test_workbook_output_path_matches_config():
    assert WORKBOOK_PATH.name == OUTPUT_FILENAME
    assert WORKBOOK_PATH.name == "Inventory_Optimization_Copilot.xlsx"


def test_workbook_can_be_opened_with_openpyxl(tmp_path: Path):
    output = tmp_path / "workbook.xlsx"
    build_workbook(generate_all_data(), output_path=output)
    wb = load_workbook(output)
    try:
        assert len(wb.sheetnames) == len(SHEET_ORDER)
        assert wb["README"].cell(row=1, column=1).value is not None
    finally:
        wb.close()
