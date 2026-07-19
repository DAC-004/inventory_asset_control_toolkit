"""Automated layout verification for Inventory Dashboard A:V grid."""

from __future__ import annotations

import re
from dataclasses import dataclass

from openpyxl import Workbook, load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from src.main import generate_all_data
from src.sheets.dashboard_layout import (
    CHART_AREA_END_COL,
    CHART_AREA_START_COL,
    CHART_MAX_HEIGHT_CM,
    CHART_MAX_WIDTH_CM,
    CHART_MIN_HEIGHT_CM,
    CHART_MIN_WIDTH_CM,
    CHART_TITLES,
    DASHBOARD_CANVAS_COLS,
    DASHBOARD_COLUMN_WIDTHS,
    TABLE_AREA_END_COL,
    chart_layout_specs_from_layouts,
    compute_section_layouts,
)
from src.sheets.inventory_dashboard_sheet import build


@dataclass
class ChartBounds:
    name: str
    anchor_col: int
    anchor_row: int
    end_col: int
    end_row: int


def _anchor_col_row(chart) -> tuple[int, int]:
    anchor = chart.anchor
    if hasattr(anchor, "_from"):
        return anchor._from.col + 1, anchor._from.row + 1
    text = str(anchor)
    col_letter = "".join(ch for ch in text if ch.isalpha())
    row_num = int("".join(ch for ch in text if ch.isdigit()))
    return column_index_from_string(col_letter), row_num


def _build_dashboard_ws():
    wb = Workbook()
    ws = wb.active
    ctx: dict = {"data": generate_all_data()}
    build(ws, ctx)
    return ws, ctx


def test_dashboard_layout_chart_count_and_specs():
    ws, ctx = _build_dashboard_ws()
    layouts = ctx["_dashboard_section_layouts"]
    specs = chart_layout_specs_from_layouts(layouts)
    assert len(ws._charts) == 10
    assert len(specs) == 10

    anchors = []
    for chart in ws._charts:
        col, row = _anchor_col_row(chart)
        anchors.append((col, row))
        assert col >= CHART_AREA_START_COL
        assert col <= CHART_AREA_END_COL
        assert CHART_MIN_WIDTH_CM <= chart.width <= CHART_MAX_WIDTH_CM + 1
        assert CHART_MIN_HEIGHT_CM <= chart.height <= CHART_MAX_HEIGHT_CM + 1

    assert len(set(anchors)) == 10


def test_dashboard_layout_sections_and_columns():
    ws, ctx = _build_dashboard_ws()
    layouts = ctx["_dashboard_section_layouts"]
    specs = chart_layout_specs_from_layouts(layouts)

    for section_spec, layout_spec in zip(layouts, specs, strict=True):
        assert ws.cell(row=section_spec.start_row, column=1).value == section_spec.title
        header_row = section_spec.start_row + 2
        chart_for_section = [
            c
            for c in ws._charts
            if _anchor_col_row(c)[1] == header_row
            and _anchor_col_row(c)[0] >= CHART_AREA_START_COL
        ]
        assert len(chart_for_section) == 1
        col, row = _anchor_col_row(chart_for_section[0])
        assert layout_spec.anchor_cell == f"{get_column_letter(col)}{row}"

    for col_idx in range(1, DASHBOARD_CANVAS_COLS + 1):
        letter = get_column_letter(col_idx)
        dim = ws.column_dimensions[letter]
        assert dim.hidden is False
        assert (dim.width or 0) > 0

    for letter in "EFGHIJKLMNOP":
        assert (ws.column_dimensions[letter].width or 0) > 0


def test_dashboard_layout_no_chart_table_overlap():
    ws, ctx = _build_dashboard_ws()
    layouts = ctx["_dashboard_section_layouts"]
    title_to_section = {s.chart_title: s for s in layouts}
    for chart in ws._charts:
        title = chart.title.tx.rich.paragraphs[0].r[0].t
        section = title_to_section[title]
        col, row = _anchor_col_row(chart)
        assert col >= CHART_AREA_START_COL
        assert col <= CHART_AREA_END_COL
        assert col > TABLE_AREA_END_COL
        assert section.start_row <= row <= section.end_row


def test_dashboard_layout_freeze_and_view():
    ws, _ = _build_dashboard_ws()
    assert ws.freeze_panes == "A9"
    assert ws.sheet_view.showGridLines is False
    assert ws.sheet_view.zoomScale == 85
    assert ws.print_area is not None
    assert "V" in ws.print_area


def test_dashboard_workbook_integrity(tmp_path):
    ws, _ = _build_dashboard_ws()
    path = tmp_path / "dashboard_test.xlsx"
    ws.parent.save(path)
    reloaded = load_workbook(path)
    assert len(reloaded.active._charts) == 10
    reloaded.save(path)


def test_dashboard_column_widths_match_spec():
    ws, _ = _build_dashboard_ws()
    for letter, expected in DASHBOARD_COLUMN_WIDTHS.items():
        assert ws.column_dimensions[letter].width == expected


def test_dashboard_chart_titles_are_business_focused():
    ws, _ = _build_dashboard_ws()
    titles = []
    for chart in ws._charts:
        title = chart.title.tx.rich.paragraphs[0].r[0].t
        titles.append(title)
        assert title not in {"Status", "Aging", "ABC", "Vendor", "Actions", "Summary"}
    assert set(titles) == set(CHART_TITLES.values())


def test_dashboard_chart_layout_spec_registry():
    counts = {key: 5 for key in CHART_TITLES}
    layouts = compute_section_layouts(counts)
    specs = chart_layout_specs_from_layouts(layouts)
    assert len({s.name for s in specs}) == 10
    for spec in specs:
        assert spec.allowed_start_col == CHART_AREA_START_COL
        assert spec.allowed_end_col == CHART_AREA_END_COL
        assert re.fullmatch(r"J\d+", spec.anchor_cell)


def test_replenishment_section_fits_all_status_rows():
    ws, ctx = _build_dashboard_ws()
    meta = ctx["_dashboard_section_meta"]["replenishment"]
    layouts = ctx["_dashboard_section_layouts"]
    repl_section = next(s for s in layouts if s.key == "replenishment")
    data_rows = meta["last_data_row"] - meta["first_data_row"] + 1
    assert data_rows == repl_section.data_row_count
    assert meta["last_data_row"] <= repl_section.end_row

    for row in range(meta["first_data_row"], meta["last_data_row"] + 1):
        dim = ws.row_dimensions[row]
        assert dim.hidden is not True
        assert dim.height not in (0, 0.0)
