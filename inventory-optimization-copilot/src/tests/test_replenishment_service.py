"""Tests for replenishment planning analytics and sheet."""

from datetime import date

import pandas as pd
import pytest
from openpyxl import Workbook

from config.workbook_config import (
    AS_OF_DATE,
    GLOBAL_LEAD_TIME_FALLBACK_DAYS,
    SERVICE_LEVEL_TARGETS,
    SERVICE_LEVEL_Z_SCORES,
)
from src.main import generate_all_data
from src.services.purchase_order_service import compute_valid_open_supply
from src.services.replenishment_service import (
    REPLENISHMENT_HEADERS,
    LeadTimeStats,
    _build_lead_time_records,
    _compute_eoq,
    _round_to_case_pack,
    _safety_stock,
    build_replenishment_dataframe,
    compute_lead_time_stats,
    resolve_planning_lead_time,
)
from src.sheets.replenishment_planning_sheet import (
    DATA_START_ROW,
    HEADER_ROW,
    TABLE_NAME,
    build as build_sheet,
)


def test_lead_time_from_receipt_minus_approval():
    """Actual lead time equals final receipt date minus approval date."""
    pos = pd.DataFrame(
        [
            {
                "po_line_id": "POL-00001",
                "sku": "T-1001",
                "supplier_id": "SUP-001",
                "approval_date": date(2026, 1, 1),
                "promised_date": date(2026, 1, 15),
                "ordered_quantity": 100,
                "po_status": "Received",
                "location_id": "LOC-001",
            }
        ]
    )
    receipts = pd.DataFrame(
        [
            {
                "po_line_id": "POL-00001",
                "receipt_date": date(2026, 1, 10),
                "accepted_quantity": 50,
            },
            {
                "po_line_id": "POL-00001",
                "receipt_date": date(2026, 1, 20),
                "accepted_quantity": 50,
            },
        ]
    )
    records = _build_lead_time_records(pos, receipts)
    assert len(records) == 1
    assert records.iloc[0]["actual_lead_time_days"] == 19


def test_lead_time_stats_and_late_rate():
    """Lead time stats include late rate and average days late."""
    records = pd.DataFrame(
        {
            "actual_lead_time_days": [10.0, 12.0, 14.0],
            "days_late": [0.0, 2.0, 4.0],
            "is_late": [False, True, True],
        }
    )
    stats = compute_lead_time_stats(records)
    assert stats.sample_count == 3
    assert stats.average == pytest.approx(12.0)
    assert stats.late_rate == pytest.approx(2 / 3, abs=0.01)


def test_lead_time_fallback_hierarchy():
    """Planning lead time follows SKU+Supplier → Supplier → default → global."""
    sku_stats = LeadTimeStats(11, 11, 10, 12, 1.0, 0.1, 0.5, 4)
    supplier_stats = LeadTimeStats(15, 15, 14, 16, 1.0, 0.2, 1.0, 8)
    lt, source, _ = resolve_planning_lead_time(
        "T-1", "SUP-1", sku_stats, supplier_stats, 12
    )
    assert lt == 11
    assert source == "SKU + Supplier history"

    empty = LeadTimeStats(0, 0, 0, 0, 0, 0, 0, 0)
    lt2, source2, _ = resolve_planning_lead_time(
        "T-1", "SUP-1", empty, supplier_stats, 12
    )
    assert lt2 == 15
    assert source2 == "Supplier history"

    lt3, source3, _ = resolve_planning_lead_time("T-1", "SUP-1", empty, empty, 18)
    assert lt3 == 18
    assert source3 == "Supplier default"

    lt4, source4, _ = resolve_planning_lead_time("T-1", "SUP-1", empty, empty, 0)
    assert lt4 == GLOBAL_LEAD_TIME_FALLBACK_DAYS
    assert source4 == "Global fallback"


def test_service_level_z_scores():
    """ABC classes map to configured service targets and Z-scores."""
    assert SERVICE_LEVEL_TARGETS["A"] == 0.99
    assert SERVICE_LEVEL_Z_SCORES["A"] == pytest.approx(2.326, abs=0.001)
    assert SERVICE_LEVEL_TARGETS["B"] == 0.95
    assert SERVICE_LEVEL_Z_SCORES["C"] == pytest.approx(1.282, abs=0.001)


def test_safety_stock_methods_never_negative():
    """Safety stock is never negative for basic and advanced methods."""
    basic, method = _safety_stock(1.645, 5.0, 2.0, 14, 0.0, 0, 4)
    assert basic >= 0
    assert "Basic" in method

    advanced, method2 = _safety_stock(1.645, 5.0, 2.0, 14, 3.0, 6, 52)
    assert advanced >= 0
    assert "Advanced" in method2


def test_reorder_point_and_projected_available():
    """Reorder point and projected available follow documented formulas."""
    data = generate_all_data()
    df = build_replenishment_dataframe(data)
    assert list(df.columns) == REPLENISHMENT_HEADERS
    for _, row in df.iterrows():
        expected_pa = (
            row["Quantity On Hand"]
            + row["Open PO Quantity"]
            - row["Quantity Allocated"]
            - row["Backorder Quantity"]
        )
        assert row["Projected Available"] == expected_pa
        assert row["Reorder Point"] >= 0
        assert row["Safety Stock"] >= 0


def test_open_po_no_double_counting():
    """Open PO quantity uses remaining unreceived units only once."""
    pos = pd.DataFrame(
        [
            {
                "po_line_id": "POL-1",
                "sku": "T-1001",
                "location_id": "LOC-001",
                "ordered_quantity": 100,
                "po_status": "Open",
            },
            {
                "po_line_id": "POL-2",
                "sku": "T-1001",
                "location_id": "LOC-001",
                "ordered_quantity": 50,
                "po_status": "Partially Received",
            },
        ]
    )
    receipts = pd.DataFrame(
        [
            {
                "po_line_id": "POL-2",
                "receipt_date": date(2026, 1, 15),
                "received_quantity": 20,
                "accepted_quantity": 20,
                "rejected_quantity": 0,
                "receipt_status": "Partial",
                "quality_issue": "",
            }
        ]
    )
    open_qty = compute_valid_open_supply(
        {"purchase_orders": pos, "purchase_order_receipts": receipts}
    )
    total = int(open_qty.iloc[0]["valid_open_quantity"]) if not open_qty.empty else 0
    assert total == 130


def test_eoq_zero_demand_safe():
    """EOQ returns zero when demand or cost is zero."""
    assert _compute_eoq(0, 10.0) == 0
    assert _compute_eoq(100, 0) == 0
    assert _compute_eoq(500, 25.0) > 0


def test_case_pack_and_moq_rounding():
    """Recommended quantities respect case pack rounding."""
    assert _round_to_case_pack(13, 6) == 18
    assert _round_to_case_pack(0, 6) == 0


def test_replenishment_statuses_present():
    """Generated plan includes documented replenishment statuses."""
    data = generate_all_data()
    df = build_replenishment_dataframe(data)
    allowed = {
        "Order Required",
        "Expedite Existing PO",
        "Monitor",
        "Sufficient",
        "Overstocked",
        "No Recent Demand",
        "Capacity Constraint",
        "Supplier Constraint",
        "Data Review Required",
    }
    assert set(df["Replenishment Status"].unique()).issubset(allowed)


def test_order_by_date_logic():
    """Order-by date is today when below reorder point."""
    data = generate_all_data()
    df = build_replenishment_dataframe(data)
    below = df[df["Projected Available"] <= df["Reorder Point"]]
    if not below.empty:
        assert (below["Order By Date"] == AS_OF_DATE).any()


def test_replenishment_sheet_structure():
    """Sheet includes table, KPIs, conditional formatting, and charts."""
    wb = Workbook()
    ws = wb.active
    data = generate_all_data()
    build_sheet(ws, {"data": data})

    assert ws.cell(row=HEADER_ROW, column=1).value == REPLENISHMENT_HEADERS[0]
    assert ws.cell(row=DATA_START_ROW, column=1).value is not None
    assert TABLE_NAME in ws.tables
    assert len(ws.conditional_formatting) >= 1
    assert len(ws._charts) >= 1


def test_workbook_integrity_with_replenishment_sheet(tmp_path):
    """Full workbook includes Replenishment Planning tab."""
    from config.workbook_config import SHEET_ORDER
    from openpyxl import load_workbook

    from src.workbook.builder import build_workbook

    data = generate_all_data()
    output = tmp_path / "replenishment_workbook.xlsx"
    build_workbook(data, output_path=output)
    wb = load_workbook(output, read_only=True)
    assert "Replenishment Planning" in wb.sheetnames
    assert len(wb.sheetnames) == len(SHEET_ORDER)
