"""Tests for Management Summary sheet builder."""

from openpyxl import Workbook

from src.data_generation.generate_inventory import generate_inventory_data
from src.main import generate_all_data
from src.sheets.management_summary_sheet import (
    KPI1_VALUE_ROW,
    RISK_DATA_START_ROW,
    SECTION1_ROW,
    SECTION4_ROW,
    _top_inventory_risks,
    build,
)


def test_top_inventory_risk_helper_returns_rows():
    inventory = generate_inventory_data(row_count=120, seed=42)
    inv_risks = _top_inventory_risks(inventory)
    assert 1 <= len(inv_risks) <= 5
    assert "SKU" in inv_risks[0]


def test_management_summary_sheet_structure():
    """Sheet should include inventory executive sections and print settings."""
    wb = Workbook()
    ws = wb.active
    data = generate_all_data()
    build(ws, {"data": data})

    assert "Inventory Health" in str(ws.cell(row=SECTION1_ROW, column=1).value)
    assert "Recommended Action" in str(ws.cell(row=SECTION1_ROW + 4, column=1).value)
    assert "Top Inventory Risks" in str(ws.cell(row=SECTION1_ROW + 11, column=1).value)
    assert "30 / 60 / 90" in str(ws.cell(row=SECTION4_ROW, column=1).value)
    assert str(ws.cell(row=KPI1_VALUE_ROW, column=1).value).startswith("=")
    assert ws.cell(row=RISK_DATA_START_ROW, column=1).value
    assert ws.sheet_view.showGridLines is False
    assert ws.page_setup.orientation == ws.ORIENTATION_LANDSCAPE
