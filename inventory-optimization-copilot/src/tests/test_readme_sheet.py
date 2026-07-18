"""Tests for README sheet content and structure."""

from openpyxl import Workbook

from config import style_config as sc
from config.workbook_config import AUTHOR, VERSION, WORKBOOK_TITLE
from src.sheets.readme_sheet import build


def test_readme_sheet_contains_required_sections():
    """README sheet should include title, disclaimer, demo path, and capabilities."""
    wb = Workbook()
    ws = wb.active
    ws.title = "README"

    build(ws, {"data": {}})

    values = []
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=4):
        for cell in row:
            if cell.value:
                values.append(str(cell.value))

    combined = " ".join(values)
    assert WORKBOOK_TITLE in combined
    assert VERSION in combined
    assert AUTHOR in combined
    assert "fictional sample data" in combined.lower()
    assert "Recommended Demo Path" in combined
    assert "Key Capabilities" in combined
    assert "Inventory Dashboard" in combined
    assert "IT Asset Register" not in combined


def test_readme_sheet_has_formatted_title_area():
    """README title cell should use navy banner styling."""
    wb = Workbook()
    ws = wb.active
    build(ws, {"data": {}})

    title_cell = ws.cell(row=1, column=1)
    assert title_cell.value == WORKBOOK_TITLE
    assert title_cell.font.bold is True
    assert sc.COLORS["navy"] in str(title_cell.fill.start_color.rgb)
