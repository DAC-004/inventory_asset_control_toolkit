"""Tests for Aged Excess Analysis sheet builder."""

import pandas as pd
from openpyxl import Workbook

from src.data_generation.generate_inventory import generate_inventory_data
from src.services.inventory_health_service import (
    assign_risk_level,
    build_analysis_dataframe,
)
from src.sheets.aged_excess_sheet import (
    DATA_START_ROW,
    HEADER_ROW,
    HEADERS,
    TABLE_NAME,
    build,
)


def test_assign_risk_level_mapping():
    """Risk levels should map correctly from inventory status."""
    assert assign_risk_level("Obsolete") == "High"
    assert assign_risk_level("Excess / Aged") == "High"
    assert assign_risk_level("Stockout Risk") == "High"
    assert assign_risk_level("Excess") == "Medium"
    assert assign_risk_level("Slow-Moving") == "Medium"
    assert assign_risk_level("Healthy") == "Low"


def test_build_analysis_dataframe_calculations():
    """Analysis DataFrame should compute excess quantity and exclude healthy rows."""
    df = generate_inventory_data(row_count=120, seed=42)
    analysis = build_analysis_dataframe(df)

    assert list(analysis.columns) == HEADERS
    assert "Healthy" not in analysis["Risk Level"].values
    assert (analysis["Excess Quantity"] >= 0).all()
    merged = analysis.merge(
        df[["sku", "location", "quantity_on_hand", "max_stock"]],
        left_on=["SKU", "Location", "Quantity On Hand", "Max Stock"],
        right_on=["sku", "location", "quantity_on_hand", "max_stock"],
    )
    expected = (merged["quantity_on_hand"] - merged["max_stock"]).clip(lower=0)
    pd.testing.assert_series_equal(
        merged["Excess Quantity"].reset_index(drop=True),
        expected.reset_index(drop=True),
        check_names=False,
    )


def test_aged_excess_sheet_structure():
    """Sheet should include table, conditional formatting, and validations."""
    wb = Workbook()
    ws = wb.active
    df = generate_inventory_data(row_count=50, seed=42)
    build(ws, {"data": {"inventory": df}})

    assert ws.cell(row=HEADER_ROW, column=1).value == HEADERS[0]
    assert ws.cell(row=DATA_START_ROW, column=1).value is not None
    assert TABLE_NAME in ws.tables
    assert len(ws.conditional_formatting) >= 1
    assert len(ws.data_validations.dataValidation) >= 1
    assert ws.sheet_view.showGridLines is False
