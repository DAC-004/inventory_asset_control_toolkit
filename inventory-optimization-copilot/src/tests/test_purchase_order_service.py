"""Tests for purchase order tracking and vendor scorecards."""

from datetime import date

import pandas as pd
import pytest
from openpyxl import Workbook

from config.workbook_config import AS_OF_DATE, VENDOR_SCORE_WEIGHTS
from src.main import generate_all_data
from src.services.purchase_order_service import (
    PO_TRACKER_HEADERS,
    aggregate_receipts,
    build_po_tracker_dataframe,
    compute_valid_open_supply,
)
from src.services.replenishment_service import (
    REPLENISHMENT_HEADERS,
    build_replenishment_dataframe,
)
from src.services.vendor_scorecard_service import (
    VENDOR_SCORECARD_HEADERS,
    build_vendor_risk_lookup,
    build_vendor_scorecard_dataframe,
    validate_vendor_score_weights,
)
from src.sheets.purchase_order_tracker_sheet import TABLE_NAME as PO_TABLE
from src.sheets.purchase_order_tracker_sheet import build as build_po_sheet
from src.sheets.vendor_scorecards_sheet import TABLE_NAME as VENDOR_TABLE
from src.sheets.vendor_scorecards_sheet import build as build_vendor_sheet


def test_receipt_aggregation():
    """Receipts aggregate received, accepted, rejected, and dates."""
    receipts = pd.DataFrame(
        [
            {
                "po_line_id": "POL-1",
                "receipt_date": date(2026, 1, 10),
                "received_quantity": 50,
                "accepted_quantity": 48,
                "rejected_quantity": 2,
                "receipt_status": "Partial",
                "quality_issue": "",
            },
            {
                "po_line_id": "POL-1",
                "receipt_date": date(2026, 1, 20),
                "received_quantity": 50,
                "accepted_quantity": 50,
                "rejected_quantity": 0,
                "receipt_status": "Complete",
                "quality_issue": "",
            },
        ]
    )
    agg = aggregate_receipts(receipts)
    assert int(agg.iloc[0]["received_quantity"]) == 100
    assert int(agg.iloc[0]["accepted_quantity"]) == 98
    assert int(agg.iloc[0]["rejected_quantity"]) == 2


def test_open_quantity_and_value():
    """Open quantity and value reflect ordered minus accepted."""
    data = generate_all_data()
    tracker = build_po_tracker_dataframe(data)
    assert list(tracker.columns) == PO_TRACKER_HEADERS
    open_lines = tracker[tracker["Open Quantity"] > 0]
    for _, row in open_lines.iterrows():
        assert row["Open PO Value"] == pytest.approx(
            row["Open Quantity"] * row["Unit Cost"], abs=0.01
        )


def test_cancelled_not_open_supply():
    """Cancelled PO lines have zero open quantity."""
    pos = pd.DataFrame(
        [
            {
                "po_line_id": "POL-1",
                "po_number": "PO-1",
                "supplier_id": "SUP-001",
                "sku": "T-1001",
                "location_id": "LOC-001",
                "buyer": "Buyer A",
                "order_date": date(2026, 1, 1),
                "approval_date": date(2026, 1, 2),
                "promised_date": date(2026, 1, 20),
                "expected_date": date(2026, 1, 22),
                "ordered_quantity": 100,
                "unit_cost": 10.0,
                "extended_cost": 1000.0,
                "case_pack": 1,
                "minimum_order_quantity": 1,
                "po_status": "Cancelled",
            }
        ]
    )
    tracker = build_po_tracker_dataframe({"purchase_orders": pos})
    assert tracker.iloc[0]["Open Quantity"] == 0
    supply = compute_valid_open_supply({"purchase_orders": pos})
    assert supply.empty


def test_days_late_open_orders_use_as_of_date():
    """Open late orders calculate days late from AS_OF_DATE."""
    pos = pd.DataFrame(
        [
            {
                "po_line_id": "POL-2",
                "po_number": "PO-2",
                "supplier_id": "SUP-001",
                "sku": "T-1001",
                "location_id": "LOC-001",
                "buyer": "Buyer A",
                "order_date": date(2026, 5, 1),
                "approval_date": date(2026, 5, 2),
                "promised_date": date(2026, 6, 1),
                "expected_date": date(2026, 6, 5),
                "ordered_quantity": 50,
                "unit_cost": 10.0,
                "extended_cost": 500.0,
                "case_pack": 1,
                "minimum_order_quantity": 1,
                "po_status": "Open",
            }
        ]
    )
    tracker = build_po_tracker_dataframe({"purchase_orders": pos})
    assert tracker.iloc[0]["PO Status"] == "Late"
    assert tracker.iloc[0]["Days Late"] == (AS_OF_DATE - date(2026, 6, 1)).days


def test_over_receipt_status():
    """Accepted quantity exceeding ordered yields Over-Received status."""
    pos = pd.DataFrame(
        [
            {
                "po_line_id": "POL-3",
                "po_number": "PO-3",
                "supplier_id": "SUP-001",
                "sku": "T-1001",
                "location_id": "LOC-001",
                "buyer": "Buyer A",
                "order_date": date(2026, 1, 1),
                "approval_date": date(2026, 1, 2),
                "promised_date": date(2026, 1, 20),
                "expected_date": date(2026, 1, 22),
                "ordered_quantity": 10,
                "unit_cost": 10.0,
                "extended_cost": 100.0,
                "case_pack": 1,
                "minimum_order_quantity": 1,
                "po_status": "Received",
            }
        ]
    )
    receipts = pd.DataFrame(
        [
            {
                "po_line_id": "POL-3",
                "receipt_date": date(2026, 1, 18),
                "received_quantity": 15,
                "accepted_quantity": 15,
                "rejected_quantity": 0,
                "receipt_status": "Complete",
                "quality_issue": "",
            }
        ]
    )
    tracker = build_po_tracker_dataframe(
        {"purchase_orders": pos, "purchase_order_receipts": receipts}
    )
    assert tracker.iloc[0]["PO Status"] == "Over-Received"


def test_quality_hold_excluded_from_valid_supply():
    """Quality hold lines are excluded from valid replenishment supply."""
    pos = pd.DataFrame(
        [
            {
                "po_line_id": "POL-4",
                "po_number": "PO-4",
                "supplier_id": "SUP-001",
                "sku": "T-1001",
                "location_id": "LOC-001",
                "buyer": "Buyer A",
                "order_date": date(2026, 5, 1),
                "approval_date": date(2026, 5, 2),
                "promised_date": date(2026, 6, 15),
                "expected_date": date(2026, 6, 20),
                "ordered_quantity": 40,
                "unit_cost": 10.0,
                "extended_cost": 400.0,
                "case_pack": 1,
                "minimum_order_quantity": 1,
                "po_status": "Open",
            }
        ]
    )
    receipts = pd.DataFrame(
        [
            {
                "po_line_id": "POL-4",
                "receipt_date": date(2026, 6, 10),
                "received_quantity": 40,
                "accepted_quantity": 0,
                "rejected_quantity": 40,
                "receipt_status": "Quality Hold",
                "quality_issue": "Failed inspection",
            }
        ]
    )
    supply = compute_valid_open_supply(
        {"purchase_orders": pos, "purchase_order_receipts": receipts}
    )
    assert supply.empty


def test_purchase_price_variance():
    """PPV compares PO unit cost to inventory baseline."""
    data = generate_all_data()
    tracker = build_po_tracker_dataframe(data)
    assert "Purchase Price Variance" in tracker.columns
    assert "Purchase Price Variance %" in tracker.columns


def test_vendor_score_weights_sum_to_100():
    """Vendor score weights must total 100%."""
    validate_vendor_score_weights()
    assert sum(VENDOR_SCORE_WEIGHTS.values()) == pytest.approx(1.0)


def test_vendor_otif_and_quality_metrics():
    """Vendor scorecard includes OTIF and quality acceptance rates."""
    data = generate_all_data()
    scorecards = build_vendor_scorecard_dataframe(data)
    assert list(scorecards.columns) == VENDOR_SCORECARD_HEADERS
    assert (scorecards["OTIF %"] >= 0).all()
    assert (scorecards["OTIF %"] <= 1).all()
    assert (scorecards["Quality Acceptance %"] >= 0).all()
    assert (scorecards["Quality Acceptance %"] <= 1).all()


def test_vendor_score_and_risk_class():
    """Total vendor score is 0-100 with documented risk classes."""
    data = generate_all_data()
    scorecards = build_vendor_scorecard_dataframe(data)
    allowed = {"Preferred", "Approved", "Watch", "High Risk", "Data Insufficient"}
    assert set(scorecards["Risk Class"].unique()).issubset(allowed)
    assert (scorecards["Total Vendor Score"] >= 0).all()
    assert (scorecards["Total Vendor Score"] <= 100).all()


def test_replenishment_po_integration():
    """Replenishment uses valid open supply and supplier risk."""
    data = generate_all_data()
    df = build_replenishment_dataframe(data)
    assert "Supplier Risk Class" in REPLENISHMENT_HEADERS
    assert set(df["Supplier Risk Class"].unique()).issubset(
        {"Preferred", "Approved", "Watch", "High Risk", "Data Insufficient"}
    )


def test_vendor_risk_lookup():
    """Vendor risk lookup supports replenishment."""
    data = generate_all_data()
    lookup = build_vendor_risk_lookup(data)
    assert lookup
    assert all(isinstance(v, str) for v in lookup.values())


def test_po_and_vendor_sheets(tmp_path):
    """Sheets and workbook include PO tracker and vendor scorecards."""
    from config.workbook_config import SHEET_ORDER
    from openpyxl import load_workbook

    from src.workbook.builder import build_workbook

    data = generate_all_data()
    wb = Workbook()
    build_po_sheet(wb.active, {"data": data})
    assert PO_TABLE in wb.active.tables

    wb2 = Workbook()
    build_vendor_sheet(wb2.active, {"data": data})
    assert VENDOR_TABLE in wb2.active.tables

    output = tmp_path / "po_vendor.xlsx"
    build_workbook(data, output_path=output)
    loaded = load_workbook(output, read_only=True)
    assert "Purchase Order Tracker" in loaded.sheetnames
    assert "Vendor Scorecards" in loaded.sheetnames
    assert len(loaded.sheetnames) == len(SHEET_ORDER)
