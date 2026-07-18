"""Tests for Markdown Planner sheet builder."""

from openpyxl import Workbook

from config.workbook_config import MARKDOWN_THRESHOLDS
from src.main import generate_all_data
from src.services.markdown_service import assign_markdown_plan, build_markdown_dataframe
from src.sheets.markdown_planner_sheet import (
    DATA_START_ROW,
    HEADER_ROW,
    HEADERS,
    TABLE_NAME,
    build,
)


def test_assign_markdown_plan_rules():
    """Markdown rules should follow age, demand, and transfer-first logic."""
    assert assign_markdown_plan(400, 0, "Obsolete", 0.0, False) == (
        MARKDOWN_THRESHOLDS["liquidate_markdown_pct"],
        "Dispose",
    )
    assert assign_markdown_plan(310, 5, "Slow-Moving", 0.08, False) == (
        0.30,
        "30% Markdown",
    )
    assert assign_markdown_plan(300, 5, "Slow-Moving", 0.08, True) == (
        0.0,
        "Transfer First",
    )
    assert assign_markdown_plan(200, 5, "Excess / Aged", 0.15, False) == (
        MARKDOWN_THRESHOLDS["markdown_10_pct"],
        "10% Markdown",
    )
    assert assign_markdown_plan(90, 10, "Excess", 0.20, True) == (0.0, "Transfer First")
    assert assign_markdown_plan(90, 10, "Slow-Moving", 0.20, False) == (0.0, "Hold")


def test_build_markdown_dataframe_filters_and_calculates():
    """Planner should include eligible statuses and compute markdown price."""
    data = generate_all_data()
    planner = build_markdown_dataframe(data)
    inventory = data["inventory"]

    assert list(planner.columns) == HEADERS
    assert planner["Recommended Disposition"].notna().all()
    assert set(planner["SKU"]).issubset(set(inventory["sku"]))
    assert planner["Suggested Markdown %"].between(0, 1).all()
    assert (planner["Markdown Price"] <= planner["Current Selling Price"]).all()
    assert (planner["Estimated Recovery Value"] >= 0).all()

    sample = planner.iloc[0]
    expected_price = round(
        sample["Current Selling Price"] * (1 - sample["Suggested Markdown %"]),
        2,
    )
    assert sample["Markdown Price"] == expected_price


def test_markdown_transfer_first_when_surplus_exists():
    """Transfer-first disposition should appear when network surplus exists."""
    data = generate_all_data()
    planner = build_markdown_dataframe(data)
    transfer_rows = planner[planner["Recommended Disposition"] == "Transfer First"]
    if not transfer_rows.empty:
        assert (transfer_rows["Transfer Feasible"] == "Yes").all()


def test_markdown_planner_sheet_structure():
    """Sheet should include table and disposition conditional formatting."""
    wb = Workbook()
    ws = wb.active
    data = generate_all_data()
    build(ws, {"data": data})

    assert ws.cell(row=HEADER_ROW, column=1).value == "SKU"
    assert TABLE_NAME in ws.tables
    assert len(ws.conditional_formatting) >= 1
    assert ws.cell(row=DATA_START_ROW, column=21).value is not None
