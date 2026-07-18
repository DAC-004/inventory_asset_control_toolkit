"""Tests for Transfer Planner sheet builder."""

from openpyxl import Workbook

from config.workbook_config import TRANSFER_MARGIN_RATE
from src.main import generate_all_data
from src.services.transfer_service import (
    build_transfer_dataframe,
    estimate_transfer_cost,
    is_allowed_route,
    transfer_lead_time_days,
)
from src.sheets.transfer_planner_sheet import (
    DATA_START_ROW,
    HEADER_ROW,
    HEADERS,
    TABLE_NAME,
    build,
)


def test_estimate_transfer_cost_increases_with_quantity():
    """Transfer cost should scale with quantity moved."""
    low = estimate_transfer_cost("DC-NY", "Store-Bronx", 5)
    high = estimate_transfer_cost("DC-NY", "Store-Bronx", 20)
    assert high > low


def test_route_and_lead_time_rules():
    """Routes require different locations and respect lead-time limits."""
    assert is_allowed_route("DC-NY", "Store-Bronx")
    assert not is_allowed_route("DC-NY", "DC-NY")
    assert transfer_lead_time_days("DC-NY", "Store-Bronx") <= 7


def test_build_transfer_dataframe_net_benefit_model():
    """Transfers should include economics and positive quantities."""
    data = generate_all_data()
    transfers = build_transfer_dataframe(data)

    assert list(transfers.columns) == HEADERS
    if transfers.empty:
        return

    assert (transfers["Transfer Quantity"] > 0).all()
    assert (transfers["Source Location"] != transfers["Destination Location"]).all()
    assert (
        transfers["Action"]
        .isin(
            {
                "Transfer",
                "Purchase",
                "Expedite Existing PO",
                "Transfer Then Purchase",
                "Monitor",
                "No Action",
                "Data Review Required",
            }
        )
        .all()
    )

    sample = transfers.iloc[0]
    qty = int(sample["Transfer Quantity"])
    unit_cost = float(sample["Unit Cost"])
    expected_margin_floor = round(qty * unit_cost * TRANSFER_MARGIN_RATE, 2)
    assert float(sample["Estimated Margin Protected"]) >= expected_margin_floor * 0.9
    assert float(sample["Net Benefit"]) == round(
        float(sample["Estimated Margin Protected"])
        - float(sample["Transfer Cost"])
        - float(sample["Expected Source Risk Cost"]),
        2,
    )


def test_transfer_allocation_invariants():
    """Allocated quantities must not exceed initial surplus or requirement."""
    data = generate_all_data()
    transfers = build_transfer_dataframe(data)
    if transfers.empty:
        return

    grouped = transfers.groupby(["SKU", "Source Location", "Destination Location"])
    for _, group in grouped:
        assert (
            group["Transfer Quantity"].sum() <= group["Source Initial Surplus"].iloc[0]
        )
        assert (
            group["Transfer Quantity"].sum()
            <= group["Destination Initial Requirement"].iloc[0]
        )


def test_transfer_planner_sheet_structure():
    """Sheet should include table and conditional formatting rules."""
    wb = Workbook()
    ws = wb.active
    data = generate_all_data()
    build(ws, {"data": data})

    assert ws.cell(row=HEADER_ROW, column=1).value == "Rank"
    assert TABLE_NAME in ws.tables
    assert len(ws.conditional_formatting) >= 2
    if ws.max_row >= DATA_START_ROW:
        assert ws.cell(row=DATA_START_ROW, column=26).value is not None
