"""Generate workbook QA report data for docs/FINAL_WORKBOOK_QA.md."""

from __future__ import annotations

import json
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.chart import BarChart

REQUIRED_SHEETS = [
    "README",
    "Master Inventory",
    "Inventory Dashboard",
    "Inventory Classification",
    "Aged Excess Analysis",
    "Cycle Count Plan",
    "Replenishment Planning",
    "Transfer Planner",
    "Markdown Planner",
    "Demand Forecast",
    "Service Level Analysis",
    "Purchase Order Tracker",
    "Vendor Scorecards",
    "Management Summary",
]

WORKBOOK = (
    Path(__file__).resolve().parents[1] / "dist" / "Inventory_Optimization_Copilot.xlsx"
)


def _chart_title(chart) -> str:
    if chart.title and chart.title.tx and chart.title.tx.rich:
        return chart.title.tx.rich.paragraphs[0].r[0].t
    return str(chart.title)


def _table_stats(ws) -> dict:
    tables = []
    for name, table in ws.tables.items():
        ref = table if isinstance(table, str) else table.ref
        min_row = int("".join(c for c in ref.split(":")[0] if c.isdigit()))
        max_row = int("".join(c for c in ref.split(":")[1] if c.isdigit()))
        hidden = sum(
            1
            for r in range(min_row, max_row + 1)
            if ws.row_dimensions[r].hidden or ws.row_dimensions[r].height in (0, 0.0)
        )
        tables.append(
            {
                "name": name,
                "ref": ref,
                "rows": max_row - min_row,
                "hidden_rows": hidden,
            }
        )
    return {"count": len(tables), "tables": tables}


def _chart_stats(chart) -> dict:
    legend_pos = chart.legend.position if chart.legend else None
    legend_overlay = chart.legend.overlay if chart.legend else None
    labels = getattr(chart, "dataLabels", None)
    info = {
        "title": _chart_title(chart),
        "type": chart.type if isinstance(chart, BarChart) else type(chart).__name__,
        "legend_position": legend_pos,
        "legend_overlay": legend_overlay,
        "data_labels_value_only": (
            labels.showVal and not labels.showSerName and not labels.showCatName
            if labels
            else None
        ),
    }
    if isinstance(chart, BarChart):
        info["category_orientation"] = (
            chart.y_axis.scaling.orientation if chart.type == "bar" else None
        )
        info["x_axis_title"] = str(chart.x_axis.title)
        info["y_axis_title"] = str(chart.y_axis.title)
    return info


def inspect_workbook(path: Path) -> dict:
    wb = load_workbook(path)
    report: dict = {
        "path": str(path),
        "size_bytes": path.stat().st_size,
        "sheet_count": len(wb.sheetnames),
        "sheets": [],
        "charts": [],
        "round_trip_ok": False,
    }

    for sheet_name in REQUIRED_SHEETS:
        assert sheet_name in wb.sheetnames
        ws = wb[sheet_name]
        tstats = _table_stats(ws)
        charts = [_chart_stats(c) for c in ws._charts]
        report["sheets"].append(
            {
                "name": sheet_name,
                "tables_inspected": tstats["count"],
                "table_rows": sum(t["rows"] for t in tstats["tables"]),
                "charts_inspected": len(charts),
                "max_col": ws.max_column,
                "print_area": ws.print_area,
            }
        )
        for c in charts:
            c["sheet"] = sheet_name
            report["charts"].append(c)

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp_path = Path(tmp.name)
    try:
        wb.save(tmp_path)
        reloaded = load_workbook(tmp_path)
        reloaded.save(tmp_path)
        report["round_trip_ok"] = len(reloaded.sheetnames) == 14
    finally:
        tmp_path.unlink(missing_ok=True)

    return report


def main() -> None:
    report = inspect_workbook(WORKBOOK)
    out = Path(__file__).resolve().parents[1] / "docs" / "final_workbook_qa_data.json"
    out.write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(
        json.dumps(
            {"sheets": len(report["sheets"]), "charts": len(report["charts"])}, indent=2
        )
    )


if __name__ == "__main__":
    main()
