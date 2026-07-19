#!/usr/bin/env python3
"""Render Inventory Dashboard ranked charts for visual QA (matplotlib proxy)."""

from __future__ import annotations

from pathlib import Path

import matplotlib.pyplot as plt
from openpyxl import load_workbook

OUT = Path(__file__).resolve().parents[1] / "docs" / "screenshots" / "workbook-final-qa"
WORKBOOK = (
    Path(__file__).resolve().parents[1] / "dist" / "Inventory_Optimization_Copilot.xlsx"
)


def _section_rows(ws, title: str) -> tuple[int, int, int]:
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 1).value == title:
            header = r + 2
            first = header + 1
            row = first
            while ws.cell(row, 1).value is not None:
                row += 1
            return header, first, row - 1
    raise ValueError(title)


def _render_horizontal(
    title: str, labels: list, values: list, x_label: str, out: Path
) -> None:
    fig, ax = plt.subplots(figsize=(10, 6))
    y_pos = range(len(labels))
    ax.barh(list(y_pos), values, color="#1f4e79")
    ax.set_yticks(list(y_pos))
    ax.set_yticklabels(labels, fontsize=8)
    ax.invert_yaxis()
    ax.set_xlabel(x_label)
    ax.set_title(title)
    ax.grid(axis="x", alpha=0.3)
    fig.tight_layout()
    fig.savefig(out, dpi=120)
    plt.close(fig)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(WORKBOOK)
    ws = wb["Inventory Dashboard"]

    _, first, last = _section_rows(ws, "Top 10 Excess Inventory Items")
    labels = [ws.cell(r, 8).value for r in range(first, last + 1)]
    values = [float(ws.cell(r, 6).value or 0) for r in range(first, last + 1)]
    _render_horizontal(
        "Top 10 Excess Inventory Items",
        labels,
        values,
        "Excess Inventory Value ($)",
        OUT / "top_10_excess_inventory_items.png",
    )

    _, first, last = _section_rows(ws, "Top Transfer Opportunities by Net Benefit")
    labels = [ws.cell(r, 8).value for r in range(first, last + 1)]
    values = [float(ws.cell(r, 6).value or 0) for r in range(first, last + 1)]
    _render_horizontal(
        "Top Transfer Opportunities by Net Benefit",
        labels,
        values,
        "Net Benefit ($)",
        OUT / "top_transfer_opportunities.png",
    )

    _, first, last = _section_rows(ws, "Replenishment Requirements by Status")
    labels = [ws.cell(r, 1).value for r in range(first, last + 1)]
    values = [float(ws.cell(r, 4).value or 0) for r in range(first, last + 1)]
    _render_horizontal(
        "Replenishment Requirements by Status",
        labels,
        values,
        "Recommended Order Value ($)",
        OUT / "replenishment_requirements_by_status.png",
    )
    print(f"Saved QA renders to {OUT}")


if __name__ == "__main__":
    main()
